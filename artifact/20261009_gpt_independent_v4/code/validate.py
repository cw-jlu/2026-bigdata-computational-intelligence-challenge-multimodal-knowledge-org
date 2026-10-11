"""Independent packaging/schema/date checks. Not a content-accuracy score.

Does not import generation, matching, or schema constants from the builder.
Manual raw evidence labels are checked separately by check_cases.py.
"""
import argparse,hashlib,json,re,zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from PIL import Image
from docx import Document
from openpyxl import load_workbook

DATES=['2025.5.10','2025.5.18','2025.5.30','2025.6.9','2025.6.17','2025.6.29','2025.7.9','2025.7.17']
HEADERS={
 '飞机':['日期','星期','国家地区','机型','批号','批次','架次','起飞基地','呼号','起飞时间','降落时间','活动方向','活动区域','任务描述','任务类型','备注'],
 '舰艇':['日期','国家地区','舰船','批号','艘次','活动时间','活动方向','活动区域','任务','母港','备注']}
PNGS=['01_舰机侦巡数量曲线图.png','02_中国主要机舰活动示意图.png','03_美国主要机舰活动示意图.png','04_其他国家主要机舰活动示意图.png']

def run(inp,out):
    errors=[];records=[];expected=set();known={}
    for line in (out/'evidence/all_events.jsonl').read_text(encoding='utf-8').splitlines():
        e=json.loads(line);known[e['id']]=e
    root=inp.parents[1];source_atoms=[]
    for day in DATES:
        prefix=f'日报/{day}/';files=[f'{day}日报.docx',f'{day}机舰活动明细.xlsx']+['附件图表/'+s for s in PNGS]
        expected.update(prefix+x for x in files);d=out/'日报'/day;local=[];ee=[e for e in known.values() if e['date']==day and e['decision']=='纳入']
        for e in ee:
            if len(e['claims'])!=2 or e['issues']:local.append('纳入记录违反自身证据门控 '+e['id'])
            for f in ['country','key','quantity','task','start','end','region']:
                a=e['field_audit'][f]
                if len(a['support_types'])<2 or a['conflicts']:local.append('关键字段支持不足 '+e['id']+' '+f)
            for a in e['claims'].values():
                source_atoms.append(a['id']);p=root/a['path']
                if a['date']!=day:local.append('文本日期泄露 '+a['path'])
                if hashlib.sha256(p.read_bytes()).hexdigest()!=a['sha256']:local.append('源文件哈希改变')
        try:
            w=load_workbook(d/f'{day}机舰活动明细.xlsx',data_only=False)
            if w.sheetnames!=['飞机','舰艇']:local.append('工作表结构不符')
            workbook_ids=[];counts={'飞机':0,'舰艇':0}
            for name,h in HEADERS.items():
                s=w[name]
                if [c.value for c in s[1]]!=h or s.max_column!=len(h):local.append(name+'字段或顺序不符')
                for row in list(s.values)[1:]:
                    r=dict(zip(h,row));v=r['日期']
                    if not isinstance(v,datetime) or v.date()!=datetime.strptime(day,'%Y.%m.%d').date():local.append('XLSX日期或数据类型不符')
                    key='架次' if name=='飞机' else '艘次';q=r[key]
                    if not isinstance(q,(int,float)) or q<1 or int(q)!=q:local.append('数量不是正整数')
                    counts[name]+=q or 0
                    eid=re.search(r'2025\.\d{1,2}\.\d{1,2}-(?:AIR|SHIP)-[A-F0-9]+',r['备注'] or '')
                    if not eid or eid[0] not in known:local.append('明细缺少可定位事件ID');continue
                    e=known[eid[0]];workbook_ids.append(eid[0]);f=e['facts']
                    if r['国家地区']!=f['country'] or r['活動方向' if '活動方向' in r else '活动方向']!=f['direction'] or r['活动区域']!=f['region'] or q!=f['quantity']:local.append('明细与证据不一致')
                    if name=='飞机' and (r['起飞时间'] is not None or r['降落时间'] is not None):local.append('无证据将观测时段写入起降字段')
            if sorted(workbook_ids)!=sorted(e['id'] for e in ee):local.append('XLSX活动集合或去重不符')
            w.close()
            doc=Document(d/f'{day}日报.docx');texts=[p.text for p in doc.paragraphs];txt='\n'.join(texts)
            m,dd=map(int,day.split('.')[1:])
            if texts[0]!=f'{m}月{dd}日飞机舰船活动情况':local.append('DOCX标题不符')
            for heading in ['一、重要军事动向','二、重点态势分析','三、外（台）机舰活动情况','四、美重要武器平台部署情况','五、次日关注重点']:
                if heading not in texts:local.append('缺章节 '+heading)
            if len(doc.inline_shapes)!=4:local.append('DOCX附图不齐')
            in_trend=False
            for t in texts:
                if t=='二、重点态势分析':in_trend=True
                if t=='三、外（台）机舰活动情况':in_trend=False
                for ds in re.findall(r'2025\.\d{1,2}\.\d{1,2}',t):
                    dt=datetime.strptime(ds,'%Y.%m.%d')
                    if dt>datetime.strptime(day,'%Y.%m.%d'):local.append('DOCX未来数据')
                    if ds!=day and not in_trend:local.append('DOCX非主线章节用了跨日数据')
            for i in range(1,5):
                if f'图{i}' not in txt:local.append('附图未引用')
            if f'飞机{counts["飞机"]}架次、舰艇{counts["舰艇"]}艘次' not in txt:local.append('正文合计与明细不一致')
            for f in PNGS:
                path=d/'附件图表'/f
                with Image.open(path) as img:
                    if img.format!='PNG' or min(img.size)<600:local.append('PNG尺寸或格式不符')
                    img.verify()
            chart=json.loads((out/'payloads'/(day+'.chart_data.json')).read_text(encoding='utf-8'))
            if sorted(chart['diagrams'])!=sorted(e['id'] for e in ee):local.append('图表数据集合不符')
            if any(c['date']!=day for c in chart['curve']):local.append('曲线用了非当日数据')
            for kind in ['飞机','舰艇']:
                if sum(c['quantity'] for c in chart['curve'] if c['kind']==kind)!=counts[kind]:local.append('图表数量与明细不一致')
            p=json.loads((out/'payloads'/(day+'.json')).read_text(encoding='utf-8'));chain=p['chain']
            if not chain:local.append('无满足三日规则的主线')
            else:
                if len(set(chain['dates']))<3 or len(chain['stages']) not in [2,3]:local.append('主线日期或阶段数量不符')
                if day not in chain['dates'] or any(datetime.strptime(s,'%Y.%m.%d')>datetime.strptime(day,'%Y.%m.%d') for s in chain['dates']):local.append('主线不在节点日或泄露未来')
                for eid in chain['events']:
                    if known[eid]['decision']!='纳入':local.append('主线使用排除活动')
                if max(chain['gaps'])>=15:local.append('主线跨越15日断点')
        except Exception as exc:local.append(repr(exc))
        errors.extend(day+' '+x for x in local);records.append({'date':day,'errors':local,'included':len(ee),'structural_checks':'DOCX/XLSX/PNG/日界/交叉合计/主线条件'})
    if len(source_atoms)!=len(set(source_atoms)):errors.append('纳入记录重复使用同一文本原子事实')
    actual={p.relative_to(out).as_posix() for p in (out/'日报').rglob('*') if p.is_file()}
    if actual!=expected:errors.append('日报实体目录文件集合不符 '+str(sorted(actual.symmetric_difference(expected))))
    # Package only six required files per day, with POSIX names and deterministic timestamps.
    archive=out/'赛题2_初赛_八节点_提交.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for name in sorted(expected):
            info=zipfile.ZipInfo(name,(2026,10,9,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;z.writestr(info,(out/name).read_bytes())
    with zipfile.ZipFile(archive) as z:
        if set(z.namelist())!=expected or len(z.namelist())!=48:errors.append('ZIP内容或数量不符')
        if z.testzip() is not None:errors.append('ZIP CRC校验失败')
        for name in z.namelist():
            if '\\' in name or '..' in name.split('/') or name.endswith('.pdf'):errors.append('ZIP路径或类型不符')
            if z.read(name)!=(out/name).read_bytes():errors.append('ZIP字节与最终输出不符')
    result={'ok':not errors,'errors':errors,'nodes':records,'zip_entries':48,'zip_sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),
            'limitation':'文件结构与内部一致性验证；不等于内容真值验证或官方评测。独立原始案例见audit/cases_final.json。'}
    (out/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2))
    if errors:raise SystemExit(1)
    return result

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--input',type=Path,required=True);a.add_argument('--output',type=Path,required=True);x=a.parse_args();run(x.input.resolve(),x.output.resolve())
