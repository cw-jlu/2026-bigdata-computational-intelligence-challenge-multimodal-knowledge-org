"""Author Word documents with bundled Python, using finalized payloads only."""
import argparse,json
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

def citation(e):
    ss=[]
    for s,a in e['claims'].items():ss.append(('军情' if s=='开源军情' else '舆论')+Path(a['path']).name.split('_')[0])
    rr=','.join(str(r['row']) for r in e['rows'])
    return f'{e.get("ref",e["batch"])}：'+ '、'.join(ss)+f'；雷达{e["sheet"]}行{rr}'

def event_sentence(e,reference=True):
    f=e['facts'];unit='架次' if e['kind']=='飞机' else '艘次'
    return f'{f["country"]}{e["platform"]}{f["quantity"]}{unit}，{f["start"]}至{f["end"]}在{f["region"]}执行{f["task"]}。'+(f'（{e["ref"]}）' if reference else '')

def explain_issue(e):
    """Plain Chinese from field audit; full raw conflicting values stay in evidence."""
    issue=e['issues'][0]
    for field,label in [('country','主体'),('key','型号'),('quantity','数量'),('task','任务'),('start','开始时间'),('end','结束时间')]:
        if '关键字段'+field in issue:
            audit=e['field_audit'][field]
            if audit['conflicts']:
                details='、'.join(f'{source}记为{value}' for source,value in audit['values'].items())
                return label+'存在明确冲突（'+details+'）'
            return label+'尚不足两类来源支持'
    if '关键区域冲突' in issue:
        conflicts=e['field_audit']['region']['conflicts']
        details=[]
        for item in conflicts:
            if item.get('region'):details.append(item['source']+'记为'+item['region'])
            else:details.append(item['source']+'轨迹与文本具体区域不相容')
        return '区域存在明确冲突（'+'、'.join(details)+'）'
    return issue

def style(doc):
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    for st in doc.styles:
        for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(5);doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 2'].font.size=Pt(11)
    doc.styles['Caption'].font.bold=False;doc.styles['Caption'].font.size=Pt(9)
    for n in ['Heading 1','Heading 2']:doc.styles[n].paragraph_format.keep_with_next=True
    doc.core_properties.author='GPT independent v4';doc.core_properties.created=__import__('datetime').datetime(2026,10,9,0,0,0)
    doc.core_properties.modified=doc.core_properties.created

def write(p,out):
    doc=Document();style(doc);day=p['day'];m,d=map(int,day.split('.')[1:]);events=p['events'];c=p['counts'];ch=p['chain']
    title=doc.add_paragraph(f'{m}月{d}日飞机舰船活动情况',style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    seas='、'.join(f'{s}{c["directions"].get(s,0)}起' for s in ['东海','南海','黄海'])
    doc.add_paragraph(f'当日活动总览：三源核验后纳入{c["included"]}起活动，飞机{c["air"]}架次、舰艇{c["ship"]}艘次，分布为{seas}。另有{c["excluded"]}项雷达候选因关键字段冲突、单源事实或无法唯一对应而排除，{c["outside"]}项完成核验但位于三海以外。纳入数量是可核实记录合计，不代表全海域实际兵力总量。')
    doc.add_paragraph('数量按事件记录计次，同型平台多次活动分别统计。图1展示当日开始小时的数量分布；图2为中国、图3为美国、图4为其他国家及台湾的活动时段与具体区域示意。活动观测起止不等同于起飞降落；单源基地和母港线索不填写为核实事实。')
    doc.add_heading('一、重要军事动向',level=1)
    important=sorted(events,key=lambda e:(-int(any(x in e['platform'] for x in ['航空母舰','攻击核潜艇','战略轰炸机','电子侦察','预警','导弹驱逐舰'])),-e['facts']['quantity'],e['facts']['start']))[:8]
    for e in important:doc.add_paragraph('核实事实：'+event_sentence(e))
    missions=Counter(e['facts']['task'] for e in events);top=missions.most_common(2)
    doc.add_paragraph('分析性判断：当日可信记录主要涉及'+ '、'.join(f'{k}{v}起' for k,v in top)+'。任务种类反映已核实活动构成；未据单一侦巡或训练记录推断作战意图。')
    doc.add_heading('二、重点态势分析',level=1)
    if ch:
        regions={e['facts']['region'] for e in ch['evidence']};seas=set(e['facts']['direction'] for e in ch['evidence'])
        pattern='跨海域活动区域变化与间歇出现' if len(seas)>1 else ('同海域不同具体区域的间歇活动' if len(regions)>1 else '相同区域的重复活动线索')
        doc.add_paragraph(f'总体判断：截至{day}，{ch["country"]}{ch["platform"]}主线覆盖{len(ch["dates"])}个不同日期，表现为{pattern}。证据日期为'+ '、'.join(ch['dates'])+f'；相邻证据最大间隔{max(ch["gaps"])}日。'+ch['identity_scope']+'。未出现证据的日期不能认定为连续在场，也不能据缺报确认离开。')
        for i,g in enumerate(ch['stages'],1):
            ss=Counter(e['facts']['region'] for e in g['evidence']);tt=Counter(e['facts']['task'] for e in g['evidence']);qs=sum(e['facts']['quantity'] for e in g['evidence']);hours=[e['facts']['start']+'—'+e['facts']['end'] for e in g['evidence']]
            prev=ch['stages'][i-2] if i>1 else None
            if prev:
                ps=set(e['facts']['direction'] for e in prev['evidence']);ns=set(e['facts']['direction'] for e in g['evidence']);pt=set(e['facts']['task'] for e in prev['evidence']);nt=set(tt)
                gap=(__import__('datetime').datetime.strptime(g['dates'][0],'%Y.%m.%d')-__import__('datetime').datetime.strptime(prev['dates'][-1],'%Y.%m.%d')).days
                change=('活动方向由'+ '、'.join(sorted(ps))+'转至'+ '、'.join(sorted(ns))) if ps!=ns else ('任务由'+'、'.join(sorted(pt))+'变化为'+'、'.join(sorted(nt)) if pt!=nt else ('具体区域由'+'、'.join(sorted({e['facts']['region'] for e in prev['evidence']}))+'变化为'+'、'.join(sorted(ss)) if {e['facts']['region'] for e in prev['evidence']}!=set(ss) else '区域任务组合在间隔后重复出现'))
                change+=f'，与前阶段证据相隔{gap}日。'
            else:change='形成该段主线的起始核实记录。'
            doc.add_paragraph(f'第{i}阶段（{g["dates"][0]}至{g["dates"][-1]}）：{change}具体区域为'+ '、'.join(f'{k}{v}起' for k,v in ss.items())+'，任务为'+'、'.join(f'{k}{v}起' for k,v in tt.items())+f'，合计{qs}'+('架次' if g['evidence'][0]['kind']=='飞机' else '艘次')+'；观测时段包括'+ '、'.join(dict.fromkeys(hours))+'。')
            for e in g['evidence']:doc.add_paragraph(f'主线证据 {e["date"]}：'+event_sentence(e,False)+' '+citation(e),style='Caption')
        doc.add_paragraph('分析边界：阶段依据实际区域、任务组合与证据间隔划分。任务转换可以描述记录中的活动类型变化，但缺少跨日同机体或舰体标识，不能确认具体平台航迹持续连通。')
    else:doc.add_paragraph('截至节点日，同一外方主体同型平台没有满足至少三个证据日期且时间间隔可衔接的可信主线，不作确定的跨日阶段判断。')
    doc.add_heading('三、外（台）机舰活动情况',level=1)
    for sea in ['东海','南海','黄海']:
        foreign=[e for e in events if e['facts']['country']!='中国' and e['facts']['direction']==sea]
        doc.add_heading(sea,level=2)
        if not foreign:doc.add_paragraph('当日无满足三源核验的外（台）记录。')
        for e in foreign:doc.add_paragraph(event_sentence(e))
    doc.add_heading('四、美重要武器平台部署情况',level=1)
    us=[e for e in events if e['facts']['country']=='美国']
    if us:
        for e in us:doc.add_paragraph(event_sentence(e))
        doc.add_paragraph('部署判断限于上述当日活动时段和区域；基地、编队组成与连续部署期如缺乏两类来源支持，不扩大为确定结论。参见图3。')
    else:doc.add_paragraph('当日无满足三源核验的美国三海活动记录，不能由此认定美国平台未在相关区域活动。参见图3。')
    doc.add_heading('五、次日关注重点',level=1)
    seeds=us or [e for e in events if e['facts']['country']!='中国']
    for e in seeds[:3]:
        f=e['facts'];doc.add_paragraph(f'基于当日{e["ref"]}，关注{f["country"]}{e["platform"]}在{f["region"]}是否再次执行{f["task"]}，核对新记录与{f["start"]}—{f["end"]}时段的衔接，以及数量是否仍为{f["quantity"]}。这是待验证关注项，不是次日事实或行动预测。')
    if p['excluded']:
        e=p['excluded'][0];reason=explain_issue(e)
        doc.add_paragraph(f'核查{e["platform"]}候选的证据缺口：'+reason+'。只有补足关联或消解明确冲突后才可纳入后续统计。')
    doc.add_heading('当日证据索引',level=1)
    doc.add_paragraph('索引的军情、舆论编号仅定位当日原始文件；不参与事件关联。批号与完整路径、证据片段、候选及排除理由保存于独立中间证据。')
    for e in events:doc.add_paragraph(citation(e)+'；'+e['platform']+'；'+e['batch'],style='Caption')
    # Predictable pages keep figures and their captions together.
    for i,name in enumerate(p['charts'],1):
        doc.add_page_break();doc.add_heading(f'图{i} '+name[3:-4],level=1)
        pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.add_run().add_picture(str(out/'日报'/day/'附件图表'/name),width=Cm(16.5))
        doc.add_paragraph(f'图{i}仅采用{day}的纳入活动。'+('每个事件在开始小时计入一次，飞机与舰艇分别使用架次、艘次。' if i==1 else '横轴为当日小时，条目给出平台、区域、任务与数量；线段表示活动证据时段，不表示未经核实的地理航线。'),style='Caption')
    dest=out/'日报'/day/f'{day}日报.docx';doc.save(dest)

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--output',type=Path,required=True);x=a.parse_args()
    for f in sorted((x.output/'payloads').glob('*.json')):
        if '.chart_data' not in f.name:write(json.loads(f.read_text(encoding='utf-8')),x.output)
