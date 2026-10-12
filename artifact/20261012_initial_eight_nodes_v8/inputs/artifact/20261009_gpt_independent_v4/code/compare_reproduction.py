"""Compare clean-environment results; do not use generation helpers."""
import argparse, json, hashlib
from pathlib import Path
from docx import Document
from PIL import Image
from io import BytesIO
import zipfile
from openpyxl import load_workbook

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()

def compare(first, second):
    checks=[]
    for p in sorted((first/'evidence').glob('*.jsonl')):
        q=second/'evidence'/p.name
        checks.append({'path':p.relative_to(first).as_posix(),'kind':'raw evidence bytes','equal':q.exists() and sha(p)==sha(q)})
    for p in sorted((first/'payloads').glob('*.json')):
        q=second/'payloads'/p.name
        checks.append({'path':p.relative_to(first).as_posix(),'kind':'payload values','equal':json.loads(p.read_text(encoding='utf-8'))==json.loads(q.read_text(encoding='utf-8'))})
    for p in sorted((first/'日报').rglob('*')):
        if not p.is_file():continue
        q=second/p.relative_to(first)
        if p.suffix=='.docx':
            # Container timestamps are not document content.
            left=Document(p);right=Document(q)
            a=[s.text for s in left.paragraphs];b=[s.text for s in right.paragraphs]
            eq=a==b and len(left.inline_shapes)==len(right.inline_shapes)==4
            with zipfile.ZipFile(p) as zl, zipfile.ZipFile(q) as zr:
                media=sorted(n for n in zl.namelist() if n.startswith('word/media/'))
                for n in media:
                    ia=Image.open(BytesIO(zl.read(n))).convert('RGBA');ib=Image.open(BytesIO(zr.read(n))).convert('RGBA')
                    eq=eq and ia.size==ib.size and ia.tobytes()==ib.tobytes()
            kind='DOCX paragraphs + decoded pixels of four embedded figures' 
        elif p.suffix=='.xlsx':
            left=load_workbook(p);right=load_workbook(q)
            eq=left.sheetnames==right.sheetnames and all(list(left[n].values)==list(right[n].values) for n in left.sheetnames)
            left.close();right.close();kind='XLSX typed values and sheet order'
        else:
            a=Image.open(p).convert('RGBA');b=Image.open(q).convert('RGBA')
            eq=a.size==b.size and a.tobytes()==b.tobytes();kind='PNG decoded pixels'
            checks.append({'path':p.relative_to(first).as_posix(),'kind':'PNG compression/container bytes (informational)','bytes_equal':sha(p)==sha(q),'equal':True})
        checks.append({'path':p.relative_to(first).as_posix(),'kind':kind,'equal':eq})
    return {'ok':all(c['equal'] for c in checks),'checks':checks,'failures':[c for c in checks if not c['equal']],
            'limitation':'复现稳定性，不是内容准确性；DOCX/XLSX容器时间戳及PNG压缩流不要求字节相同；图片逐像素比较。'}

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--first',type=Path,required=True);a.add_argument('--second',type=Path,required=True);a.add_argument('--record',type=Path,required=True);x=a.parse_args()
    r=compare(x.first,x.second);x.record.write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'ok':r['ok'],'checks':len(r['checks']),'failures':r['failures']},ensure_ascii=False))
    if not r['ok']:raise SystemExit(1)
