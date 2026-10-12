"""Review metadata, clean-run equivalence and isolated validator fault probes."""
import hashlib,json,shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree
from openpyxl import load_workbook
from validate import check,DATES

V=Path(__file__).resolve().parents[1]
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def save(p,x):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def metadata():
    config=read(V/'config/experiment.json');config['parameters']['time_endpoint_tolerance_minutes']=1
    config['amendment']='原冻结记录误写时间容差5分钟；代码实际始终为1分钟，未修改该参数。原记录保留。'
    config['runtime_dependencies']=read(V/'audit/full_run.json')['dependencies']
    config['source_sha256']={p.name:sha(p) for p in sorted((V/'code').glob('*')) if p.is_file()}
    save(V/'config/experiment_final.json',config)
    first=read(V/'audit/new_development_cases_initial.json')['cases'];final=read(V/'audit/new_development_cases.json')['cases']
    wrong=[c for c in first if c['label_status']!='human_reviewed']
    save(V/'audit/annotation_changes.json',{'old_42_labels':'沿用v2人工标签及已保留的初稿；本次未更改','new_development_locator_correction':wrong,'corrected_case':[c for c in final if c['date']=='2025.5.30' and c['batch']=='AIR-D259B796A74070'],'timing':'定位错误在首次新案例评估前发现；初稿与原预期保留；不能计作算法失败或修复收益。后加开发案例见最终冻结，均属于开发回归。','reserved':'24例冻结后只评估，未用于事实修复；未宣称IID独立测试；与开发有同日及原文重叠。'})
    diagnostic=read(V/'audit/diagnostic_raw_bundles.json')
    save(V/'audit/sampling_report.json',{'v2_diagnostic_count':len(diagnostic),'diagnostic_strata':dict(Counter((b['baseline']['decision']+'/'+('飞机' if b['batch'].startswith('AIR') else '舰船')) for b in diagnostic)),'diagnostic_nodes':dict(Counter(b['date'] for b in diagnostic)),'old_development':42,'new_development':26,'reserved':24,'selection':'八节点v2纳入/排除，各抽普通飞机/舰船；增加同型同日、数量/任务/时间/细区冲突、未知辅助字段及竞争轨迹。预期从完整原始记录人工录入，v2输出仅决定抽样分层，不作为真值。','limitations':'目的抽样，无概率总体代表性；人工预期可错；开发和预留有同日及文件重叠；只报告分母内结果，不外推全量正确率。'})
    tracks=[t for t in map(json.loads,(V/'output/中间证据/radar_tracks.jsonl').open(encoding='utf-8')) if t['date'] in DATES]
    edge=[t for t in tracks if t['competing_tracks']>1 or t['jumps'] or t['anomalies'] or len(t['region_values'])>1]
    save(V/'audit/trajectory_review.json',{'scope':'8节点所有雷达候选完整记录；下列边界组保留全部行、顺序与点位，不以中位点或多数消解持续冲突。','candidate_tracks':len(tracks),'edge_tracks':len(edge),'engineering_tolerances':{'anchor_branch_km':80,'same_time_km':5,'jump_speed_kmh':2400,'window_endpoint_minutes':1},'basis':'保持v2容差作宽松可行性筛查；80km只区分初始位置候选，不能证明细区；2400km/h不是型号实测速度界限；1分钟是源记录分钟精度容差。本次未用调整阈值改善样本。','noise_rule':'舰船孤立区域标签只有在同批其他完整字段和时间一致、两类文本共同指向主区域且其余行区域一致时可解释为记录噪声；持续两个区域或飞机连续分支不作噪点删除。无效时刻/越界点缺独立隔离依据时维持冲突。未建立经过独立噪声标签验证的通用飞机去噪算法。','boundaries':edge})
    save(V/'audit/official_sample_search.json',{'date':'2026-10-09','queries':['多模态融合的知识组织与决策生成 2026 样例','2026 大数据与计算智能挑战赛 赛题2 初赛 日报'],'public_pages':[{'url':'https://d2.nudt.edu.cn/tzs/','finding':'公开首页赛事公告及联系信息，未取得可下载DOCX/XLSX官方成品样例'},{'url':'https://finance.sina.com.cn/wm/2026-08-25/doc-inipnyqs9723875.shtml','finding':'找到赛题名称及赛事官网入口；不是格式样例'}],'local_scope':'README、docs/01—07与原始数据结构；未读取来源不明提交目录成品','obtained_sample':None,'limitation':'只说明本次检索未取得样例，不声称官网登录区或其他位置不存在样例。依据docs/04逐项核验。'})

def clean_compare():
    primary=V/'output';clean=V/'audit/repro_final/output';results=[]
    for name in ['events.jsonl','claims.jsonl','radar_tracks.jsonl','candidate_links.jsonl','field_evidence.jsonl','source_coverage.jsonl','unallocated_claims.jsonl','chart_datasets.json','workbook_rows.json']:
        a=primary/'中间证据'/name;b=clean/'中间证据'/name;results.append({'file':name,'equal':sha(a)==sha(b),'primary_sha256':sha(a),'clean_sha256':sha(b)})
    results.append({'file':'raw_snapshot.jsonl','equal':sha(V/'audit/raw_snapshot.jsonl')==sha(V/'audit/repro_final/raw_snapshot.jsonl')})
    for day in DATES:
        for name in [day+'日报.docx',day+'机舰活动明细.xlsx']:
            a=primary/'日报'/day/name;b=clean/'日报'/day/name
            if name.endswith('.docx'):
                def text(p):
                    with ZipFile(p) as z:return [''.join(x.itertext()) for x in etree.fromstring(z.read('word/document.xml')).findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')]
                equal=text(a)==text(b)
            else:
                def cells(p):
                    w=load_workbook(p,data_only=False);v={s.title:[list(r) for r in s.iter_rows(values_only=True)] for s in w};w.close();return v
                equal=cells(a)==cells(b)
            results.append({'file':day+'/'+name,'equal':equal,'scope':'DOCX所有正文/表格文本，或XLSX所有工作表全部单元格；包元数据及不同Excel引擎样式不作字节相等要求'})
        for p in sorted((primary/'日报'/day/'附件图表').glob('*.png')):results.append({'file':day+'/'+p.name,'equal':sha(p)==sha(clean/'日报'/day/'附件图表'/p.name)})
    result={'all_equal':all(r['equal'] for r in results),'comparisons':results,'clean_run':'audit/repro_final/full_run.json','note':'主交付XLSX使用artifact-tool，纯Python复现使用openpyxl；事实、全部单元格、报告文本和PNG一致。可视化审核针对主交付文件。'}
    save(V/'audit/clean_reproduction_comparison.json',result)
    if not result['all_equal']:raise SystemExit('clean reproduction differs')

def validator_probes():
    root=V/'audit/validator_probes';root.mkdir(exist_ok=True);results=[]
    for mode in ['header','date','curve_quantity','chart_field','future_history','missing_reference']:
        out=root/mode;shutil.copytree(V/'output/日报',out/'日报',dirs_exist_ok=True);inter=out/'中间证据';inter.mkdir(exist_ok=True)
        for p in (V/'output/中间证据').glob('factpack_*.json'):shutil.copy2(p,inter/p.name)
        for name in ['chart_datasets.json','workbook_rows.json']:shutil.copy2(V/'output/中间证据'/name,inter/name)
        book=out/'日报/2025.5.10/2025.5.10机舰活动明细.xlsx'
        if mode in ['header','date']:
            w=load_workbook(book)
            if mode=='header':w['飞机']['C1']='错误主体列'
            else:w['飞机']['A2']=datetime(2025,5,11)
            w.save(book);w.close()
        elif mode in ['curve_quantity','chart_field']:
            x=read(inter/'chart_datasets.json');c=next(c for c in x if c['date']=='2025.5.10' and c['filename'].startswith('01' if mode=='curve_quantity' else '02'))
            if mode=='curve_quantity':c['air'][0]+=1
            else:c['records'][0]['quantity']+=1
            save(inter/'chart_datasets.json',x)
        elif mode=='future_history':
            p=inter/'factpack_2025.5.10.json';x=read(p);x['trend']['evidence'][0]['date']='2025.5.11';save(p,x)
        else:
            p=out/'日报/2025.5.10/2025.5.10日报.docx'
            with ZipFile(p) as z:parts={n:z.read(n) for n in z.namelist()}
            parts['word/document.xml']=parts['word/document.xml'].decode('utf-8').replace('附图4','附件四').encode('utf-8')
            with ZipFile(p,'w',ZIP_DEFLATED) as z:
                for n,b in parts.items():z.writestr(n,b)
        r=check(out,V/'audit/raw_snapshot.jsonl');results.append({'probe':mode,'rejected':not r['ok'],'errors':[e for n in r['nodes'] for e in n['errors']]})
    save(V/'audit/validator_probe_results.json',{'probes':results,'all_rejected':all(r['rejected'] for r in results),'scope':'仅复制到新版本audit子目录的故意错误副本；未改动交付或基线。校验器独立字面schema、日期及重算统计能检出列/日期/数量/字段/未来历史/正文引用错误。'})
    if not all(r['rejected'] for r in results):raise SystemExit('validator probe unexpectedly passed')

if __name__=='__main__':
    import sys
    action=sys.argv[1] if len(sys.argv)>1 else 'metadata'
    {'metadata':metadata,'clean':clean_compare,'probes':validator_probes}[action]()
    print(action+' complete')
