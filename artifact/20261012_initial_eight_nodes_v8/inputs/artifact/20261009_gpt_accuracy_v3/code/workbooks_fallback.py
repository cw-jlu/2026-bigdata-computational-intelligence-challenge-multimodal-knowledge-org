"""Portable fallback used ONLY when @oai/artifact-tool is unavailable.

Main delivered workbooks were authored by the bundled JS artifact tool. This
fallback preserves the official data schema for machines without that package.
"""
import json
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo

source,target=map(Path,sys.argv[1:3])
for book in json.loads(source.read_text(encoding='utf-8')):
    wb=Workbook(); wb.remove(wb.active)
    for data in book['sheets']:
        ws=wb.create_sheet(data['name']); ws.append(data['headers'])
        for row in data['rows']:
            r=list(row); r[0]=datetime.strptime(r[0],'%Y-%m-%d');ws.append(r);ws.cell(ws.max_row,1).number_format='yyyy.mm.dd'
        ws.sheet_view.showGridLines=False;ws.freeze_panes='D2'
        for c in ws[1]: c.fill=PatternFill('solid',fgColor='294C67');c.font=Font(name='Microsoft YaHei',size=10,bold=True,color='FFFFFF');c.alignment=Alignment(horizontal='center',vertical='center')
        ws.row_dimensions[1].height=30
        widths=[14,12,14,36,28,9,9,18,14,14,14,12,24,24,24,70] if data['name']=='飞机' else [14,14,36,29,10,20,12,25,25,20,70]
        for col,w in enumerate(widths,1):ws.column_dimensions[ws.cell(1,col).column_letter].width=w
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=82
            for c in row:c.font=Font(name='Microsoft YaHei',size=10);c.alignment=Alignment(vertical='center',wrap_text=True)
        table=Table(displayName='AircraftActivities' if data['name']=='飞机' else 'ShipActivities',ref=ws.dimensions)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    dest=target/book['date']/(book['date']+'机舰活动明细.xlsx');dest.parent.mkdir(parents=True,exist_ok=True);wb.save(dest)
    print('Portable XLSX fallback saved',book['date'])
