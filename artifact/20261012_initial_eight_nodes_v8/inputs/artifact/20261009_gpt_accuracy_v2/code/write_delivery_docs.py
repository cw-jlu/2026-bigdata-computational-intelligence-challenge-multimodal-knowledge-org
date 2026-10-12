"""Produce an auditable engineering account; never alters event decisions."""
import hashlib
import importlib.metadata
import json
import platform
import re
import sys
from pathlib import Path
from datetime import datetime,timezone
from openpyxl import load_workbook

V=Path(__file__).resolve().parents[1]; ROOT=V.parents[1]; A=V/'audit'; O=V/'output'
def read(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def save(p,x): p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
stats=read(O/'中间证据/reconstruction_summary.json'); comparison=read(A/'coverage_comparison.json')
cases=read(A/'independent_cases.json'); hold=read(A/'holdout_cases.json')
raw={x['date']:x for x in (json.loads(s) for s in (A/'raw_snapshot.jsonl').open(encoding='utf-8'))}
records=[]
for c in cases['cases']+hold['cases']:
    rows=[r for r in raw[c['date']]['radar'] if r['data']['批号']==c['batch']]
    if 'should_associate' in c:
        relation='应关联指定文本与该雷达候选' if c['should_associate'] else '不应关联指定文本与该雷达候选'
    elif c.get('expect_competing_tracks'):
        relation='不应把同批号两组异地同时轨迹合并；分支与文本的关系须另核验'
    elif c['id'] in ['C14','C15','C16']:
        relation='已有两源可作候选关联；不应确认为三源同一事件，也不纳入'
    elif c['id'] in ['H06','H14']:
        relation='三源应关联同一活动；非核心海域，不应纳入日报'
    elif c.get('expect_inclusion',c.get('included',False)):
        relation='三源应关联并纳入；只认正文明确或双源支持的字段'
    else:
        relation='不应直接确认三源同一事件；保留冲突和候选，排除日报'
    prefixes=[]
    if c.get('file_prefix'): prefixes.append((c['source'],c['file_prefix']))
    for source,pattern in [('开源军情',r'(?:开源|军情)(\d{3})'),('网上舆论',r'(?:舆论|舆情)(\d{3})')]:
        prefixes.extend((source,m+'_') for m in re.findall(pattern,c['basis']))
    text_refs=[{'file':t['file'],'text':t['text']} for t in raw[c['date']]['texts'] if any(t['source']==s and Path(t['file']).name.startswith(n) for s,n in prefixes)]
    records.append({**c,'expected_relation':relation,'raw_radar_file':rows[0]['file'],'raw_radar_rows':rows,'direct_text_locators_from_annotation':text_refs,
                    'locator_limit':'未明确指定TXT编号的案例，应按同日期原始快照及人工依据查阅；不由程序的匹配输出补写真值。'})
save(A/'cases_raw_records.json',{'method':'关系标签来自人工原始材料审阅；仅复制原始定位和行，不读取融合输出生成标签','cases':records})
lines=['# 原始材料人工案例审阅','',
       '下面42条是工程师根据原始文本与Excel行制定的预期。它们覆盖特意抽取的易错案例，不能用于估计总体准确率。指定文本的应关联与最终三源可纳入是不同判断；两源候选、非核心活动和无法消解的分支分别说明。完整原始行在 cases_raw_records.json，初始标签和首次失败保留。','',
       '| 案例 | 日期 | 类型 | 预期关系 | 原始雷达定位 | 关键字段依据 |','|---|---|---|---|---|---|']
for c in records:
    rr=c['raw_radar_rows']; rows=','.join(str(r['row']) for r in rr)
    lines.append(f"| {c['id']} | {c['date']} | {c.get('category','跨节点三源复核')} | {c['expected_relation']} | {rr[0]['sheet']} 行{rows}；批号{c['batch']} | {c['basis'].replace('|','／')} |")
lines+=['','标签修订：C03的原始TXT定位由034改为033；034实际上为中国歼-16。C19、C20最初只看局部行，误认为单行区域污染；逐条核对全部原始行发现两组持续区域记录，因此改为不可消解并排除。C26明确把电子侦察机名称与本次电子侦察任务区分。这些修订保留初始版本，没有把模型输出自动当作真值。']
(A/'manual_case_review.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')

portable=[]
for mode in ['portable_reproduction','clean_reproduction']:
    port=A/mode/'output'
    for n in stats['nodes']:
        day=n['date']; main=O/'日报'/day/(day+'机舰活动明细.xlsx'); alternate=port/'日报'/day/main.name
        a=load_workbook(main,data_only=True);b=load_workbook(alternate,data_only=True)
        equal=a.sheetnames==b.sheetnames and all(list(a[s].values)==list(b[s].values) for s in a.sheetnames)
        portable.append({'mode':mode,'date':day,'cell_values_equal':equal});a.close();b.close()
assert all(x['cell_values_equal'] for x in portable)
save(A/'portable_content_comparison.json',{'ok':True,'checks':portable,'note':'Compare all two-sheet cell values and dates, not byte identity or font metrics.'})

render=read(A/'docx_visual/render_log.json'); rasters=read(A/'docx_visual/raster_checks.json'); excel=read(A/'xlsx_native/native_excel_log.json')
assert all(not x['blank_pages'] and not x['text_outside_page'] for x in rasters)
assert all(x['docx_sha256']==sha(Path(x['docx'])) for x in render)
assert all(x['xlsx_sha256']==sha(O/'日报'/x['date']/(x['date']+'机舰活动明细.xlsx')) for x in excel)
assert len(excel)==16 and all(x['cell_A2_display'].startswith('2025.') for x in excel)
validation=read(O/'validation.json'); development=read(A/'cases_final.json'); cross=read(A/'holdout_results.json'); invariant=read(A/'invariance_results.json')
assert validation['ok'] and development['passed']==development['cases'] and cross['passed']==cross['cases'] and invariant['ok']
run=read(A/'full_run.json');prun=read(A/'portable_reproduction/full_run.json');clean=read(A/'clean_reproduction/full_run.json')
assert run['original_input_reread'] and prun['original_input_reread'] and prun['validation_ok'] and clean['original_input_reread'] and clean['validation_ok']
assert sha(A/'raw_snapshot.jsonl')==sha(A/'clean_reproduction/raw_snapshot.jsonl')
assert sha(O/'中间证据/events.jsonl')==sha(A/'clean_reproduction/output/中间证据/events.jsonl')
docs=[ROOT/'README.md']+sorted((ROOT/'docs').glob('0[1-5]-*.md'))
save(A/'reference_manifest.json',{'read_before_implementation':[{'file':p.relative_to(ROOT).as_posix(),'sha256_at_delivery':sha(p)} for p in docs],
    'old_gpt_only':['artifact/20261008_001028_gpt-5.6solhigh/solution.py','artifact/20261008_001028_gpt-5.6solhigh/README.md','artifact/20261008_001028_gpt-5.6solhigh/output/中间结果/evidence.jsonl'],
    'restriction':'No Claude implementation or artifacts read/copied. Shared docs05 comparison table is a required rules document, not a competitor artifact.'})
verification={'generated_at':datetime.now(timezone.utc).isoformat(),'official_re_evaluation':None,
              'structural_validation':validation['ok'],'raw_case_development':[development['passed'],development['cases']],
              'raw_case_cross_node_regression':[cross['passed'],cross['cases']],'invariance_checks':len(invariant['checks']),
              'visual_pages':sum(x['pages'] for x in rasters),'workbook_native_previews':len(excel),
              'visual_human_review':'All 8 reports reviewed via page contact images; all 16 native sheet previews reviewed. Focus: Chinese glyphs, headings, indexes, 4 chart references, plot labels, table breaks. Preview sheets show first 4 data rows; full workbook cell checks cover all rows.',
              'renderer_fallback':'Skill render_docx.py failed because LibreOffice executable unavailable; used installed Microsoft Word read-only hidden native PDF export. QA PDFs excluded from submission.',
              'artifact_renderer_limit':'Artifact-tool preview prints Excel date serials despite saved yyyy.mm.dd format. Native Excel display and openpyxl date type verified separately.',
              'portable_content_equal':True,'clean_venv_all_eight_from_original':True,'clean_raw_snapshot_and_events_sha_equal':True,
              'submission_zip_sha256':validation['zip']['sha256']}
save(A/'verification_record.json',verification)
rows=['| 节点 | 新版纳入 | 排除候选 | 飞机架次/舰艇艘次 | 旧版纳入 | 新增/移除批次 | 主线日期数/阶段数 |','|---|---:|---:|---|---:|---|---|']
for n,x in zip(stats['nodes'],comparison['comparison']):
    rows.append(f"| {n['date']} | {n['included']} | {n['excluded']} | {n['air_quantity']}/{n['ship_quantity']} | {x['old']} | {len(x['added_batches'])}/{len(x['removed_batches'])} | {n['trend']['dates']}/{n['trend']['stages']} |")
analysis='''# 重做结果与实际风险

八节点共433个雷达轨迹候选，纳入212起、排除221个；飞机118架次、舰艇93艘次之外，飞机数量须以表中逐节点汇总为准（本段实际总数由脚本下方计算）。轨迹点不作为架次/艘次。排除包括已核验非核心海域、缺第三类同事件证据、关键字段矛盾、异常轨迹及竞争文本不能唯一分配。排除是保守判定，不代表所有排除活动都是假事件。

TABLE

旧版纳入269起，新版212起；按同一原始雷达批号比较，共保留193个、新增19个、移除76个。新增补足明确别名、可拆分多平台正文、文本限定轨迹分支和实名舰艇；移除涉及主体、任务、数量、时段与区域冲突。完整名单在 audit/coverage_comparison.json。统计单位是候选活动，不是文本篇数。八节点三条活动的具体任务只得到雷达单源，输出“任务未核实”，不把平台用途或标题词汇升级为任务事实。未双源核实的起降、基地、呼号、母港留空。

## 旧风险核实与重构

旧 solution.py 的157行采用分组首行；176—183行整篇关键词及标题相似度选文本，没有平台子证据的独占约束；185—187行粗略坐标筛选；277行取首个历史同型事件；304行固定半段叙述。原始证据核查确认这些并非仅理论风险：

- 5月10日台湾P-3C 11:21—13:46（AIR-C498D3CE9B3184）旧版使用舆论029的14:46—16:16联合巡逻两架活动。本版用舆论037的同时间反潜活动；当日晚段另一P-3C（AIR-74BB12F6148F7E）数量1/2冲突，排除。
- 空警-600（AIR-6A5EB9A4E1BBA4）开源049空中警戒，与雷达/舆论预警指挥冲突。旧版纳入预警指挥；新版保留矛盾并排除。
- 无敌级（SHIP-C8D66FB15B8C0B）与056A（SHIP-F4B78A88D06154）旧版使用首行南海/东海，而文本分别指宁波/连云港。逐行审阅发现各有两组持续区域记录，不是单行偶然错字；新版不靠文本多数或雷达点数压过矛盾，排除。
- 903A（SHIP-1BEADE90A1114F）雷达东海南部，旧版关联军情011菲律宾海护航及舆论011黄海成山头补给。相同序号只是巧合。本版保存所有同型候选、区域和任务排除理由。

新版只从正文事实段提取平台子事件；先检查国家、型号、活动时段、具体区域、任务、数量及连续观测，再进行逐字段三源裁决。每个肯定关键字段有至少两种来源；任务无法核实可保留未知。明确第三源矛盾不以二比一投票消解。候选强度只在已通过事实检查的证据组合之间排序，不能代替真实性判定；多个组合同强度或同一子证据被不同雷达事件竞争时保留不确定并排除。单篇明确描述不同平台允许拆分，各自承担自己的数量证据。

## 原始案例发现与修复

先直接读取原始TXT和Excel标注28条易错案例，首次23条通过。五条失败分别包含：C02的非核心地理包络缺失导致孟加拉湾A400M可误连钓鱼岛轨迹，已补独立区域排除；C03人工原始定位034应为033，已修正定位；C19/C20人工误认为孤立一行，读全批次行后改为两组持续冲突；C26把“电子侦察机”用途当任务，已区分正文任务，保留同事件但任务未知。初始标签和首次失败全部保留，这些不是五个算法准确率改进。

另从其余七节点原始三源提取14条案例，首次13条符合。H12山东号是单一实名舰艇，正文不一定写“一艘”；增加明确“号”且具舰艇实体词的数量1推断，限定具名舰体，不能推广到舰级或机型。该轮还发现“19时”缺分钟的结束时间未解析、075同日多活动需要更具体证据、单文本细分区域不能由宽容坐标包络证实，均已修复。全部42条最终符合人工预期，结果分别在 cases_final.json、holdout_results.json；它们是发展与回归案例，有人工误标风险，不是官方真值，也不是未触碰测试集准确率。

独立检查8节点原始雷达顺序/TXT顺序/文件重新编号，以及8节点未来材料加入后历史主线不变，共16项通过。另对已提取文本做同型同完整时段同区域的额外数量/主体/任务矛盾扫描，标记0条；该扫描仍会漏掉未成功提取或部分窗口矛盾，不能声称无内容错误。结构验证重读所有单元格、原始引用和图表数据，与人工原始案例分开记录。

## 分析和日期边界

正文、当日D证据、XLSX和4图只用节点当日。历史主线从当天已核验外（台）平台回溯至本节点，按主体/完整型号聚合，至少3日期；各节点实际有5—30个日期，均形成3阶段。阶段按实际任务、方向改变和相邻证据间隔设边界，列出具体海域、任务次数、时间窗口及原始TXT/Excel行。它们表示同主体同型平台活动线索，不证明同一机体或舰体迁移，不因资料空缺推断离开。

舰船原始记录缺少真实经纬度，活动图采用时间—具体海域示意，飞机/舰艇分别标记，图中不补造舰艇坐标。数量曲线仅统计已核实侦巡任务，按活动开始小时记架次/艘次，不解释为同时在场数量。四图均有正文引用和绘图数据。

## 未解决风险与评价边界

规则提取受现有正文格式、词典和任务表达限制；未提取文本保留待查。坐标包络是宽松排除规则，不是精确海区边界，80km分支及2400km/h阈值是工程假设。持续冲突、观测子时段与完整活动时段关系不明时保守排除，可能降低真实覆盖；未知任务和留空基地等也可能损失完整度。具有相同型号、时段、海域的不同机体/舰体仍可能难以辨别；时间图也不能替代真实航迹地图。跨日长跨度相同型号只表示线索，资料不足以给出离开或未来活动预测。

没有本版官方重新评测结果。旧版用户提供的完整性83.2458、准确性17.8993、格式73.3715、分析25.3750、总分44.0519仅作背景。本版不声称分数或总体准确率提高。
'''
air=sum(n['air_quantity'] for n in stats['nodes']); ship=sum(n['ship_quantity'] for n in stats['nodes'])
analysis=analysis.replace('飞机118架次、舰艇93艘次之外，飞机数量须以表中逐节点汇总为准（本段实际总数由脚本下方计算）',f'飞机{air}架次、舰艇{ship}艘次')
analysis=analysis.replace('TABLE','\n'.join(rows))
(V/'ANALYSIS.md').write_text(analysis,encoding='utf-8')
audit=f'''# 最终验证记录

主流程从原始输入重读并重建：{run['seconds']}秒；便携Python流程同样重读：{prun['seconds']}秒；全新隔离虚拟环境仅按requirements安装后的原始数据重建：{clean['seconds']}秒。日志与命令在 audit/full_reproduction.log、full_run.json、portable_reproduction/ 和 clean_reproduction/。后者记录安装日志和完整requirements.lock.txt。这是当前实测，不保证其他机器耗时一致。

- 8节点文件结构、DOCX标题和五章、正文四图引用、32 PNG尺寸及绘图数据、16工作表固定列序/日期/数量类型/全部单元格、当日与历史截止边界：通过。
- 每个纳入事件三类来源定位、肯定关键字段至少两类支持、未知任务例外注明、子证据唯一分配、原始正文片段与Excel行：通过证据链一致性检查；该检查不提供官方内容准确率。
- 原始人工开发案例：28/28；其他七节点回归案例：14/14。首次失败和标签更正保留。原始案例详细依据在 audit/manual_case_review.md、cases_raw_records.json。
- 乱序/文件重新编号不影响当日事实，未来记录不影响截至节点主线：16/16。
- Word原生渲染{verification['visual_pages']}页，逐页检查，无空白页、无页外文字。所有节点页面拼图已查看；修复默认标题蓝色边框残留、证据索引孤行续页和强制分页造成的空白页。保留PDF/PNG、render_log.json和raster_checks.json，渲染日志记录最终DOCX SHA256。
- XLSX由artifact-tool生成，并进行内部检查和预览。其日期预览器显示序列值的局限已记录；openpyxl确认原生日期类型，Excel实际显示yyyy.mm.dd。16个原生工作表前4条数据预览已查看，列宽、换行、表头及全部单元格另作程序验证。原生日志含最终XLSX SHA256。
- 便携流程及全新虚拟环境各8个工作簿的所有单元格与主交付一致：通过，见 portable_content_comparison.json。全新环境的原始快照与全部事件JSONL SHA256也与主流程一致。
- 提交ZIP48文件，根目录日报/，无PDF、无辅助证据，CRC及48项SHA256逐个一致：通过。

提交ZIP SHA256：`{validation['zip']['sha256']}`。

LibreOffice不可用，文档技能脚本渲染尝试失败后使用本机Microsoft Word隐藏只读导出，没有修改日报。QA PDF只用于视觉验证，不进入可提交ZIP。完整复制的原始数据快照属于复现证据，也不进入日报ZIP。

以上结构、证据链及42个特选人工案例分别记录；没有官方重新评分，不把任一校验通过解释为整体事实正确。
'''
(V/'AUDIT.md').write_text(audit,encoding='utf-8')
save(A/'environment.json',{'python':sys.version,'os':platform.platform(),'dependencies':{n:importlib.metadata.version(n) for n in ['openpyxl','python-docx','matplotlib','Pillow','lxml','PyMuPDF']},
                         'docx_runtime_python':'3.12.14','artifact_tool':'2.8.89','random_seeds':[4601,4602],'external_models':[],'external_data':[],'gpu_required':False})
print('Analysis, raw case records, portable comparison and final audit written')
