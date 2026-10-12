"""Submission format/date/evidence consistency checks, NOT an accuracy metric.

Schema constants are deliberately independent of the report builder.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path
from datetime import datetime
from collections import Counter
from zipfile import ZipFile,ZIP_DEFLATED,ZipInfo
from openpyxl import load_workbook
from PIL import Image
from lxml import etree

DATES=['2025.5.10','2025.5.18','2025.5.30','2025.6.9','2025.6.17','2025.6.29','2025.7.9','2025.7.17']
AIR=['日期','星期','国家地区','机型','批号','批次','架次','起飞基地','呼号','起飞时间','降落时间','活动方向','活动区域','任务描述','任务类型','备注']
SHIP=['日期','国家地区','舰船','批号','艘次','活动时间','活动方向','活动区域','任务','母港','备注']
PNG=['01_舰机侦巡数量曲线图.png','02_中国主要机舰活动示意图.png','03_美国主要机舰活动示意图.png','04_其他国家主要机舰活动示意图.png']
HEAD=['一、重要军事动向','二、重点态势分析','三、外（台）机舰活动情况','四、美重要武器平台部署情况','五、次日关注重点']
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def require(condition,description,errors):
    if not condition: errors.append(description)
def normalize_date(x):
    if isinstance(x,datetime): return f'{x.year}.{x.month}.{x.day}'
    return str(x)

def check(out,raw_snapshot):
    daily=out/'日报'; intermediate=out/'中间证据'; results=[]; manifests=[]
    raw={d['date']:d for d in (json.loads(x) for x in raw_snapshot.open(encoding='utf-8'))}
    charts=json.loads((intermediate/'chart_datasets.json').read_text(encoding='utf-8'))
    data={p['date']:p for p in json.loads((intermediate/'workbook_rows.json').read_text(encoding='utf-8'))}
    for day in DATES:
        errors=[]; folder=daily/day; pack=json.loads((intermediate/f'factpack_{day}.json').read_text(encoding='utf-8')); ee=pack['events']
        expected={day+'日报.docx',day+'机舰活动明细.xlsx',*['附件图表/'+x for x in PNG]}
        actual={p.relative_to(folder).as_posix() for p in folder.rglob('*') if p.is_file()}
        require(actual==expected,'节点文件集合不符:'+str(actual^expected),errors)
        source_files={t['file']:t for t in raw[day]['texts']}; source_rows={(r['sheet'],r['row']):r for r in raw[day]['radar']}
        owned=[]
        for e in ee:
            require(e['date']==day,'纳入事件非当日:'+e['event_id'],errors)
            cc=[pack['claims'][c] for c in e['claims']]
            require({c['source'] for c in cc}=={'开源军情','网上舆论'},'事件三源种类不足',errors)
            require(len(e['claims'])==2,'文本子证据数量错误',errors)
            for c in cc:
                require(c['date']==day,'文本事实日期非当日',errors)
                require(c['file'] in source_files,'原始文本不存在',errors)
                if c['file'] in source_files:
                    require(c['quote'].replace(' ','') in source_files[c['file']]['fields']['正文'].replace(' ',''),'原文摘录不在原始正文',errors)
                owned.append(c['claim_id'])
            for row in e['radar_rows']:
                r=source_rows.get((e['kind'],row)); require(r is not None,'雷达行不存在',errors)
                if r:
                    require(r['file']==e['radar_file'] and r['data']['批号']==e['batch'],'雷达原始定位/批号不符',errors)
                    require(r['data']['国家地区']==e['country'],'主体与原始雷达行不符',errors)
                    require(int(r['data'].get('目标数量',r['data'].get('数量')))==e['quantity'],'数量与原始雷达行不符',errors)
            for key,fact in e['facts'].items():
                if fact.get('status')=='unknown':
                    require(key=='task' and fact['value']=='任务未核实' and not fact['sources'],'未知字段伪装为核实事实',errors)
                else: require(len(set(fact['sources']))>=2,'关键字段不足两类:'+e['event_id']+'/'+key,errors)
            require(e['direction'] in ['东海','南海','黄海'],'纳入非核心海域',errors)
            require(all(datetime.fromisoformat(p['time']).date()==datetime.strptime(day,'%Y.%m.%d').date() for p in e['points']),'轨迹含非当日点',errors)
        require(len(owned)==len(set(owned)),'同一平台子证据重复分配',errors)
        tr=pack['trend']
        require(tr is not None,'缺少3日期主线',errors)
        if tr:
            require(len(set(tr['dates']))>=3 and len(tr['phases']) in [2,3],'主线日期或阶段数不符',errors)
            require(day in tr['dates'],'主线节点当日无活动',errors)
            for e in tr['evidence']:
                require(datetime.strptime(e['date'],'%Y.%m.%d')<=datetime.strptime(day,'%Y.%m.%d'),'主线使用未来信息',errors)
                require(e['country']!='中国' and e['decision']=='纳入','主线主体或核验状态错误',errors)
            ids={e['event_id'] for e in tr['evidence']}
            phase_ids=[x for phase in tr['phases'] for x in phase['event_ids']]
            require(set(phase_ids)==ids and len(phase_ids)==len(ids),'主线阶段遗漏/重复',errors)
        book=load_workbook(folder/(day+'机舰活动明细.xlsx'),data_only=False)
        require(book.sheetnames==['飞机','舰艇'],'工作表名称/数量错误',errors)
        for s,headers in [('飞机',AIR),('舰艇',SHIP)]:
            ws=book[s]; require(ws.max_column==len(headers),'列数错误:'+s,errors)
            require([c.value for c in ws[1]]==headers,'列顺序错误:'+s,errors)
            target=next(x for x in data[day]['sheets'] if x['name']==s)['rows']
            rows=list(ws.iter_rows(min_row=2,values_only=True)); require(len(rows)==len(target),'明细行数错误:'+s,errors)
            for rr,tt in zip(rows,target):
                require(isinstance(rr[0],datetime) and normalize_date(rr[0])==day,'日期不是原生日期/非当日:'+s,errors)
                expect=list(tt); expect[0]=day; actual=list(rr); actual[0]=normalize_date(rr[0])
                require([x if x is not None else '' for x in actual]==expect,'明细字段与核验事实包不一致:'+s,errors)
                q=rr[6] if s=='飞机' else rr[4]; require(type(q) is int and q>0,'数量类型/单位错误',errors)
                require(not any(isinstance(v,str) and v.startswith('=') for v in rr),'明细意外含公式',errors)
        book.close()
        doc=folder/(day+'日报.docx')
        with ZipFile(doc) as z:
            require(z.testzip() is None,'DOCX损坏',errors)
            xml=etree.fromstring(z.read('word/document.xml')); text=''.join(xml.itertext())
            body=xml.find('w:body',NS); paras=[''.join(x.itertext()) for x in body]
            dt=datetime.strptime(day,'%Y.%m.%d'); title=f'{dt.month}月{dt.day}日飞机舰船活动情况'
            require(title in text,'DOCX固定标题缺失',errors)
            require('当日活动总览' in text and all(h in text for h in HEAD),'DOCX章节缺失',errors)
            require(len(xml.findall('.//w:drawing',NS))==4,'DOCX附图嵌入数错误',errors)
            for i in range(1,5): require(text.count(f'附图{i}')>=2,'正文未引用附图'+str(i),errors)
            require(not etree.fromstring(z.read('word/styles.xml')).findall('.//w:pBdr',NS),'标题残留蓝色边框',errors)
            section='daily'
            for line in paras:
                if line==HEAD[1]: section='history'
                elif line in HEAD and line!=HEAD[1]: section='daily'
                for ds in re.findall(r'2025\.\d{1,2}\.\d{1,2}',line):
                    dd=datetime.strptime(ds,'%Y.%m.%d')
                    require(dd<=dt if section=='history' else dd==dt,'DOCX章节日期越界:'+ds,errors)
        for name in PNG:
            p=folder/'附件图表'/name
            with Image.open(p) as im:
                require(im.format=='PNG' and im.width>=1200 and im.height>=600,'PNG类型/尺寸错误:'+name,errors); im.verify()
            datasets=[x for x in charts if x['date']==day and x['filename']==name]
            require(len(datasets)==1,'PNG数据集缺失/重复',errors)
            if datasets:
                c=datasets[0]; ids={e['event_id'] for e in ee}; require(set(c['event_ids'])<=ids,'图表使用未核验/非当日事件',errors)
                if name==PNG[0]:
                    aa=Counter(); ss=Counter()
                    for e in ee:
                        if e['task'] in c['recon_tasks']: (aa if e['kind']=='飞机' else ss)[e['interval'][0]//60]+=e['quantity']
                    require(c['air']==[aa[h] for h in range(24)] and c['ships']==[ss[h] for h in range(24)],'侦巡曲线数量不符',errors)
                else:
                    group=c['group']; select={e['event_id'] for e in ee if (e['country']==group if group in ['中国','美国'] else e['country'] not in ['中国','美国'])}
                    require(set(c['event_ids'])==select,'活动示意图分组/覆盖不符',errors)
                    for record in c['records']: require(record['date']==day,'活动图含跨日数据',errors)
        for relative in sorted(expected):
            p=folder/relative; manifests.append({'path':p.relative_to(out).as_posix(),'bytes':p.stat().st_size,'sha256':sha(p)})
        results.append({'date':day,'ok':not errors,'included':len(ee),'errors':errors,'format_checks':'文件、章节、表结构/原生日期/数值、PNG解析/数据/引用、DOCX章节日期、跨日日期/阶段','evidence_checks':'原文片段/原始行定位、来源种类、字段两源、子证据唯一性；不等于官方内容准确率'})
    allok=all(x['ok'] for x in results)
    zip_path=out/'日报_gpt_accuracy_v2.zip'
    if allok:
        with ZipFile(zip_path,'w',ZIP_DEFLATED,compresslevel=9) as z:
            for entry in manifests:
                zi=ZipInfo(entry['path'],date_time=(2026,10,9,0,0,0)); zi.compress_type=ZIP_DEFLATED; zi.external_attr=0o644<<16
                z.writestr(zi,(out/entry['path']).read_bytes())
        with ZipFile(zip_path) as z:
            names=z.namelist(); expected_names=[x['path'] for x in manifests]
            zip_errors=[]
            require(sorted(names)==sorted(expected_names) and len(names)==48,'ZIP结构/文件数量错误',zip_errors)
            require(z.testzip() is None,'ZIP CRC失败',zip_errors)
            require(not any(x.endswith('.pdf') for x in names),'ZIP含PDF',zip_errors)
            for entry in manifests: require(hashlib.sha256(z.read(entry['path'])).hexdigest()==entry['sha256'],'ZIP内容哈希不符',zip_errors)
        allok=allok and not zip_errors
    else: zip_errors=['节点校验失败，未创建提交ZIP']
    result={'ok':allok,'nodes':results,'zip':{'path':str(zip_path),'files':48 if allok else None,'errors':zip_errors,'sha256':sha(zip_path) if zip_path.exists() else None},'note':'仅结构、时间边界和证据链一致性；独立原始案例结果另存，二者均无官方重新评分'}
    (out/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    (out/'deliverable_manifest.json').write_text(json.dumps(manifests,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'ok':allok,'errors':[{k:r[k] for k in ['date','errors']} for r in results if not r['ok']],'zip_errors':zip_errors},ensure_ascii=False))
    return result

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--output',type=Path,default=Path(__file__).resolve().parents[1]/'output'); p.add_argument('--snapshot',type=Path,default=Path(__file__).resolve().parents[1]/'audit/raw_snapshot.jsonl'); a=p.parse_args()
    result=check(a.output,a.snapshot)
    raise SystemExit(0 if result['ok'] else 1)
