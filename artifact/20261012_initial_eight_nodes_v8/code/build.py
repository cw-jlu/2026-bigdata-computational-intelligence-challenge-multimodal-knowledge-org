from common import *
from copy import deepcopy
from collections import Counter
import argparse,re,subprocess,sys
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
import charts,source_config
POINTS={
'2025.5.10':('观测区域和核实任务组合变化，比活动强度上升更有依据。','前段东海、黄海两地训练和巡逻，随后南海联合巡逻，再在节点日东海空中警戒。','第一段是训练与巡逻组合；第二段只有5月6日一个日期，第三段只有节点日。两次空间切换都缺少中间转场资料。'),
'2025.5.18':('东海编队警戒的多日记载之后，出现南海防空反导训练；可见任务与区域变化，连续驻留尚无依据。','前段三个日期都为东海编队警戒，末段在南海出现防空反导训练。','5月7日至18日相隔11日，应区分前后阶段；不能把这11日补为驻留或用它构造直接航线。'),
'2025.5.30':('记录从南海侦察、黄海联合巡逻转为黄海反潜巡逻的多日重复，最清楚的变化是任务集中度。','5月5日南海海上侦察巡逻，5月12日黄海联合巡逻，5月21日起多日黄海反潜巡逻。','后段六个证据日期反复出现相同核实任务，支持同主体同型活动线索的重复性；不能证明同一机体持续值班。'),
'2025.6.9':('三次核实记录的地点和任务不同，支持阶段性变化，尚不足以形成持续行动判断。','东海联合搜救训练、东海自由航行，随后节点日南海防空反导训练。','前两次虽然同属东海，具体区域和任务不同；5月30日至6月9日相隔10日，末段单日记录不能填补间隔。'),
'2025.6.17':('南海反潜巡逻反复出现，其间和节点日有其他方向记录；这是一条区域分布变化的同型平台线索。','南海反潜为主，5月底黄海与东海出现记录，6月初后再次以南海反潜记录为主。','节点日同时有南海和黄海核实活动。将主线概括为全部回到南海会漏掉当日黄海记录；同型平台的不同批次不能视为同一机体。'),
'2025.6.29':('东海情报监视侦察的重复记录后出现空中警戒，再在黄海出现情报监视侦察，变化涉及任务与区域两个维度。','前段三个日期东海情报监视侦察，中段单日东海空中警戒，末段两个日期黄海情报监视侦察。','6月6日至16日相隔10日；任务变化有核实依据，但没有足够身份资料证明由同一架飞机切换任务。'),
'2025.7.9':('黄海编队警戒多日出现后，节点前一日出现南海战备巡逻，节点日再次记录黄海编队警戒。','前段五个日期黄海编队警戒，7月8日南海战备巡逻，7月9日黄海编队警戒。','主线对象是加拿大同级舰艇活动，不是具名舰体。两日不同方向可能来自不同舰艇，不能据此推断一艘舰艇快速往返。'),
'2025.7.17':('由单日黄海战备巡逻记载转为南海编队警戒的三个日期记载，支持任务和区域变化，未证持续增强。','7月11日黄海战备巡逻，7月14、16、17日南海编队警戒。','本节分析7月11日至17日最近一段核实线索；较早同名活动不能直接用来证明本段连续驻留。')}
EXTERNAL={
'2025.5.10':'德国联邦政府2024年9月25日固定版报告记载，2024年空军印太活动包含在阿拉斯加、日本、澳大利亚和印度的系列演习。[B04，PDF第8页3.1] 这使“系列训练与转场”成为有历史依据的替代解释。报告没有核验本次德国台风活动，也没有证明这些2025记录属于上述2024部署；本节只借此约束把跨区域出现直接解释为任务升级的判断。',
'2025.5.18':'德国联邦政府2024年9月25日报告在3.1节明确记载，巴登-符腾堡号与法兰克福号的2024部署包含RIMPAC演习及制裁监测，并与多国合作。[B04，PDF第8页3.1] 同一具名平台在一个公开部署框架内可承担不同任务，这为“计划内科目或合作任务转换”提供背景。该2024报告不证明2025节点活动的原因、身份连续性或具体任务。',
'2025.6.9':'日本海上幕僚监部2025年4月16日固定新闻稿把与马来西亚的MALPAN活动分为港内潜水训练、海上机动与通信训练。[B09，PDF第1页2、4、5项] 这是任务变化也可能来自训练科目切换的近期公开案例。参加的马方平台是MAHAMIRU、KINABALU等，区域是马六甲海峡及巴生港，均不同于本主线；它只提供替代解释，不能验证莱库号的比赛事件或把不同区域直接同比。',
'2025.6.17':'2024年4月6日四国国防部长固定版联合声明，将包含澳大利亚的海空联合活动安排与互操作性目标相联系。[B01，PDF第1页第2—3段] 合作训练因此是解释反复任务记载的一种历史背景，而不是由P-8A的一般能力推导其本次任务。声明是四国共同表述，计为一个外部来源组；它不能证明本轮任何反潜活动属于该联合活动，也不与本节点架次作规模同比。',
'2025.7.9':'加拿大2022年固定版《印太战略》把单国与多国行动、前沿海军存在列为地区参与方式。[B07，PDF第17页／印刷页14第2段] 这一政策背景使“既定地区参与下的多任务活动”成为合理解释之一。政策公开立场不能证明本轮同级舰艇的具体意图，更不能由一般部署政策识别某艘舰艇或推断7月8日至9日的航线。',
'2025.7.17':'日本海上幕僚监部2025年2月6日固定公告，将戴高乐号列为Pacific Steller演习拟参加平台，科目包括防空、反潜、数据链和跨甲板活动。[B10，PDF第1页2—5项] 这份节点前公告表明航母出现可置于多国训练安排的背景下，为“合作演训而非单一意图”提供替代解释。公告的区域为菲律宾以东海空域、日期为2月10日至18日，不能验证7月本比赛三海活动；拟安排也不等于各科目已全部实施。'}
def item(e,ref=None):
    f=e['facts'];unit='架次' if e['kind']=='飞机' else '艘次';tm=f"，核实活动窗口{f['start']}—{f['end']}" if f['start'] and f['end'] else '，活动时间未充分核实'
    return f"{f['platform']}{f['quantity']}{unit}，在{f['region']}"+(f"执行{f['task']}" if f['task'] else '活动，任务未充分核实')+tm+f"[{ref or e.get('ref','')}]"
def source_locator(e):return [{'file':a['path'],'locator':a['locator'],'sha256':a['sha256'],'snippet':a['snippet']} for a in e['claims'].values()]+[{'file':e['radar_file'],'sheet':e['sheet'],'rows':[r['row'] for r in e['rows']]}]
def analysis(f,n,used):
    m=f['mainline'];ids=set(m['events']);hh=sorted([deepcopy(e) for e in f['verified_events'] if e['id'] in ids],key=lambda e:(dt(e['date']),e['start'] or 0,e['id']));refs={e['id']:f'H{i:02d}' for i,e in enumerate(hh,1)};allrefs=list(refs.values());claims=[]
    for e in hh:claims.append({'id':refs[e['id']],'type':'provided_fact','text':e['date']+' '+item(e,refs[e['id']]),'source_ids':[e['id']],'locator':source_locator(e),'premise_claim_ids':[],'limits':['仅这一证据日期的核实活动，不代表间隔日期在场'],'event_id':e['id']})
    q='分析问题：'+m['country']+m['platform']+'在节点前的变化，主要发生在区域、任务还是参与数量？多日记录能否支持同一实体连续活动？'
    p=POINTS[n];blocks=[('A01','总体判断：'+p[0]+f" 截至{dt(n).month}月{dt(n).day}日，本条主线有{len(m['dates'])}个不同证据日期、{len(hh)}条独立活动。"+p[1]+f"[{allrefs[0]}—{allrefs[-1]}]")]
    stages=[]
    for i,s in enumerate(m['stages'],1):
        es=[e for e in hh if e['id'] in s['events']];ds=s['dates'];region=sorted({e['facts']['region'] for e in es});tasks=Counter(e['facts']['task'] or '任务未知' for e in es);qty=sum(e['facts']['quantity'] for e in es);unit='架次' if es[0]['kind']=='飞机' else '艘次';rr=[refs[e['id']] for e in es]
        obs='；'.join(f"{e['date'].replace('2025.','').replace('.', '月')}日，{e['facts']['region']}，{e['facts']['task'] or '任务未知'}，{e['facts']['quantity']}{unit}[{refs[e['id']]}]" for e in es)
        if len(es)>5:obs='证据日期为'+'、'.join(d.replace('2025.','').replace('.','月')+'日' for d in ds)+'；区域包括'+'、'.join(region)+'；按独立活动数计，'+ '、'.join(f'{k}{v}条' for k,v in sorted(tasks.items()))+f"。[{rr[0]}—{rr[-1]}]"
        explanation=(s.get('boundary_basis') or '以核实的区域、任务与证据间隔为边界')+'。'
        if len(ds)==1:explanation+='本段只有一个证据日期，只描述当日变化，不写成持续状态。'
        else:explanation+='多个日期的相同或相近任务支持重复出现，但未观测日不补为在场。'
        if i>1:
            prev=[e for e in hh if e['id'] in m['stages'][i-2]['events']];pr=set(e['facts']['direction'] for e in prev);cur=set(e['facts']['direction'] for e in es);pt=set(e['facts']['task'] for e in prev);ct=set(e['facts']['task'] for e in es)
            explanation+=('与前段相比，核实海域方向集合发生变化。' if cur!=pr else '与前段相比，海域方向集合一致。')+('任务集合也有变化；这支持任务组合改变，但未证明行动意图改变。' if pt!=ct else '核实任务集合相同，不能称新增任务。')
        blocks.append((f'A{i+1:02d}',f'第{i}阶段（'+('仅'+ds[0] if len(ds)==1 else ds[0]+'至'+ds[-1]+'，仅下列证据日期')+f'）：{obs}。本段{len(es)}条独立活动、参与数量合计{qty}{unit}。'+explanation));stages.append({'id':f'S{i}','evidence_dates':ds,'event_ids':s['events'],'boundary_basis':s.get('boundary_basis','区域、任务签名或8—14日证据间隔'),'premise_claim_ids':rr})
    blocks.append(('A05','阶段含义与数量边界：'+p[2]+' 各证据日期的参与数量为'+ '、'.join(f"{d.replace('2025.','').replace('.', '月')}日{sum(e['facts']['quantity'] for e in hh if e['date']==d)}{'架次' if hh[0]['kind']=='飞机' else '艘次'}" for d in m['dates'])+'。各阶段长度与观测覆盖不同，不能直接比较阶段合计并称活动增强。记录数、独立事件数、参与数量和去重实体数分别处理。'))
    if n in EXTERNAL and used:blocks.append(('E01',EXTERNAL[n]))
    else:blocks.append(('E01','外部证据边界：本轮检索未取得与本主线直接相关且能确认节点前具体版本的新增外部实证材料。对相关当前网页的旧署名日期和装备一般用途均不作替代验证；分析限于比赛内核实的历史活动，不填入装饰性宏观统计。'))
    blocks.append(('A06','替代解释与综合判断：'+p[0]+' 既定科目转换、合作训练或不同同型平台轮换均可能解释部分记录。当前三源能够核实单次活动，但对这些原因的区分能力有限；需要节点前公开的训练日程、具名身份或机体标识与进出区域资料才能作更强判断。'+('上述外部材料补充了训练或地区参与的可行解释，并未证明本次原因。' if used else '相关外部版本证据仍是缺口，保留并列解释。')+' 报道选择和观测覆盖差异也可能影响可见频次；没有证据说明比赛雷达具有AIS相同限制，不直接移植这种假设。'))
    for cid,text in blocks:
        external=cid=='E01' and bool(used);claims.append({'id':cid,'type':'attributed_view' if external else 'inference' if cid!='E01' else 'method','text':text,'source_ids':used if external else [],'locator':[source_config.SOURCES[s]['locator'] for s in used] if external else [],'premise_claim_ids':[] if external else allrefs,'limits':['无记录日期不填补；未知身份不视为同一实体；不推断确定意图或后续部署']})
    return q,blocks,claims,stages,hh,refs
def add_para(doc,text,role='body',model=None):
    p=doc.add_paragraph()
    if model is not None and model._p.pPr is not None:p._p.insert(0,deepcopy(model._p.pPr))
    r=p.add_run(text)
    if model is not None and model.runs and model.runs[0]._r.rPr is not None:r._r.insert(0,deepcopy(model.runs[0]._r.rPr))
    else:
        r.font.name='黑体' if role in ['heading','title'] else '仿宋_GB2312';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),r.font.name)
    r.font.size=Pt(22 if role=='title' else 16);r.bold=role in ['heading','title'];r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    if role=='title':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    if role=='heading':p.paragraph_format.keep_with_next=True
    return p
def build(out=BASE):
    source_config.catalog(out);structure=load(out/'evidence/sample_structure.json');sample=ROOT/'提交示例文件/日报/2025.10.20/2025.10.20日报.docx';refdoc=Document(sample);specs=[]
    for n in NODES:
        ad=out/f'audits/{n}';fp=ad/'frozen_ledger.json';before=sha(fp);f=load(fp);ee=f['today_events'];used=source_config.USED[n];dest=out/f'official/{n}';dest.mkdir(parents=True,exist_ok=True)
        q,blocks,claims,stages,hh,hrefs=analysis(f,n,used);(ad/'analysis.md').write_text('## 二、重点态势分析\n\n'+q+'\n\n'+'\n\n'.join(t for _,t in blocks),encoding='utf8')
        chartsdata=charts.draw(f,dest/'附件图表');save(ad/'chart_data.json',chartsdata)
        doc=Document(sample)
        for child in list(doc._element.body):
            if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
        date=dt(n);title=f'{date.month}月{date.day}日飞机舰船活动情况';add_para(doc,title,'title',refdoc.paragraphs[0]);foreign=[e for e in ee if e['facts']['country']!='中国'];air=[e for e in foreign if e['kind']=='飞机'];ship=[e for e in foreign if e['kind']=='舰艇']
        add_para(doc,f"{date.month}月{date.day}日，共核实外（台）飞机{len(air)}批{sum(e['facts']['quantity'] for e in air)}架次、舰艇{sum(e['facts']['quantity'] for e in ship)}艘次；含中国的全部三海活动共{len(ee)}条独立记录。统计按活动参与数量计，不等同于去重机体或舰体数。附图1展示当日核实窗口的整点参与量，时间未知记录不补造窗口。",model=refdoc.paragraphs[1])
        add_para(doc,'一、重要军事动向','heading',refdoc.paragraphs[2]);important=sorted(ee,key=lambda e:(e['id'] not in set(f['mainline']['events']),not('号' in e['facts']['platform']),e['facts']['country']!='中国',e['ref']))[:4]
        for i,e in enumerate(important,1):add_para(doc,f'（{i}）'+e['facts']['country']+item(e)+'。',model=refdoc.paragraphs[1])
        china=[e for e in ee if e['facts']['country']=='中国'];add_para(doc,'中国当日核实活动：'+('；'.join(item(e) for e in china) if china else '无满足三源核验条件的三海活动记录')+'。中国主要机舰活动见附图2。',model=refdoc.paragraphs[1])
        add_para(doc,'二、重点态势分析','heading',refdoc.paragraphs[7]);add_para(doc,q,model=refdoc.paragraphs[1])
        for _,text in blocks:add_para(doc,text,model=refdoc.paragraphs[1])
        add_para(doc,'三、外（台）机舰活动情况','heading',refdoc.paragraphs[12])
        for i,c in enumerate(sorted({e['facts']['country'] for e in foreign}),1):
            es=[e for e in foreign if e['facts']['country']==c];aa=[e for e in es if e['kind']=='飞机'];ss=[e for e in es if e['kind']=='舰艇'];add_para(doc,f'（{i}）{c}：飞机{len(aa)}批{sum(e["facts"]["quantity"] for e in aa)}架次，舰艇{sum(e["facts"]["quantity"] for e in ss)}艘次。',model=refdoc.paragraphs[1])
            for e in es:add_para(doc,item(e)+'。',model=refdoc.paragraphs[1])
        add_para(doc,'其他国家及台湾活动见附图4；美国活动另见附图3。所有活动均按节点当日核实集合统计。',model=refdoc.paragraphs[1])
        add_para(doc,'四、美重要武器平台部署情况','heading',refdoc.paragraphs[29]);us=[e for e in ee if e['facts']['country']=='美国']
        add_para(doc,('当日核实的美国平台活动如下：'+'；'.join(item(e) for e in us) if us else '当日没有满足三源核验条件的美国三海平台活动')+'。这里只报告活动记载；同级舰艇名称不能识别单一舰体，现有记录不足以推断完整部署范围或确定意图。美国主要机舰活动见附图3。',model=refdoc.paragraphs[1])
        add_para(doc,'五、次日关注重点','heading',refdoc.paragraphs[31]);seeds=[e for e in ee if e['id'] in set(f['mainline']['events'])];seed=seeds[0]
        questions=f"以本节点{seed['facts']['country']}{seed['facts']['platform']}在{seed['facts']['region']}的记录[{seed['ref']}]为起点，后续公开资料核查重点为：一是能否补充具名身份或机体标识，区分延续活动与不同同型平台；二是不同来源的区域、任务和时间表述能否继续相互印证；三是能否取得同口径参与数量，避免把重复报道当作活动增长。这些是核查问题，未使用次日材料，不写成次日事件或部署预测。"
        add_para(doc,questions,model=refdoc.paragraphs[1]);claims.append({'id':'V01','type':'inference','text':questions,'source_ids':[],'locator':[],'premise_claim_ids':[seed['ref']],'limits':['核查问题，不是已知未来事实']})
        add_para(doc,'引用与跨日证据索引','heading',refdoc.paragraphs[2])
        for sid in used:
            s=source_config.SOURCES[sid];add_para(doc,f"[{sid}] {s['publisher']}：《{s['title']}》，公开上界{s['published_by']}，{s['locator']}。{s['url']}",model=refdoc.paragraphs[1])
        for e in hh:add_para(doc,e['date']+' '+item(e,hrefs[e['id']])+'。',model=refdoc.paragraphs[1])
        add_para(doc,'当日F编号对应明细备注及独立事实审计；跨日H编号对应上述证据索引。正文与附图1—4使用同一节点冻结台账。基地、呼号、起降等未充分核实的辅助字段留空。',model=refdoc.paragraphs[1])
        for i,name in enumerate(charts.NAMES,1):
            doc.add_page_break();add_para(doc,f'附图{i}','heading',refdoc.paragraphs[40]);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            # The external PNG is authoritative; fit the complete graphic inside
            # the sample page without cropping or stretching.
            from PIL import Image
            w,h=Image.open(dest/'附件图表'/name).size;maxw=15.6;maxh=21.1;scale=min(maxw/w,maxh/h);p.add_run().add_picture(str(dest/'附件图表'/name),width=Cm(w*scale),height=Cm(h*scale))
        doc.save(dest/f'{n}日报.docx')
        workbook=[]
        for spec in structure['sheets']:
            rows=[]
            for e in ee:
                if e['kind']!=spec['name']:continue
                x=e['facts'];note=f"{e['ref']}；三源核验；辅助信息不足两类支持留空";desc=f"{n}，{x['country']}{x['platform']}在{x['region']}"+(f"执行{x['task']}" if x['task'] else '活动，任务未充分核实')
                if e['kind']=='飞机':
                    if x['start'] and x['end']:note+=f"；观测窗口{x['start']}-{x['end']}，不作为起降时间"
                    rows.append([n,'星期'+'一二三四五六日'[date.weekday()],x['country'],x['platform'],e['batch'],1,x['quantity'],None,None,None,None,x['direction'],x['region'],desc,x['task'],note])
                else:rows.append([n,x['country'],x['platform'],e['batch'],x['quantity'],x['start'].replace(':','')+'-'+x['end'].replace(':','') if x['start'] and x['end'] else None,x['direction'],x['region'],x['task'],None,note])
            workbook.append({'name':spec['name'],'headers':spec['headers'],'rows':rows,'sample_widths':spec['widths']})
        save(ad/'workbook_input.json',workbook);specs.append({'node':n,'sheets':workbook})
        for e in ee:claims.append({'id':e['ref'],'type':'provided_fact','text':item(e),'source_ids':[e['id']],'locator':source_locator(e),'premise_claim_ids':[],'limits':['当日核实活动；辅助未知不补造'],'event_id':e['id']})
        audit={'as_of':f['as_of'],'mode':'historical_public_research','status':'generated_pending_independent_and_visual_review','read_files':load(ad/'input_manifest.json')+load(ad/'historical_read_files.json'),'frozen_fact_input':{'path':f'audits/{n}/frozen_ledger.json','sha256':before,'verification_scope':f['verification_scope']},'used_source_ids':used,'sources':[source_config.SOURCES[s] for s in used],'claims':claims,'stages':stages,'mainline_candidate_comparison':f['candidate_comparison'],'mainline_selection':f['mainline'],'missing_evidence':(['未取得直接相关且历史版本已确认的新增外部实证资料'] if not used else [])+['辅助基地、呼号、起降与母港不足两源支持；留空','缺少间隔日期完整观察及实体标识，不能证明持续驻留'],'excluded_sources': [{'id':s['id'],'reason':'本轮未核读支持正文；未采用'} for s in load(ad/'source_catalog.json')['sources'] if s['id'] not in used]+[{'id':'MINDEF_JMT2024','reason':'已读当前原发布页；明确迁移且网页2026更新，未确认节点前文字版本，不引用'}],'rule_conflicts':['本次示例ZIP根日期目录；当前平台页面示例含日报外层，缺少两者更新时间可判定先后'],'suspected_upstream_issues':['宽松坐标相容性窗口是工程假设；未对所有记录取得外部专家复核'],'checks':{'source_metadata':'not_run','semantic_support':'reviewed_external_passages_and_generated_claim_links; final_review_pending','unchanged_fact_layer':'PASS' if before==sha(fp) else 'FAIL','visual_review':'pending'},'model_parameters':{'implementation':'deterministic Python and artifact-tool; model API calls 0','authoring_model':'Codex session; exact backend version and sampling unavailable','generation_seed':'not_applicable'},'future_guard':{'max_evidence_date':max(dt(e['date']).isoformat() for e in f['verified_events']),'as_of':f['as_of']}}
        save(ad/'audit.json',audit);assert sha(fp)==before
        print('BUILT',n,'events',len(ee),'mainline_dates',len(f['mainline']['dates']),flush=True)
    save(out/'workbooks.json',specs)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=BASE);args=ap.parse_args();build(args.output.resolve())
