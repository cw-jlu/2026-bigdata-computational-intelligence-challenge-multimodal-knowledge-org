"""Record reviews actually performed in this Codex session, then seal delivery.

This is an audit recorder, not an automatic human-review oracle. It only accepts
the final output hashes saved before review. Freshly modified outputs need a new
visual/semantic review; Run-all.ps1 intentionally leaves their reviews pending.
"""
from common import *
from pathlib import PurePosixPath
import zipfile,sys,platform,importlib.metadata as md,shutil,collections
from docx import Document
from openpyxl import load_workbook
from PIL import Image
NOTES={
'2025.5.10': ['台风4证据日期5活动；第一段东黄海、第二段5月6日南海、第三段节点日东海，两个单日段没有写成持续状态。B04仅2024系列训练背景。',
 ['093A纳土纳1艘战备，原文及完整批次相容；辅助单源细节未补','英国P-8A联合巡逻1架，单源细地名降至南海','台湾经国平台孟加拉湾/安达曼海冲突，域外且不纳入','菲律宾舰艇搜救/自由航行明确任务冲突，不取多数']],
'2025.5.18': ['具名巴登舰4证据日期；5月7日至18日11日间隔单列后段。B04记载2024部署，不能证明2025持续驻留。',
 ['台湾磐石南海例行1艘三源相容','法国地平线钓鱼岛反潜1艘三源相容；单源离开措辞不入正文','中国福建舰缺第三类关联证据，不采用单文报道','法国絮弗伦缺第三类关联证据，不采用单文报道']],
'2025.5.30': ['韩国P-8A8日期9活动；5月21日起黄海反潜重复，任务集中度为有限判断；两条同日批次独立。外部历史版缺口保留。',
 ['空警500南海空中警戒1架三源相容','空警600南海侦察1架；北部湾仅单类支持，保留南海','中国歼20对海突击/情报监视侦察冲突，不采用','日本E-2D中太平洋原文与雷达坐标不相容，域外且不采用']],
'2025.6.9': ['莱库号3次单日记录；两次东海具体区域不同、10日间隔后南海。只支持已观测区域/任务差异，持续阶段证据薄弱。B09马来西亚不同舰艇的训练科目替代解释。',
 ['无侦7黄海电子侦察1架三源相容','意大利P72南海海上侦察1架，单源细地名没有补造','法国大西洋平台孟加拉湾/安达曼海冲突，域外且排除','澳大利亚供应舰补给/护航明确任务冲突，排除']],
'2025.6.17': ['澳大利亚P-8A12日期14活动；5月底方向转换及6月再次南海反潜为主，但节点日黄海另有独立活动。B01只给合作互操作目标背景，四国联合声明计一组。',
 ['加拿大哈里级连云港实弹1艘三源相容','俄罗斯图95东海对海训练1架三源相容','新加坡福克50正文国家新加坡/印度冲突，排除','美国阿利伯克缺第三类证据及数量两类支持，排除']],
'2025.6.29': ['新加坡G5506日期；东海侦察—单日警戒—黄海侦察，6月6日至16日10日间隔。无同机体证据。2026更新的新加坡网页历史版本不合格，未用于正文。',
 ['中国055黄海联合搜救1艘三源相容','中国039A南沙例行1艘三源相容','直20联合巡逻/情报监视侦察任务冲突，排除','马来西亚苏30正文出现俄罗斯国别，关键国别冲突排除']],
'2025.7.9': ['加拿大哈利法克斯级7日期；黄海重复、7月8日南海单日、节点黄海单日。级别记录不推出同舰快速往返。B07政策前沿海军存在不能识别舰体或确定意图。',
 ['意大利P72东海侦察1架三源相容','菲律宾AW159东海反潜1架三源相容','中国歼35数量1/2明确冲突，排除','新西兰P8联合/反潜任务冲突，排除']],
'2025.7.17': ['戴高乐4日期7月11日至17日两段；较早6月23日与7月11日18日间隔不反向填补。B10二月演习公告明确为拟安排，不能验证七月事件。',
 ['马来西亚苏30南海对海训练1架三源相容','中国歼16D中沙侦察1架三源相容','台湾F16V完整12雷达行同一时刻存在相距约700公里位置，未任意删点，保留排除','中国无侦7缺第三类关联及数量两类支持，排除']]
}
PURPOSE={
'B04':'德国2024系列部署/训练提供跨区域科目转换的历史解释，不能验证本轮2025事件',
'B09':'日马近期合作训练分港内/海上科目，提供任务轮换替代解释；平台/区域均不同于莱库号',
'B01':'澳大利亚参与海空合作的互操作目标作为明确归属的历史背景；不证明本次活动属于该安排',
'B07':'加拿大政策地区参与与前沿海军存在作为背景；不证明具体舰体/路线或意图',
'B10':'戴高乐拟参加多科目合作演习的公开历史安排；不写成已实施，也不验证七月比赛记录'}

def finalize():
    original=load(BASE/'submission_manifest.json');manifest=original['files']
    for row in manifest:assert sha(BASE/'official'/row['path'])==row['sha256'],('Unreviewed output modification',row['path'])
    reproduce=load(BASE/'reproduction_comparison.json');assert reproduce['status']=='PASS' and len(reproduce['nodes'])==8
    raster={x['node']:x for x in load(BASE/'raster_records.json')};matrix=load(BASE/'node_status_matrix.json')
    for n in NODES:
        ad=BASE/f'audits/{n}';a=load(ad/'audit.json');f=load(ad/'frozen_ledger.json');cards=load(ad/'independent_review_cards.json')
        a['sources']=[s for s in load(ad/'source_catalog.json')['sources'] if s.get('source_id',s.get('id')) in a['used_source_ids']]
        assert len(cards)==4 and len(NOTES[n][1])==4
        verdicts=[]
        for c,note in zip(cards,NOTES[n][1]):
            verdicts.append({'event_id':c['event_id'],'pipeline_decision':c['pipeline_decision'],'judgment':'agree_with_original_evidence','basis':note,'original_material':'independent_review_cards.json contains original full TXT and complete radar batch rows','limits':['每节点两纳入两排除，确定性哈希选样；不是盲测，不是独立外部专家审查，不估计整体正确率']})
        save(ad/'independent_review.json',{'as_of':a['as_of'],'reviewer':'Codex session direct re-reading of original TXT/Excel evidence, distinct from extraction code','sample_count':4,'status':'PASS_sample_only','cards_sha256':sha(ad/'independent_review_cards.json'),'cases':verdicts})
        pages=raster[n]['page_pngs'];assert len(pages)==len(load(ad/'rendered_text.json'))
        proof_paths=[BASE/p for p in pages]+sorted((BASE/f'qa/xlsx_native/{n}').glob('*.png'))+sorted((BASE/f'official/{n}/附件图表').glob('*.png'))
        review={'as_of':a['as_of'],'status':'PASS_full_node_visual_review','reviewer':'Codex session using view_image, all pages/full sheets and all four PNGs','docx_pages':len(pages),'reviewed_docx_page_numbers':list(range(1,len(pages)+1)),'xlsx_sheets':['飞机','舰艇'],'xlsx_ranges':'all populated rows and all 16/11 columns; native Excel four column bands per node','png_count':4,'method':'Microsoft Word/Excel read-only native PDF export, PDFium full raster; every page/contact viewed','formal_hashes':[r for r in manifest if r['node']==n],'visual_evidence':[{'path':p.relative_to(BASE).as_posix(),'sha256':sha(p)} for p in proof_paths],'findings':['正文无溢出/截断/乱码；五章及总览完整；图1—4引用和图件齐全','完整表格行列可见，长文字换行，数量是数值，日期/星期/时间是文本，辅助未知留空','全部PNG对象、标签、图例和数量可见；无核实美国活动的节点图3明确说明缺少核实记录','长URL在正文自动换行；不影响原文定位；单条图留白是未补造活动的结果'],'renderer_issue':{'artifact_preview':'飞机活动方向预览曾重复首个方向，导出的XLSX所有单元格比对及原生Excel渲染正常','impact':'仅工程预览不可靠；正式XLSX经逐值/类型校验和原生Excel完整检查','render_docx_py':'缺少LibreOffice，改用已安装Microsoft Word原生全页渲染'}}
        save(ad/'visual_review.json',review)
        semantic={'as_of':a['as_of'],'status':'PASS_with_documented_evidence_limits','mainline_review':NOTES[n][0],'stage_evidence':a['stages'],'checks':['逐条正文F事实回指冻结台账，所有H历史日期/事件ID回指独立节点快照','所有推断有明确前提，未由计数推算评分，未把同型/同级等同同一实体','逐阶段日期、任务/区域差异和间隔核对；单日/未观测日边界明确','外部支持段落实际核读；贡献为历史背景/归属观点/替代解释，未当比赛第三源','引用版本日期不晚于节点；次日部分仅列公开核查问题；未使用节点后内容'],'limits':a['missing_evidence']+a['suspected_upstream_issues']}
        save(ad/'semantic_review.json',semantic)
        for s in a['sources']:s['purpose']=PURPOSE[s['source_id']];s['reading_status_final']='actual_passage_review_completed_in_this_session';s['content_statistics_period']=s['period']
        a['status']='content_complete_with_documented_limits_platform_upload_pending'
        a['checks'].update({'visual_review':review['status'],'semantic_support':semantic['status'],'independent_review':'PASS_sample_only','raw_reproduction':'PASS_all_eight_from_raw_inputs','source_eligibility':'reviewed_fixed_versions_before_node' if a['used_source_ids'] else 'no_qualified_external_source_used_gap_documented','future_material_guard':'PASS_per_node_cutoff','statistics_and_unknown_fields':'PASS','unchanged_fact_layer':'PASS'})
        a['output_hashes']=[r for r in manifest if r['node']==n]
        a['read_files'].append({'path':'evidence/sample_structure.json','reading_status':'actual_full_DOCX_paragraph_and_XLSX_schema_read; design_only_not_facts','sha256':sha(BASE/'evidence/sample_structure.json')})
        for sid in a['used_source_ids']:a['read_files'].append({'path':next(s['local_path'] for s in a['sources'] if s['source_id']==sid),'reading_status':'actual_supporting_passage_read','source_id':sid})
        a['read_files']=list({(r.get('path'),r.get('source_id'),r.get('version')):r for r in a['read_files']}.values())
        save(ad/'audit.json',a)
        x=next(x for x in matrix if x['date']==n);x.update({'visual_review':review['status'],'independent_review':'PASS_sample_only','raw_reproduction':'PASS_all_eight_from_raw_inputs','docx_pages':len(pages),'source_eligibility':a['checks']['source_eligibility'],'content_status':'complete_with_documented_limits','uploaded':False})
    save(BASE/'node_status_matrix.json',matrix)
    # Independent actual extraction: strict paths, exact member set, full reopen.
    archive=BASE/'赛题2_初赛_八节点_完整候选提交包.zip';target=BASE/'qa/unpacked_submission';target.mkdir(parents=True,exist_ok=True)
    checked=[]
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None and len(z.namelist())==48 and set(z.namelist())=={r['path'] for r in manifest}
        assert {PurePosixPath(p).parts[0] for p in z.namelist()}==set(NODES)
        for r in manifest:
            pp=PurePosixPath(r['path']);assert not pp.is_absolute() and '..' not in pp.parts
            p=(target/Path(*pp.parts)).resolve();assert p.is_relative_to(target.resolve());p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(z.read(r['path']));assert sha(p)==r['sha256']
            if p.suffix=='.docx':assert len(Document(p).paragraphs)>0
            elif p.suffix=='.xlsx':w=load_workbook(p);assert w.sheetnames==['飞机','舰艇'];w.close()
            else:Image.open(p).verify()
            checked.append({'path':r['path'],'sha256':sha(p),'reopened':True})
    save(BASE/'independent_unpack_validation.json',{'status':'PASS','archive_sha256':sha(archive),'extraction_directory':target.relative_to(BASE).as_posix(),'nodes':NODES,'actual_extracted_files':48,'files':checked,'extra_engineering_files':[]})
    access={'accessed_at':'2026-10-12 Asia/Shanghai','url':'https://www.heywhale.com/home/competition/6aba1c4f6e5beff2fca33b42/content/3','official_api':'evidence/platform/competition.json','official_submission_text':'evidence/platform/tab_3.md','state':'IAB public page shows 登录/注册 and 报名; no authenticated upload session','available_browsers':['IAB','MCP Apps'],'upload_attempted':False,'receipt':None,'reason':'可用访问条件没有已登录参赛会话；目录层级冲突未解决','submission_method':'官方在线结果提交入口上传全赛段ZIP；每日最多3次，取最高分，初赛按8节点平均评分'}
    save(BASE/'evidence/platform/access_and_submission_status.json',access)
    env={'python':sys.version,'platform':platform.platform(),'python_executable':sys.executable,'node_runtime':'C:/Users/Andy1/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe','fonts':['C:/Windows/Fonts/msyh.ttc','Microsoft Word sample font styles','宋体 workbook'],'python_packages':{p:md.version(p) for p in ['python-docx','openpyxl','Pillow','pypdf','pypdfium2']},'model':'Codex agent session; backend/version/sampling not exposed; no model API calls inside generator','parameters':'deterministic parser and writer; no random sampling; mainline choice configuration in select_lines.py; per-node cutoff checks'}
    npm=Path('C:/Users/Andy1/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules/@oai/artifact-tool/package.json');env['artifact_tool_version']=load(npm)['version'];save(BASE/'dependencies.json',env)
    save(BASE/'code_manifest.json',[{'path':p.relative_to(BASE).as_posix(),'sha256':sha(p)} for p in sorted((BASE/'code').glob('*')) if p.is_file()])
    stage={'target_dates':[dt(n).isoformat() for n in NODES],'mode':'historical_public_research','status':'eight_node_content_and_candidate_package_complete','content_generated':True,'complete_packaging':True,'rules_confirmed':{'eight_nodes_together':True,'file_set':'本次提供实际示例：每节点DOCX、XLSX、4PNG','directory_layout':False,'conflict':'组织方当前公开页含日报外层，本次提供说明要求根日期目录；无修订日期可判定先后，依用户明确口径制作候选包','actual_upload_method':'在线全赛段ZIP上传，但本轮无已登录上传会话'},'actual_upload':False,'submission_receipt':None,'platform_evaluation':None,'rule_evidence':['evidence/platform/competition.json','evidence/platform/tab_3.md','evidence/platform/tab_4.md','inputs/提交示例文件/提交格式说明.md'],'nodes':matrix,'formal_members':manifest,'formal_archive':{'path':archive.name,'sha256':sha(archive),'files':48,'root_directories':NODES,'extra_members':[],'independent_actual_extraction':'PASS'},'raw_input_files':len(load(BASE/'evidence/input_manifest.json')),'frozen_ledger_hashes':{n:sha(BASE/f'audits/{n}/frozen_ledger.json') for n in NODES},'reproduction':reproduce,'checks':{'structure_and_cross_artifact_consistency':'PASS_all_eight','source_metadata':'PASS_metadata_only_all_eight','actual_passage_semantic_review':'completed_all_eight_with_limits','independent_fact_sample':'32 cases, 4 per node, original reread, limited deterministic sample','visual':'96 full Word pages, 16 entire worksheets, 32 PNGs viewed','raw_reproduction':'PASS_all_eight','archive':'PASS_exact48_actual_extracted'},'unresolved_gaps':['平台ZIP外层规则冲突未确认','可用浏览器未登录，不能取得上传回执','5月30日/6月29日直接相关外部历史版本资料缺口','辅助基地/呼号/起降/母港不足支持留空；时间未注明时区','未观测日期及实体身份缺口；宽松坐标相容性窗口为工程假设，未获全部记录外部专家复核'],'scores':'本轮没有平台评测，不推测或承诺任何得分提高','history_preserved':True,'storage':'D盘无剩余空间；实际新目录在C盘，项目artifact下junction指向该目录'}
    save(BASE/'stage_audit.json',stage)
    print('FINAL AUDITS SEALED: 8 nodes, 96 pages, 16 sheets, 32 PNG, 48 extracted files')

if __name__=='__main__':finalize()
