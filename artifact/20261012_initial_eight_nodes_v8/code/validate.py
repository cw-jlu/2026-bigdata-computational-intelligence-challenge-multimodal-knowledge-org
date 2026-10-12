from common import *
from openpyxl import load_workbook
from docx import Document
from PIL import Image
from collections import Counter
import zipfile,re,argparse,subprocess,sys
import charts
def validate(out=BASE,package=True):
    nodes=[];manifest=[];checks=[]
    for n in NODES:
        ad=out/f'audits/{n}';f=load(ad/'frozen_ledger.json');audit=load(ad/'audit.json');fr=load(ad/'freeze_record.json');ee=f['today_events'];ids={e['id'] for e in ee};errors=[]
        assert f['as_of']==dt(n).isoformat()==audit['as_of'];assert sha(ad/'frozen_ledger.json')==fr['sha256']==audit['frozen_fact_input']['sha256'];assert max(dt(e['date']) for e in f['verified_events'])<=dt(n)
        for e in f['verified_events']:
            assert e['decision']=='纳入' and len(e['claims'])==2 and not e['issues'];assert e['facts']['quantity'] is not None
            for field in ['country','key','quantity','region']:
                fa=e['field_audit'][field];assert len(fa['support_types'])>=2 and not fa['conflicts']
            for field in ['task','start','end']:
                if e['facts'][field]:assert len(e['field_audit'][field]['support_types'])>=2 and not e['field_audit'][field]['conflicts']
            for s,a in e['claims'].items():assert dt(a['date'])<=dt(n)
        claims={c['id']:c for c in audit['claims']}
        for c in claims.values():
            assert c['type'] in ['provided_fact','external_background','attributed_view','inference','method']
            for premise in c['premise_claim_ids']:assert premise in claims
            if c['type']=='inference':assert c['premise_claim_ids']
        m=f['mainline'];assert n in m['dates'] and len(m['dates'])>=3 and 2<=len(m['stages'])<=3
        assert sorted(e for s in m['stages'] for e in s['events'])==sorted(m['events']);assert all(dt(d)<=dt(n) for d in m['dates'])
        docfile=out/f'official/{n}/{n}日报.docx';doc=Document(docfile);txt='\n'.join(p.text for p in doc.paragraphs);headings=['一、重要军事动向','二、重点态势分析','三、外（台）机舰活动情况','四、美重要武器平台部署情况','五、次日关注重点'];assert all(h in txt for h in headings);assert [txt.index(h) for h in headings]==sorted(txt.index(h) for h in headings);assert '10月20日' not in txt and '2025.10.20' not in txt
        for i in range(1,5):assert f'附图{i}' in txt
        for e in ee:assert f"[{e['ref']}]" in txt and e['facts']['platform'] in txt and e['facts']['region'] in txt
        xfile=out/f'official/{n}/{n}机舰活动明细.xlsx';w=load_workbook(xfile,data_only=False);specs=load(ad/'workbook_input.json');assert w.sheetnames==['飞机','舰艇']
        for spec in specs:
            s=w[spec['name']];assert [c.value for c in s[1]]==spec['headers'];assert s.max_row==len(spec['rows'])+1 and s.max_column==len(spec['headers'])
            actual=[[c.value for c in row] for row in list(s.rows)[1:]];assert actual==spec['rows'];qcol=7 if s.title=='飞机' else 5
            assert sum(r[qcol-1] for r in actual)==sum(e['facts']['quantity'] for e in ee if e['kind']==s.title)
            for row in list(s.rows)[1:]:
                assert row[0].value==n and row[0].data_type=='s';assert row[qcol-1].data_type=='n'
                if s.title=='飞机':assert row[5].data_type=='n';assert all(row[k].value is None for k in [7,8,9,10])
                else:assert row[9].value is None;assert row[5].value is None or re.fullmatch(r'\d{4}-\d{4}',row[5].value)
                assert not any(c.data_type=='f' for c in row)
        w.close();cdata=load(ad/'chart_data.json');assert set(cdata['event_ids'])==ids;assert set(e for g in cdata['groups'] for e in g['event_ids'])==ids
        for r in cdata['hourly']:
            for kind,key in [('飞机','air'),('舰艇','ship')]:
                expect=sum(e['facts']['quantity'] for e in ee if e['kind']==kind and e['facts']['start'] and e['facts']['end'] and charts.minute(e['facts']['start'])<=r['hour']*60<=charts.minute(e['facts']['end']));assert r[key]==expect
        expected={f'{n}日报.docx',f'{n}机舰活动明细.xlsx'}|{'附件图表/'+x for x in charts.NAMES};files=list((out/f'official/{n}').rglob('*'));actual={p.relative_to(out/f'official/{n}').as_posix() for p in files if p.is_file()};assert actual==expected,(n,actual-expected)
        for p in sorted(p for p in files if p.is_file()):
            if p.suffix=='.png':
                im=Image.open(p);im.verify();assert im.format=='PNG'
            manifest.append({'node':n,'as_of':f['as_of'],'path':p.relative_to(out/'official').as_posix(),'sha256':sha(p),'bytes':p.stat().st_size})
        cmd=[sys.executable,str(BASE/'code/audit_sources.py'),'--catalog',str(ad/'source_catalog.json'),'--as-of',f['as_of'],'--used',*audit['used_source_ids'],'--substantive',*audit['used_source_ids'],'--output',str(ad/'source_metadata_audit.json')];r=subprocess.run(cmd,capture_output=True,text=True,encoding='utf8');assert r.returncode==0,r.stdout+r.stderr
        audit['checks']['source_metadata']='PASS_metadata_only';audit['checks']['structure_and_consistency']='PASS';audit['checks']['unchanged_fact_layer']='PASS';save(ad/'audit.json',audit)
        nodes.append({'date':n,'as_of':f['as_of'],'formal_files':6,'included_events':len(ee),'air_batches':sum(e['kind']=='飞机' for e in ee),'air_sorties':sum(e['facts']['quantity'] for e in ee if e['kind']=='飞机'),'ship_sorties':sum(e['facts']['quantity'] for e in ee if e['kind']=='舰艇'),'frozen_sha256':sha(ad/'frozen_ledger.json'),'structure':'PASS','source_metadata':'PASS_metadata_only','visual_review':audit['checks']['visual_review'],'independent_review':audit['checks'].get('independent_review','pending'),'raw_reproduction':'pending','used_source_ids':audit['used_source_ids'],'mainline':{'country':m['country'],'platform':m['platform'],'dates':m['dates']},'missing_evidence':audit['missing_evidence']})
    save(out/'submission_manifest.json',{'nodes':NODES,'file_count':len(manifest),'files':manifest,'layout':'date_dirs_at_zip_root','uploaded':False});save(out/'node_status_matrix.json',nodes)
    if package:
        p=out/'赛题2_初赛_八节点_完整候选提交包.zip'
        with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
            for row in manifest:z.write(out/'official'/row['path'],row['path'])
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None;assert set(z.namelist())=={r['path'] for r in manifest};assert len(z.namelist())==48;assert {n.split('/')[0] for n in z.namelist()}==set(NODES)
            for row in manifest:assert hashlib.sha256(z.read(row['path'])).hexdigest()==row['sha256']
        save(out/'package_validation.json',{'status':'PASS','archive':p.name,'sha256':sha(p),'members':48,'nodes':NODES,'extra_members':[],'root_date_dirs':True,'uploaded':False,'rule_layout_conflict':'组织方当前网页含日报外层与本次提供说明根日期目录冲突；修订先后未确认'})
    print(json.dumps(nodes,ensure_ascii=False,indent=2))
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=BASE);a=ap.parse_args();validate(a.output.resolve())
