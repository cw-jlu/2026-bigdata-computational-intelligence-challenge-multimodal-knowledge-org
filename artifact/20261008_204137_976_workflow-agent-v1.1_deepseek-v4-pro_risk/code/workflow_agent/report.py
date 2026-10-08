from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from .common import CORE, day, dumps, write_json
from .engine import run_agent, tool

SECTIONS = ["当日活动总览", "一、重要军事动向", "二、重点态势分析", "三、外（台）机舰活动情况", "四、美重要武器平台部署情况", "五、次日关注重点"]
AIR_COLS = ["日期", "星期", "国家地区", "机型", "批号", "批次", "架次", "起飞基地", "呼号", "起飞时间", "降落时间", "活动方向", "活动区域", "任务描述", "任务类型", "备注"]
SHIP_COLS = ["日期", "国家地区", "舰船", "批号", "艘次", "活动时间", "活动方向", "活动区域", "任务", "母港", "备注"]
FIG_NAMES = ["01_舰机侦巡数量曲线图.png", "02_中国主要机舰活动示意图.png", "03_美国主要机舰活动示意图.png", "04_其他国家主要机舰活动示意图.png"]


def val(e, k, default=""):
    v = e["fields"].get(k, {}).get("value")
    return default if v is None else v


def report_packet(events, node):
    accepted = [e for e in events if e.get("status") == "trusted" and e.get("gate", {}).get("eligible") and val(e, "direction") in CORE]
    today = [e for e in accepted if e["day"] == node]
    stats = {"events": len(today), "air_sorties": sum(val(e, "count") for e in today if e["kind"] == "飞机"),
             "ship_sorties": sum(val(e, "count") for e in today if e["kind"] == "舰船"),
             "by_direction": dict(Counter(val(e, "direction") for e in today)), "by_country": dict(Counter(val(e, "country") for e in today))}
    return {"node": node, "today": today, "history": accepted, "stats": stats}


def validate_mainline(line, packet):
    if line is None:
        return
    events = {e["event_id"]: e for e in packet["history"]}
    ids = line["event_ids"]
    if not ids or len(set(ids)) != len(ids) or any(e not in events for e in ids):
        raise ValueError("主线只可引用去重的已准入事件")
    selected = [events[e] for e in ids]
    dates = sorted({e["day"] for e in selected})
    if packet["node"] not in dates or len(dates) < 3:
        raise ValueError("主线必须节点当日存在且覆盖至少三个日期")
    if any(val(e, "country") in {"中国", "中华人民共和国"} for e in selected):
        raise ValueError("重点主线只能选择外方或台湾活动")
    if len({val(e, "country") for e in selected}) != 1 or len({val(e, "platform") for e in selected}) != 1:
        raise ValueError("第一版主线限定同一标准主体和平台；同型号不代表同一个体")
    if any((day(b) - day(a)).days >= 15 for a, b in zip(dates, dates[1:])):
        raise ValueError("主线证据间隔达到15日，第一版需拆开线索")
    phases = line["phases"]
    if not 2 <= len(phases) <= 3:
        raise ValueError("主线需两至三个阶段")
    phase_ids = [eid for p in phases for eid in p["event_ids"]]
    if len(set(phase_ids)) != len(phase_ids) or set(phase_ids) != set(ids):
        raise ValueError("阶段必须完整且不重复覆盖主线事件")
    prev = None
    for p in phases:
        ds = sorted(events[e]["day"] for e in p["event_ids"])
        if not ds or (prev and ds[0] <= prev):
            raise ValueError("阶段应按日期递增，不重叠")
        if any((day(b) - day(a)).days > 7 for a, b in zip(sorted(set(ds)), sorted(set(ds))[1:])):
            raise ValueError("阶段内相隔8日及以上的证据需区分阶段")
        prev = ds[-1]


def create_narrative(provider, settings, packet, trace):
    block = {"type": "object", "properties": {"text": {"type": "string"}, "tag": {"type": "string", "enum": ["fact", "analysis"]},
              "event_ids": {"type": "array", "items": {"type": "string"}}}, "required": ["text", "tag", "event_ids"], "additionalProperties": False}
    phase = {"type": "object", "properties": {"summary": {"type": "string"}, "event_ids": {"type": "array", "items": {"type": "string"}}}, "required": ["summary", "event_ids"], "additionalProperties": False}
    line_schema = {"type": ["object", "null"], "properties": {"judgement": {"type": "string"}, "event_ids": {"type": "array", "items": {"type": "string"}},
                   "phases": {"type": "array", "items": phase}}, "required": ["judgement", "event_ids", "phases"], "additionalProperties": False}
    section = {"type": "object", "properties": {"heading": {"type": "string", "enum": SECTIONS}, "blocks": {"type": "array", "items": block}},
               "required": ["heading", "blocks"], "additionalProperties": False}
    result = {}
    by_id = {e["event_id"]: e for e in packet["history"]}
    today_ids = {e["event_id"] for e in packet["today"]}
    groups = {}
    for event in packet["history"]:
        groups.setdefault((val(event, "country"), val(event, "platform")), []).append(event)
    chains = [{"country": country, "platform": platform, "dates": sorted({e["day"] for e in group}), "events": len(group)}
              for (country, platform), group in groups.items()
              if country not in {"中国", "中华人民共和国"} and packet["node"] in {e["day"] for e in group} and len({e["day"] for e in group}) >= 3]

    def list_history_events(country, platform, offset=0, limit=20):
        if offset < 0 or not 1 <= limit <= 30:
            raise ValueError("分页参数不合法")
        data = sorted(groups.get((country, platform), []), key=lambda e: (e["day"], e["event_id"]))
        return {"total": len(data), "events": [{"event_id": e["event_id"], "day": e["day"], "kind": e["kind"],
                "fields": {k: f["value"] for k, f in e["fields"].items()}} for e in data[offset:offset + limit]],
                "next_offset": offset + limit if offset + limit < len(data) else None}

    def list_history_candidates(offset=0, limit=20):
        if offset < 0 or not 1 <= limit <= 30:
            raise ValueError("分页参数不合法")
        return {"total": len(chains), "chains": chains[offset:offset + limit], "next_offset": offset + limit if offset + limit < len(chains) else None}

    def submit_report(mainline, sections):
        validate_mainline(mainline, packet)
        if [s["heading"] for s in sections] != SECTIONS or any(not s["blocks"] for s in sections):
            raise ValueError("必须按顺序生成六个规定部分，每部分至少一段")
        for index, sec in enumerate(sections):
            for b in sec["blocks"]:
                if not b["text"].strip():
                    raise ValueError("正文为空")
                ids = b["event_ids"]
                allowed = set(mainline["event_ids"]) if index == 2 and mainline else today_ids
                if index == 2 and mainline is None and ids:
                    raise ValueError("无合格主线时第二节应说明证据不足，不引用凑出的主线")
                if any(eid not in allowed for eid in ids):
                    raise ValueError("正文引用超出本节时间或主线边界")
                if not ids and packet["today"] and index != 2:
                    raise ValueError("非空日报中的陈述需要引用节点日事件")
                # Conservative numerical guard, not a semantic entailment checker.
                evidence = dumps([by_id[eid]["fields"] for eid in ids]) + dumps(packet["stats"] if index != 2 else {})
                allowed_numbers = {str(int(n)) for n in re.findall(r"\d+", evidence + packet["node"])}
                for n in re.findall(r"\d+", b["text"]):
                    if str(int(n)) not in allowed_numbers:
                        raise ValueError(f"正文出现无法回查的数字 {n}")
        result.update(mainline=mainline, sections=sections)
        return {"saved": True, "checks": "引用/章节/数字检查通过，语义真实性仍需评测"}

    specs = [tool("submit_report", "提交符合事实包的正文和合格跨日主线，校验失败时修订",
                  {"mainline": line_schema, "sections": {"type": "array", "items": section}}, ["mainline", "sections"])]
    specs += [tool("list_history_events", "分页查已准入的同主体/平台历史事件，只访问事实包；选主线需核对这些事件",
                   {"country": {"type": "string"}, "platform": {"type": "string"}, "offset": {"type": "integer"}, "limit": {"type": "integer"}}, ["country", "platform"]),
              tool("list_history_candidates", "分页查节点日出现且至少三个日期的主线候选摘要", {"offset": {"type": "integer"}, "limit": {"type": "integer"}})]
    # Keep all facts but strip bulky per-field evidence; IDs map to complete saved events.
    slim = {**packet, "today": [{"event_id": e["event_id"], "day": e["day"], "kind": e["kind"], "fields": {k: f["value"] for k, f in e["fields"].items()}} for e in packet["today"]],
            "history": [], "history_candidates": list_history_candidates(), "history_count": len(packet["history"])}
    system = """你是日报写作与跨日分析 Agent，只使用事实包，不添加原材料以外的事实或数字。
正文分核实 fact 与研判 analysis，每段引用事件 ID；除第二节外只反映节点当日。
重要性排序可由你判断。次日关注为基于当日事实的分析，不把未来活动写成已发生事实。
重点主线当日仍出现、外方或台湾、三海、至少三个日期、两至三个阶段。无合格主线返回 mainline=null，并说明不足。
主线阶段间8至14日需区分阶段，15日及以上拆成新线索，同型号不自动视为同一个体。
生成六部分，顺序由工具 heading 枚举给出。不要在正文写内部工具操作，不把引文 ID 数字写进 text；引用另存 event_ids。
历史只提供候选摘要；主动调用 list_history_candidates/list_history_events 查看完整已准入历史，不得把空 history 理解成没有历史。
调用 submit_report；错误时修订。"""
    run_agent(provider, settings, system, dumps(slim), specs, {"submit_report": submit_report, "list_history_events": list_history_events,
                                                            "list_history_candidates": list_history_candidates}, trace, "writer")
    if not result:
        raise ValueError("写作模型未提交完整正文，未生成标准日报")
    return result


def write_xlsx(packet, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    air = wb.active
    air.title = "飞机"
    ship = wb.create_sheet("舰艇")
    air.append(AIR_COLS)
    ship.append(SHIP_COLS)
    node = day(packet["node"])
    for e in packet["today"]:
        start, end = datetime.fromisoformat(val(e, "start")), datetime.fromisoformat(val(e, "end"))
        remark = f"{e['event_id']}；三源核验；" + "、".join(x["record_id"] for x in e["source_links"])
        if e["kind"] == "飞机":
            air.append([node, "星期" + "一二三四五六日"[node.weekday()], val(e, "country"), val(e, "platform"), str(val(e, "batch")),
                        1, val(e, "count"), val(e, "base"), val(e, "callsign"), start.strftime("%H:%M"), end.strftime("%H:%M"),
                        val(e, "direction"), val(e, "region"), val(e, "task"), val(e, "task"), remark])
        else:
            ship.append([node, val(e, "country"), val(e, "platform"), str(val(e, "batch")), val(e, "count"),
                         f"{start.isoformat()} — {end.isoformat()}", val(e, "direction"), val(e, "region"), val(e, "task"), val(e, "base"), remark])
    for ws in (air, ship):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            ws.column_dimensions[c.column_letter].width = 18 if c.value != "备注" else 48
        for row in ws.iter_rows(min_row=2):
            row[0].number_format = "yyyy-mm-dd"
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(c.value, str) and c.value.startswith(("=", "+", "-", "@")):
                    c.data_type = "s"  # Prevent formula injection without changing the displayed value.
    wb.save(path)


def write_figures(packet, folder):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    folder.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(10, 4))
    for kind, unit in (("飞机", "架"), ("舰船", "艘")):
        values = []
        for hour in range(24):
            h0 = datetime.combine(day(packet["node"]), datetime.min.time()).replace(hour=hour)
            h1 = h0.replace(minute=59, second=59, microsecond=999999)
            count = 0
            for e in packet["today"]:
                if e["kind"] != kind:
                    continue
                if not any(word in val(e, "task") for word in ("侦察", "巡逻", "侦巡", "监视", "反潜")):
                    continue
                a, b = datetime.fromisoformat(val(e, "start")), datetime.fromisoformat(val(e, "end"))
                # Preserve the declared timezone on interval endpoints.
                lo, hi = h0.replace(tzinfo=a.tzinfo), h1.replace(tzinfo=a.tzinfo)
                if a <= hi and b >= lo:
                    count += val(e, "count")
            values.append(count)
        ax.plot(range(24), values, marker=".", label=f"{kind}（{unit}）")
    ax.set(title=f"{packet['node']} 舰机侦巡数量曲线（活动时段覆盖，不代表逐小时独立观测）", xlabel="小时", ylabel="活动目标数量（架/艘）")
    ax.legend()
    ax.set_xticks(range(0, 24, 2))
    fig.tight_layout()
    fig.savefig(folder / FIG_NAMES[0], dpi=150)
    plt.close(fig)
    # Regional matrix avoids inventing precise coordinates or a fake coastline.
    for index, label in ((1, "中国"), (2, "美国"), (3, "其他国家（地区）")):
        group = [e for e in packet["today"] if (val(e, "country") in {"中国", "中华人民共和国"} if index == 1 else
                 val(e, "country") == "美国" if index == 2 else val(e, "country") not in {"中国", "中华人民共和国", "美国"})]
        fig, ax = plt.subplots(figsize=(10, max(3.5, len(group) * .45)))
        if group:
            regions = sorted({val(e, "region") for e in group})
            for j, e in enumerate(group):
                x = regions.index(val(e, "region"))
                ax.scatter(x, j, s=100 + 15 * val(e, "count"), label=e["kind"] if j == 0 else None)
                ax.annotate(f"{val(e, 'count')}{'架' if e['kind'] == '飞机' else '艘'}", (x, j), xytext=(8, 0), textcoords="offset points")
            ax.set_xticks(range(len(regions)), regions, rotation=20)
            ax.set_yticks(range(len(group)), [f"{val(e, 'country')} {val(e, 'platform')}" for e in group])
        else:
            ax.text(.5, .5, "当日无符合准入条件的活动", ha="center", transform=ax.transAxes)
            ax.set_xticks([])
            ax.set_yticks([])
        ax.set(title=f"{packet['node']} {label}主要机舰活动示意（区域矩阵，无地理坐标）", xlabel="活动区域（类别）", ylabel="主体 / 平台")
        ax.text(0, -0.3, "标记为活动事件；标注单位：飞机=架，舰船=艘；位置不表示经纬度", transform=ax.transAxes, fontsize=8)
        fig.tight_layout()
        fig.savefig(folder / FIG_NAMES[index], dpi=150, bbox_inches="tight")
        plt.close(fig)


def write_docx(packet, narrative, folder):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    node = day(packet["node"])
    doc.add_heading(f"{node.month}月{node.day}日飞机舰船活动情况", 0)
    for section in narrative["sections"]:
        doc.add_heading(section["heading"], 1)
        for b in section["blocks"]:
            refs = "、".join(b["event_ids"])
            doc.add_paragraph(f"【{'核实' if b['tag'] == 'fact' else '研判'}】{b['text']}" + (f"〔{refs}〕" if refs else ""))
    doc.add_heading("附件图表", 1)
    for i, filename in enumerate(FIG_NAMES, 1):
        doc.add_paragraph(f"附图{i}：{filename.removesuffix('.png')}")
        doc.add_picture(str(folder / "附件图表" / filename), width=Cm(15.5))
    doc.save(folder / f"{node.year}.{node.month}.{node.day}日报.docx")


def deliver(packet, narrative, out):
    from docx import Document
    from openpyxl import load_workbook
    from PIL import Image
    node = day(packet["node"])
    label = f"{node.year}.{node.month}.{node.day}"
    folder = Path(out) / "日报"
    folder.mkdir(parents=True, exist_ok=True)
    write_xlsx(packet, folder / f"{label}机舰活动明细.xlsx")
    write_figures(packet, folder / "附件图表")
    write_docx(packet, narrative, folder)
    checks = {}
    wb = load_workbook(folder / f"{label}机舰活动明细.xlsx")
    checks["sheets"] = wb.sheetnames == ["飞机", "舰艇"]
    checks["columns"] = [c.value for c in wb["飞机"][1]] == AIR_COLS and [c.value for c in wb["舰艇"][1]] == SHIP_COLS
    checks["air_total"] = sum(wsrow[6] for wsrow in wb["飞机"].iter_rows(min_row=2, values_only=True)) == packet["stats"]["air_sorties"]
    checks["ship_total"] = sum(wsrow[4] for wsrow in wb["舰艇"].iter_rows(min_row=2, values_only=True)) == packet["stats"]["ship_sorties"]
    wb.close()
    paragraphs = [p.text for p in Document(folder / f"{label}日报.docx").paragraphs]
    checks["sections"] = all(s in paragraphs for s in SECTIONS)
    checks["figure_references"] = all(any(f"附图{i}" in p for p in paragraphs) for i in range(1, 5))
    checks["png_count"] = len(list((folder / "附件图表").glob("*.png"))) == 4
    for name in FIG_NAMES:
        try:
            with Image.open(folder / "附件图表" / name) as im:
                im.verify()
            checks[name] = True
        except (OSError, ValueError):
            checks[name] = False
    validation = {"ok": all(checks.values()), "checks": checks,
                  "limits": ["结构/引用/数字校验不证明模型语义正确", "区域矩阵图尚未按官方地图示例对齐", "未进行 Word 页面视觉验收"]}
    write_json(Path(out) / "validation.json", validation)
    if not validation["ok"]:
        raise ValueError("交付校验失败，见 validation.json")
    return validation
