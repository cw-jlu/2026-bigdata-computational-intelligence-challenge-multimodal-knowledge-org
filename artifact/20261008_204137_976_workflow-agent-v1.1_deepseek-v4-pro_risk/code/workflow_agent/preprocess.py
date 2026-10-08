from __future__ import annotations

import hashlib
import math
import re
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from .common import day, digest, dumps, timestamp, write_json

LABELS = "标题|正文|主题词|抄收时间|发文时间|时间|发布时间|来源"


def parse_text(raw):
    hits = list(re.finditer(rf"^(?P<key>{LABELS})\s*[：:]\s*", raw, re.M))
    fields = {}
    spans = {}
    for i, hit in enumerate(hits):
        end = hits[i + 1].start() if i + 1 < len(hits) else len(raw)
        fields[hit["key"]] = raw[hit.end():end].strip()
        spans[hit["key"]] = [hit.end(), end]
    return fields, spans


def scalar(cell):
    value = cell.value
    if isinstance(value, (datetime,)):
        return value.isoformat()
    # Excel display format can preserve identifier leading zeros.
    if isinstance(value, int) and re.fullmatch(r"0{2,}", cell.number_format or ""):
        return str(value).zfill(len(cell.number_format))
    return value


def number(value):
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None


def distance_km(lon1, lat1, lon2, lat2):
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    return 6371 * 2 * math.asin(min(1, math.sqrt(a)))


def preprocess(root, store, start, node, out, config):
    from openpyxl import load_workbook
    root = Path(root).resolve()
    if not root.is_dir():
        raise ValueError(f"数据目录不存在: {root}")
    manifest, issues, records, tracks = [], [], [], {}
    selected = []
    for p in root.iterdir():
        if not p.is_dir():
            continue
        try:
            d = day(p.name).isoformat()
        except (ValueError, TypeError):
            continue
        if start <= d <= node:
            selected.append((d, p))
    if not selected:
        raise ValueError("时间范围内没有日期目录")
    present = {d for d, _ in selected}
    cursor = day(start)
    while cursor <= day(node):
        if cursor.isoformat() not in present:
            issues.append({"day": cursor.isoformat(), "error": "缺少日期目录"})
        cursor += timedelta(days=1)
    for d, folder in sorted(selected):
        files = []
        for source in ("开源军情", "网上舆论"):
            txt = sorted((folder / source).glob("*.txt"))
            if not txt:
                issues.append({"day": d, "source": source, "error": "缺少 TXT"})
            files.extend((source, p) for p in txt)
        books = sorted(folder.glob("*.xlsx"))
        if not books:
            issues.append({"day": d, "source": "雷达信息", "error": "缺少 XLSX"})
        files.extend(("雷达信息", p) for p in books)
        for source, path in files:
            data = path.read_bytes()
            sha = hashlib.sha256(data).hexdigest()
            rel = path.relative_to(root).as_posix()
            raw_copy = Path(out) / "raw" / (sha + path.suffix.lower())
            raw_copy.parent.mkdir(parents=True, exist_ok=True)
            if not raw_copy.exists():
                raw_copy.write_bytes(data)
            entry = {"path": rel, "source": source, "day": d, "sha256": sha, "bytes": len(data), "status": "ok"}
            manifest.append(entry)
            try:
                if path.suffix.lower() == ".txt":
                    raw = data.decode("utf-8-sig")
                    fields, spans = parse_text(raw)
                    rid = "T-" + digest([rel, sha])[:20]
                    records.append({"id": rid, "day": d, "source": source, "kind": "text", "raw": raw,
                                    "normalized": raw.replace("\r\n", "\n").replace("\r", "\n"),
                                    "title": fields.get("标题", path.stem), "fields": fields, "field_spans": spans,
                                    "flags": [] if fields.get("正文") else ["missing_body_field"],
                                    "locator": {"file": rel, "sha256": sha}, "duplicate_group": digest(raw)})
                else:
                    wb = load_workbook(path, data_only=False)
                    try:
                        for sheet in ("飞机", "舰船"):
                            if sheet not in wb.sheetnames:
                                issues.append({"file": rel, "error": f"缺少工作表 {sheet}"})
                                continue
                            ws = wb[sheet]
                            headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
                            active_headers = [h for h in headers if h]
                            if not active_headers or len(set(active_headers)) != len(active_headers):
                                issues.append({"file": rel, "sheet": sheet, "error": "表头为空或重复"})
                                continue
                            essential = {"批号", "国家地区"} | ({"目标类型", "位置时间", "经度", "纬度", "目标数量"} if sheet == "飞机" else {"机舰名", "发现时间", "数量", "活动区域"})
                            if essential - set(headers):
                                issues.append({"file": rel, "sheet": sheet, "error": "缺少字段", "fields": sorted(essential - set(headers))})
                            for row in ws.iter_rows(min_row=2):
                                if all(c.value is None for c in row):
                                    continue
                                values = {h: scalar(c) for h, c in zip(headers, row) if h}
                                formats = {h: c.number_format for h, c in zip(headers, row) if h}
                                rid = "R-" + digest([rel, sha, sheet, row[0].row])[:20]
                                flags = []
                                if any(c.data_type == "f" for c in row):
                                    flags.append("formula_not_evaluated")
                                times = {k: timestamp(v) for k, v in values.items() if "时间" in k}
                                for k, v in times.items():
                                    if v is None:
                                        flags.append(f"invalid_time:{k}")
                                    elif v[:10] != d:
                                        flags.append(f"cross_day_time:{k}")
                                lon, lat = number(values.get("经度")), number(values.get("纬度"))
                                if sheet == "飞机" and (lon is None or lat is None or not -180 <= lon <= 180 or not -90 <= lat <= 90):
                                    flags.append("invalid_coordinates")
                                qty = number(values.get("目标数量" if sheet == "飞机" else "数量"))
                                if qty is None or qty <= 0 or not qty.is_integer():
                                    flags.append("invalid_count")
                                batch = values.get("批号")
                                if batch is None or batch == "":
                                    flags.append("missing_batch")
                                tid = "TR-" + digest([rel, sha, sheet, batch if batch not in (None, "") else rid])[:20]
                                r = {"id": rid, "day": d, "source": source, "kind": sheet, "raw": dumps(values),
                                     "fields": values, "times": times, "number_formats": formats, "flags": flags,
                                     "track_id": tid, "locator": {"file": rel, "sheet": sheet, "row": row[0].row, "sha256": sha},
                                     "duplicate_group": digest([d, sheet, values])}
                                records.append(r)
                                tracks.setdefault(tid, {"id": tid, "day": d, "kind": sheet, "batch": batch, "record_ids": [], "flags": [], "segments": []})["record_ids"].append(rid)
                    finally:
                        wb.close()
            except (UnicodeDecodeError, ValueError, OSError, KeyError) as exc:
                entry["status"] = "error"
                issues.append({"file": rel, "error": type(exc).__name__})
    by_id = {r["id"]: r for r in records}
    duplicates = defaultdict(list)
    for r in records:
        duplicates[r["duplicate_group"]].append(r["id"])
    for t in tracks.values():
        points = [by_id[rid] for rid in t["record_ids"]]
        col = "位置时间" if t["kind"] == "飞机" else "发现时间"
        points.sort(key=lambda r: r["times"].get(col) or "9999")
        t["record_ids"] = [r["id"] for r in points]
        for field in ("国家地区", "目标类型", "机舰名", "任务", "目标数量", "数量", "活动区域"):
            choices = {dumps(r["fields"][field]) for r in points if field in r["fields"]}
            if len(choices) > 1:
                t["flags"].append(f"inconsistent:{field}")
        for a, b in zip(points, points[1:]):
            ta, tb = a["times"].get(col), b["times"].get(col)
            if not ta or not tb:
                continue
            try:
                hours = (datetime.fromisoformat(tb) - datetime.fromisoformat(ta)).total_seconds() / 3600
            except TypeError:
                t["flags"].append("mixed_timezone")
                continue
            coords = [number(r["fields"].get(k)) for r in (a, b) for k in ("经度", "纬度")]
            segment = {"from": a["id"], "to": b["id"], "hours": hours}
            if all(v is not None for v in coords) and not any("invalid_coordinates" in r["flags"] for r in (a, b)):
                dist = distance_km(*coords)
                speed = dist / hours if hours > 0 else None
                segment.update(distance_km=dist, speed_kmh=speed)
                if (hours == 0 and dist > 0.1) or (speed is not None and speed > config["speed_warning_kmh"][t["kind"]]):
                    t["flags"].append("position_jump_candidate")
            if hours == 0:
                t["flags"].append("same_timestamp")
            t["segments"].append(segment)
        t["flags"] = sorted(set(t["flags"]))
    previous_path = Path(out) / "manifest.json"
    if previous_path.exists():
        import json
        previous = json.loads(previous_path.read_text(encoding="utf-8"))
        if previous["fingerprint"] != digest(manifest) or previous["start"] != start or previous["node"] != node:
            raise ValueError("输入或时间边界已改变，请使用新的输出目录，避免混用旧事件")
    for r in records:
        r["duplicate_ids"] = duplicates[r["duplicate_group"]]
        store.put("records", r["id"], r)
    for t in tracks.values():
        store.put("tracks", t["id"], t)
    snapshot = {"start": start, "node": node, "root": str(root), "files": manifest,
                "fingerprint": digest(manifest), "records": len(records), "tracks": len(tracks), "issues": issues,
                "duplicate_groups": [ids for ids in duplicates.values() if len(ids) > 1]}
    write_json(Path(out) / "manifest.json", snapshot)
    write_json(Path(out) / "quality.json", {"issues": issues, "flagged_records": [{"id": r["id"], "flags": r["flags"]} for r in records if r["flags"]],
                                           "flagged_tracks": [t for t in tracks.values() if t["flags"]]})
    return snapshot
