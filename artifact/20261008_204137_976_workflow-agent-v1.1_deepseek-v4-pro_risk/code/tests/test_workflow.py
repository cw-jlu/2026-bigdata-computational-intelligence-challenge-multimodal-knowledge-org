"""Synthetic fixtures and scripted model responses. No paid API calls."""
from __future__ import annotations

import argparse
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from workflow_agent.analysis import analyze_seed
from workflow_agent.cli import execute
from workflow_agent.common import load_config
from workflow_agent.common import append_json, register_secret
from workflow_agent.engine import BudgetExceeded, run_agent, tool
from workflow_agent.preprocess import parse_text, preprocess
from workflow_agent.providers import Provider, ProviderError
from workflow_agent.report import SECTIONS, create_narrative, report_packet, validate_mainline
from workflow_agent.store import EvidenceTools, Store
from workflow_agent.trust import review_event
from workflow_agent.validation import check_event
from workflow_agent.observability import experiment_directory, summarize_trace

ROOT = Path(__file__).resolve().parents[1]


def config():
    return load_config(ROOT / "config.example.json")


def fixture(root, days=("2025.5.1",)):
    from openpyxl import Workbook
    for day in days:
        iso = "-".join(part.zfill(2) for part in day.split("."))
        folder = root / day
        for source in ("开源军情", "网上舆论"):
            p = folder / source
            p.mkdir(parents=True)
            (p / "001.txt").write_text(f"标题：{source}报道\n发文时间：{day} 10:00:00\n正文：据示例媒体报道，\n日本2架P-1在{iso}T08:00:00至{iso}T09:00:00于东海钓鱼岛海域巡逻。\n主题词：示例\n", encoding="utf-8")
        wb = Workbook()
        ws = wb.active
        ws.title = "飞机"
        ws.append(["批号", "国家地区", "目标类型", "位置时间", "经度", "纬度", "目标数量", "任务"])
        ws.append([1, "日本", "P-1", day + " 08:00:00", 123, 25, 2, "巡逻"])
        ws.append([1, "日本", "P-1", day + " 09:00:00", 123.1, 25.1, 2, "巡逻"])
        ws["A2"].number_format = ws["A3"].number_format = "0000"
        ship = wb.create_sheet("舰船")
        ship.append(["批号", "国家地区", "机舰名", "发现时间", "数量", "活动区域"])
        wb.save(folder / f"雷达信息表_{day}.xlsx")


def event_for(evidence, d):
    records = [r for r in evidence.records() if r["day"] == d]
    texts = [r for r in records if r["kind"] == "text"]
    refs = [{"record_id": r["id"], "quote": r["raw"]} for r in texts]
    values = {"country": "日本", "platform": "P-1", "count": 2, "start": d + "T08:00:00", "end": d + "T09:00:00",
              "region": "钓鱼岛海域", "direction": "东海", "task": "巡逻", "batch": "0001"}
    return {"day": d, "kind": "飞机", "source_links": [{"record_id": r["id"], "reason": "样例日期、平台和时段对应"} for r in records],
            "fields": {k: {"value": v, "risk": "low", "supports": copy.deepcopy(refs), "conflicts": [], "reason": "合成样例明确给出"} for k, v in values.items()},
            "unresolved_issues": [], "anomaly_handling": [], "boundary_reason": "合成样例仅有一次活动"}


class ScriptedProvider:
    protocol = "openai"
    config = {"model": "SCRIPTED_TEST_NOT_A_REAL_MODEL"}
    add_results = Provider.add_results

    def __init__(self, replies):
        self.replies = iter(replies)

    def complete(self, system, messages, tools):
        calls = next(self.replies, [])
        native = [{"id": f"call{i}", "type": "function", "function": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}} for i, (name, args) in enumerate(calls)]
        return {"message": {"role": "assistant", "content": "", "tool_calls": native},
                "calls": [{"id": c["id"], "name": c["function"]["name"], "arguments": c["function"]["arguments"]} for c in native], "text": "done", "usage": {}}


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name) / "data"
        self.out = Path(self.temp.name) / "out"
        fixture(self.root)
        self.store = Store(self.out / "state.sqlite")
        self.cfg = config()
        preprocess(self.root, self.store, "2025-05-01", "2025-05-01", self.out, self.cfg)
        self.tools = EvidenceTools(self.store, "2025-05-01", "2025-05-01")
        self.event = event_for(self.tools, "2025-05-01")

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def test_multiline_and_attribution_preserved(self):
        text = next(r for r in self.tools.records() if r["kind"] == "text")
        self.assertIn("据示例媒体报道", text["fields"]["正文"])
        self.assertIn("日本2架", text["fields"]["正文"])
        a, b = text["field_spans"]["正文"]
        self.assertEqual(text["fields"]["正文"], text["raw"][a:b].strip())

    def test_leading_zeros_and_original_copy(self):
        radar = next(r for r in self.tools.records() if r["kind"] == "飞机")
        self.assertEqual(radar["fields"]["批号"], "0001")
        self.assertTrue(list((self.out / "raw").glob("*.xlsx")))

    def test_future_files_not_read(self):
        fixture(self.root, ("2025.5.2",))
        manifest = preprocess(self.root, self.store, "2025-05-01", "2025-05-01", self.out, self.cfg)
        self.assertEqual(manifest["records"], 4)
        self.assertEqual(self.tools.search_records("2025-05-02")["total"], 0)

    def test_input_change_cannot_mix_cached_results(self):
        p = self.root / "2025.5.1" / "开源军情" / "001.txt"
        p.write_text(p.read_text(encoding="utf-8") + "changed", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "输入或时间边界"):
            preprocess(self.root, self.store, "2025-05-01", "2025-05-01", self.out, self.cfg)

    def test_evidence_gate_and_no_automatic_admission(self):
        self.assertTrue(check_event(self.event, self.tools, self.cfg, reviewed=True)["eligible"])
        self.assertFalse(check_event(self.event, self.tools, self.cfg)["eligible"])
        self.event["fields"]["count"]["risk"] = "high"
        self.assertFalse(check_event(self.event, self.tools, self.cfg, reviewed=True)["eligible"])

    def test_fabricated_quotes_rejected(self):
        self.event["fields"]["count"]["supports"][0]["quote"] = "不存在的原文"
        with self.assertRaisesRegex(ValueError, "引文不在"):
            check_event(self.event, self.tools, self.cfg)

    def test_copied_texts_not_two_independent_supports(self):
        refs = self.event["fields"]["count"]["supports"]
        a = self.store.get("records", refs[0]["record_id"])
        b = self.store.get("records", refs[1]["record_id"])
        b["duplicate_group"] = a["duplicate_group"]
        self.store.put("records", b["id"], b)
        self.assertFalse(check_event(self.event, self.tools, self.cfg, reviewed=True)["eligible"])

    def test_unresolved_conflict_not_overridden_by_low_risk(self):
        ref = self.event["fields"]["count"]["supports"][0]
        self.event["fields"]["count"]["conflicts"] = [{**ref, "resolved": False, "reason": "无法消解", "resolution_refs": []}]
        self.assertFalse(check_event(self.event, self.tools, self.cfg, reviewed=True)["eligible"])

    def test_radar_computation_and_unknown_knowledge(self):
        records = [r for r in self.tools.records() if r["kind"] == "飞机"]
        records.sort(key=lambda r: r["times"]["位置时间"])
        result = self.tools.compare_radar(records[0]["id"], records[1]["id"])
        self.assertEqual(result["hours"], 1)
        self.assertGreater(result["distance_km"], 0)
        self.assertEqual(self.tools.lookup_knowledge("THIS_PLATFORM_DOES_NOT_EXIST")["platforms"], [])

    def test_tool_loop_repairs_invalid_arguments(self):
        values = []
        provider = ScriptedProvider([[('save', {"value": "wrong"})], [('save', {"value": 2})], []])
        run_agent(provider, self.cfg, "test", "test", [tool("save", "save", {"value": {"type": "integer"}}, ["value"])],
                  {"save": lambda value: values.append(value) or {"ok": True}}, self.out / "trace.jsonl", "test")
        self.assertEqual(values, [2])

    def test_budget_stops_loop(self):
        provider = ScriptedProvider([[('save', {"value": 2})]] * 4)
        cfg = {**self.cfg, "max_tool_calls": 1}
        with self.assertRaises(BudgetExceeded):
            run_agent(provider, cfg, "test", "test", [tool("save", "save", {"value": {"type": "integer"}}, ["value"])],
                      {"save": lambda value: {"ok": True}}, self.out / "trace.jsonl", "test")

    def analyze(self):
        seed = self.tools.records()[0]
        provider = ScriptedProvider([[('read_record', {"record_id": seed["id"]})],
                                     [('submit_event', {"event": self.event}), ('finish_seed', {"status": "linked", "reason": "合成样例已关联"})], []])
        state = analyze_seed(provider, self.cfg, self.tools, seed, self.out / "trace.jsonl")
        self.assertTrue(state["finished"])
        return self.store.all("events")[0]

    def test_runtime_analysis_and_risk_review(self):
        event = self.analyze()
        clean = {k: v for k, v in event.items() if k not in {"status", "gate"}}
        result = {"event": clean, "risk_level": "low", "risk_reasons": [], "missing_evidence": [], "alternative_explanations": [], "recommendation": "include"}
        provider = ScriptedProvider([[('submit_assessment', {"assessment": result})], []])
        review_event({"risk": provider}, self.cfg, self.tools, event, self.out / "trace.jsonl")
        self.assertEqual(self.store.all("events")[0]["status"], "trusted")

    def test_debate_roles_independent_then_rebuttal_judge(self):
        event = self.analyze()
        clean = {k: v for k, v in event.items() if k not in {"status", "gate"}}
        result = {"event": clean, "risk_level": "low", "risk_reasons": [], "missing_evidence": [], "alternative_explanations": [], "recommendation": "include"}
        providers = {role: ScriptedProvider([[('submit_assessment', {"assessment": result})], []] * (2 if role == "support" else 1)) for role in ("support", "critic", "judge")}
        cfg = {**self.cfg, "trust_mode": "debate"}
        review = review_event(providers, cfg, self.tools, event, self.out / "trace.jsonl")
        self.assertEqual([r["role"] for r in review["rounds"]], ["support", "critic", "rebuttal", "judge"])
        starts = [json.loads(s) for s in (self.out / "trace.jsonl").read_text(encoding="utf-8").splitlines() if json.loads(s).get("stage") == "start"]
        debate = [s for s in starts if not s["task"].startswith("analysis:")]
        self.assertEqual(json.loads(debate[0]["input"])["prior_opinions"], [])
        self.assertEqual(json.loads(debate[1]["input"])["prior_opinions"], [])


class ProviderTests(unittest.TestCase):
    def provider(self, protocol, transport, extra=None):
        cfg = {"protocol": protocol, "base_url": "https://example.invalid/v1", "model": "TEST", "api_key_env": "TEST_API_KEY", "extra_body": extra or {}}
        with patch.dict(os.environ, {"TEST_API_KEY": "TEST_SECRET"}):
            return Provider(cfg, config(), transport)

    def test_openai_wire_and_tools(self):
        sent = []
        def transport(url, headers, body):
            sent.append((url, headers, body))
            return {"choices": [{"message": {"role": "assistant", "content": None, "tool_calls": [{"id": "c1", "type": "function", "function": {"name": "read", "arguments": "{}"}}]}, "finish_reason": "tool_calls"}]}
        p = self.provider("openai", transport)
        result = p.complete("system", [{"role": "user", "content": "test"}], [tool("read", "read", {})])
        self.assertEqual(sent[0][0], "https://example.invalid/v1/chat/completions")
        messages = [result["message"]]
        p.add_results(messages, [{"id": "c1", "result": {"ok": True}}])
        self.assertEqual(messages[-1]["role"], "tool")

    def test_anthropic_wire_and_tools(self):
        sent = []
        def transport(url, headers, body):
            sent.append((url, headers, body))
            return {"content": [{"type": "text", "text": "read"}, {"type": "tool_use", "id": "c1", "name": "read", "input": {}}], "stop_reason": "tool_use"}
        p = self.provider("anthropic", transport)
        result = p.complete("system", [{"role": "user", "content": "test"}], [tool("read", "read", {})])
        self.assertEqual(sent[0][0], "https://example.invalid/v1/messages")
        self.assertIn("input_schema", sent[0][2]["tools"][0])
        messages = [result["message"]]
        p.add_results(messages, [{"id": "c1", "result": {"error": "bad"}}])
        self.assertTrue(messages[-1]["content"][0]["is_error"])

    def test_deepseek_reasoning_roundtrip_and_max_tokens(self):
        p = self.provider("openai", lambda *args: {"choices": [{"message": {"role": "assistant", "content": "ok", "reasoning_content": "TEST_ONLY"}, "finish_reason": "stop"}]}, {"max_tokens": 2048})
        p.config["preserve_reasoning_content"] = True
        result = p.complete("s", [], [])
        self.assertEqual(result["message"]["reasoning_content"], "TEST_ONLY")
        self.assertNotIn("max_completion_tokens", p.request_payload("s", [], []))

    def test_truncation_rejected(self):
        p = self.provider("openai", lambda *args: {"choices": [{"message": {"role": "assistant", "content": "partial"}, "finish_reason": "length"}]})
        with self.assertRaises(ProviderError):
            p.complete("s", [], [])

    def test_missing_key_fails_before_network(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaisesRegex(ValueError, "缺少环境变量"):
                Provider({"protocol": "openai", "base_url": "https://example.invalid/v1", "model": "TEST", "api_key_env": "MISSING_TEST_KEY"}, config())

    def test_malformed_gateway_response_is_sanitized(self):
        p = self.provider("openai", lambda *args: {"error": "gateway echoed TEST_SECRET"})
        with self.assertRaises(ProviderError) as caught:
            p.complete("s", [], [])
        self.assertNotIn("TEST_SECRET", str(caught.exception))


class IntegrationTests(unittest.TestCase):
    def test_trace_has_timestamps_arguments_and_redacts_keys(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "traces.jsonl"
            register_secret("TRACE_TEST_SECRET_123")
            provider = ScriptedProvider([[('save', {"value": "TRACE_TEST_SECRET_123"})], []])
            run_agent(provider, config(), "system", "task", [tool("save", "test", {"value": {"type": "string"}}, ["value"])],
                      {"save": lambda value: {"ok": True}}, path, "analysis:test")
            raw = path.read_text(encoding="utf-8")
            self.assertNotIn("TRACE_TEST_SECRET_123", raw)
            records = [json.loads(s) for s in raw.splitlines()]
            self.assertTrue(all(e["timestamp"].endswith("+08:00") for e in records))
            result = next(e for e in records if e["stage"] == "tool_result")
            self.assertIn("arguments", result)
            self.assertIn("call_id", result)
            self.assertIn("latency_seconds", result)

    def test_native_and_normalized_usage_not_double_counted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "traces.jsonl"
            base = {"task_run_id": "r", "role": "analysis", "request_sequence": 1}
            usage = {"prompt_tokens": 10, "completion_tokens": 5}
            append_json(path, {**base, "stage": "provider_response", "body": {"usage": usage}})
            append_json(path, {**base, "stage": "model_response", "usage": usage, "latency_seconds": 1})
            summary = summarize_trace(tmp)
            self.assertEqual(summary["tokens_by_role"]["analysis"]["input_tokens"], 10)
            self.assertEqual(summary["tokens_by_role"]["analysis"]["output_tokens"], 5)

    def test_experiment_name_contains_version_and_model(self):
        cfg = config()
        cfg["providers"]["ds"]["model"] = "deepseek-v4-pro"
        cfg["roles"] = {k: "ds" for k in cfg["roles"]}
        path = experiment_directory(cfg, ".")
        self.assertIn("workflow-agent-v1.1_deepseek-v4-pro_risk", path.name)

    def test_cli_prepare_without_api(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, out = Path(tmp) / "data", Path(tmp) / "out"
            fixture(root)
            args = argparse.Namespace(input=str(root), start="2025.5.1", node="2025.5.1", output=str(out),
                                      config=str(ROOT / "config.example.json"), env_file=None, stage="prepare", trust_mode=None, max_seeds=None)
            with patch.dict(os.environ, {}, clear=True):
                summary = execute(args)
            self.assertEqual(summary["records"], 4)
            self.assertFalse((out / "traces.jsonl").exists())

    def test_end_to_end_cli_and_resume_with_scripted_model(self):
        with tempfile.TemporaryDirectory() as tmp:
            root, out = Path(tmp) / "data", Path(tmp) / "out"
            fixture(root, ("2025.5.1", "2025.5.2", "2025.5.3"))
            external_store = Store(out / "state.sqlite")
            evidence = EvidenceTools(external_store, "2025-05-01", "2025-05-03")
            call_count = []

            class ReplayClient(ScriptedProvider):
                def __init__(self, provider_config, settings):
                    self.config = {"model": "SCRIPTED_INTEGRATION_NOT_REAL_AI"}

                def complete(self, system, messages, tools):
                    if len(messages) > 1:
                        return ScriptedProvider([[]]).complete(system, messages, tools)
                    names = {s["name"] for s in tools}
                    task = messages[0]["content"]
                    call_count.append(sorted(names))
                    if "submit_event" in names:
                        seed = json.loads(task.split("\n", 1)[1])
                        event = event_for(evidence, seed["day"])
                        existing = [e for e in external_store.all("events") if e["day"] == seed["day"]]
                        if existing:
                            event["event_id"] = existing[0]["event_id"]
                        calls = [("read_record", {"record_id": seed["id"]}), ("submit_event", {"event": event}),
                                 ("finish_seed", {"status": "linked", "reason": "仅用于合成测试"})]
                    elif "submit_assessment" in names:
                        event = json.loads(task)["candidate"]
                        calls = [("submit_assessment", {"assessment": {"event": event, "risk_level": "low", "risk_reasons": [],
                                  "missing_evidence": [], "alternative_explanations": [], "recommendation": "include"}})]
                    elif "submit_ranking" in names:
                        calls = [("submit_ranking", {"ordered_ids": [e["event_id"] for e in json.loads(task)], "reason": "合成测试排序"})]
                    else:
                        packet = json.loads(task)
                        history = sorted(external_store.all("events"), key=lambda e: e["day"])
                        ids = [e["event_id"] for e in history]
                        today = [e["event_id"] for e in packet["today"]]
                        line = {"judgement": "样例平台连续出现，不能认定为同一个体。", "event_ids": ids,
                                "phases": [{"summary": "开始", "event_ids": ids[:1]}, {"summary": "延续", "event_ids": ids[1:]}]}
                        sections = [{"heading": heading, "blocks": [{"text": "日本P-1在东海巡逻。", "tag": "analysis" if i in {2, 5} else "fact", "event_ids": ids if i == 2 else today}]} for i, heading in enumerate(SECTIONS)]
                        calls = [("submit_report", {"mainline": line, "sections": sections})]
                    return ScriptedProvider([calls]).complete(system, messages, tools)

            args = argparse.Namespace(input=str(root), start="2025.5.1", node="2025.5.3", output=str(out),
                                      config=str(ROOT / "config.example.json"), env_file=None, stage="all", trust_mode=None, max_seeds=None)
            try:
                with patch("workflow_agent.cli.Provider", ReplayClient):
                    summary = execute(args)
                    self.assertTrue(summary["report_generated"])
                    self.assertEqual(summary["events"]["trusted"], 3)
                    validation = json.loads((out / "validation.json").read_text(encoding="utf-8"))
                    self.assertTrue(validation["ok"])
                    self.assertEqual(len(list((out / "日报" / "附件图表").glob("*.png"))), 4)
                    self.assertTrue((out / "日报" / "2025.5.3日报.docx").exists())
                    before = sum("submit_event" in calls for calls in call_count)
                    execute(args)
                    after = sum("submit_event" in calls for calls in call_count)
                    self.assertEqual(before, after)
                    traces = (out / "traces.jsonl").read_text(encoding="utf-8")
                    self.assertNotIn("TEST_SECRET", traces)
            finally:
                external_store.close()

    def test_bad_mainline_cannot_use_future_or_less_than_three_dates(self):
        packet = {"node": "2025-05-01", "history": []}
        with self.assertRaises(ValueError):
            validate_mainline({"event_ids": ["FAKE"], "phases": []}, packet)


if __name__ == "__main__":
    unittest.main()
