from common import *
from copy import deepcopy
from collections import Counter
import argparse,re,subprocess,sys
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
import charts
import source_config
def locality_note(e):
    accepted=e['facts']['region'];groups={}
    for source,a in e['claims'].items():
        reg=a.get('region')
        if reg and reg!=accepted and reg not in ['东海','黄海','南海']:
            groups.setdefault(reg,[]).append(source)
    return '；'.join('单源地点线索（未独立核实）：'+('军情' if sources[0]=='开源军情' else '舆论')+'称'+reg for reg,sources in sorted(groups.items()) if len(sources)==1)
def item(e,ref=None):
    f=e['facts'];unit='架次' if e['kind']=='飞机' else '艘次';tm=f"，核实活动窗口{f['start']}—{f['end']}" if f['start'] and f['end'] else '，活动时间未充分核实'
    return f"{f['platform']}{f['quantity']}{unit}，在{f['region']}"+(f"执行{f['task']}" if f['task'] else '活动，任务未充分核实')+tm+f"[{ref or e.get('ref','')}]"
def source_locator(e):return [{'file':a['path'],'locator':a['locator'],'sha256':a['sha256'],'snippet':a['snippet']} for a in e['claims'].values()]+[{'file':e['radar_file'],'sheet':e['sheet'],'rows':[r['row'] for r in e['rows']]}]
def compact_refs(refs):
    nums=sorted({int(x[1:]) for x in refs});groups=[]
    for x in nums:
        if groups and x==groups[-1][-1]+1:groups[-1].append(x)
        else:groups.append([x])
    return '、'.join(f'H{g[0]:02d}—H{g[-1]:02d}' if len(g)>=3 else '、'.join(f'H{x:02d}' for x in g) for g in groups)
def analysis(f,n,used,out=BASE):
    snap=load(out/f'audits/{n}/analysis_snapshot.json');assert snap['frozen_sha256']==sha(out/f'audits/{n}/frozen_ledger.json')
    hh=sorted([deepcopy(e) for e in f['verified_events'] if e['id'] in f['mainline']['events']],key=lambda e:(dt(e['date']),e['start'] or 0,e['id']));refs={e['id']:f'H{i:02d}' for i,e in enumerate(hh,1)};claims=[];blocks=[]
    for e in hh:claims.append({'id':refs[e['id']],'type':'provided_fact','text':e['date']+' '+item(e,refs[e['id']]),'source_ids':[e['id']],'locator':source_locator(e),'premise_claim_ids':[],'limits':['仅该证据日期，间隔不填补'],'event_id':e['id']})
    for row in snap['blocks']:
        pp=[refs[x] for x in row['premise_event_ids']];text=row['text']+('['+compact_refs(pp)+']' if pp else '')
        blocks.append((row['id'],text));ext=row['type']=='external_background'
        claims.append({'id':row['id'],'type':row['type'],'text':text,'source_ids':used if ext else [],'locator':[source_config.SOURCES[x]['locator'] for x in used] if ext else [],'premise_claim_ids':pp,'limits':['解释不升级为已核实事实；不补无记录日期或机舰身份']})
    stages=[{'id':f'S{i+1}','evidence_dates':st['dates'],'event_ids':st['events'],'boundary_basis':st['boundary_basis'],'premise_claim_ids':[refs[x] for x in st['events']]} for i,st in enumerate(f['mainline']['stages'])]
    return '分析问题：'+snap['question'],blocks,claims,stages,hh,refs

def add_para(doc,text,role='body',model=None):
    p=doc.add_paragraph()
    if model is not None and model._p.pPr is not None:p._p.insert(0,deepcopy(model._p.pPr))
    r=p.add_run(text)
    if model is not None and model.runs and model.runs[0]._r.rPr is not None:r._r.insert(0,deepcopy(model.runs[0]._r.rPr))
    else:
        r.font.name='黑体' if role in ['heading','title'] else '仿宋_GB2312';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),r.font.name)
    r.font.size=Pt(22 if role=='title' else 16);r.bold=role in ['heading','title'];r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    if role=='title':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
    if role=='heading':p.paragraph_format.keep_with_next=True
    return p
def build(out=BASE):
    source_config.catalog(out);structure=load(out/'evidence/sample_structure.json');sample=ROOT/'提交示例文件/日报/2025.10.20/2025.10.20日报.docx';refdoc=Document(sample);specs=[]
    for n in NODES:
        ad=out/f'audits/{n}';fp=ad/'frozen_ledger.json';before=sha(fp);f=load(fp);ee=f['today_events'];used=source_config.USED[n];dest=out/f'official/{n}';dest.mkdir(parents=True,exist_ok=True)
        f['mainline']=load(ad/'analysis_snapshot.json')['selection']
        q,blocks,claims,stages,hh,hrefs=analysis(f,n,used,out);(ad/'analysis.md').write_text('## 二、重点态势分析\n\n'+q+'\n\n'+'\n\n'.join(t for _,t in blocks),encoding='utf8')
        chartsdata=charts.draw(f,dest/'附件图表');save(ad/'chart_data.json',chartsdata)
        doc=Document(sample)
        for child in list(doc._element.body):
            if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
        date=dt(n);title=f'{date.month}月{date.day}日飞机舰船活动情况';add_para(doc,title,'title',refdoc.paragraphs[0]);foreign=[e for e in ee if e['facts']['country']!='中国'];air=[e for e in foreign if e['kind']=='飞机'];ship=[e for e in foreign if e['kind']=='舰艇']
        add_para(doc,f"{date.month}月{date.day}日，共核实外（台）飞机{len(air)}批{sum(e['facts']['quantity'] for e in air)}架次、舰艇{sum(e['facts']['quantity'] for e in ship)}艘次；含中国的全部三海活动共{len(ee)}条独立记录。统计按活动参与数量计，不等同于去重机体或舰体数。附图1展示当日核实窗口的整点参与量，时间未知记录不补造窗口。",model=refdoc.paragraphs[1])
        add_para(doc,'一、重要军事动向','heading',refdoc.paragraphs[2]);important=sorted(ee,key=lambda e:(e['id'] not in set(f['mainline']['events']),not('号' in e['facts']['platform']),e['facts']['country']!='中国',e['ref']))[:4]
        for i,e in enumerate(important,1):add_para(doc,f'（{i}）'+e['facts']['country']+item(e)+'。',model=refdoc.paragraphs[1])
        china=[e for e in ee if e['facts']['country']=='中国'];add_para(doc,'中国当日核实活动：'+('；'.join(item(e) for e in china) if china else '无满足三源核验条件的三海活动记录')+'。中国主要机舰活动见附图2。',model=refdoc.paragraphs[1])
        add_para(doc,'二、重点态势分析','heading',refdoc.paragraphs[7]);add_para(doc,q,model=refdoc.paragraphs[1])
        for _,text in blocks:add_para(doc,text,model=refdoc.paragraphs[1])
        add_para(doc,'本节引用与跨日证据索引','heading',refdoc.paragraphs[2])
        for sid in used:
            s=source_config.SOURCES[sid];loc=s['locator'].split('；')[0];p=add_para(doc,f"[{sid}] {s['publisher']}：《{s['title']}》，公开上界{s['published_by']}，{loc}。{s['url']}",model=refdoc.paragraphs[1]);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for e in hh:add_para(doc,e['date']+' '+item(e,hrefs[e['id']])+'。',model=refdoc.paragraphs[1])
        add_para(doc,'三、外（台）机舰活动情况','heading',refdoc.paragraphs[12])
        for i,c in enumerate(sorted({e['facts']['country'] for e in foreign}),1):
            es=[e for e in foreign if e['facts']['country']==c];aa=[e for e in es if e['kind']=='飞机'];ss=[e for e in es if e['kind']=='舰艇'];p=add_para(doc,f'（{i}）{c}：飞机{len(aa)}批{sum(e["facts"]["quantity"] for e in aa)}架次，舰艇{sum(e["facts"]["quantity"] for e in ss)}艘次。',model=refdoc.paragraphs[1]);p.paragraph_format.keep_with_next=True
            for e in es:add_para(doc,item(e)+'。',model=refdoc.paragraphs[1])
        add_para(doc,'其他国家及台湾活动见附图4；美国活动另见附图3。所有活动均按节点当日核实集合统计。',model=refdoc.paragraphs[1])
        add_para(doc,'四、美重要武器平台部署情况','heading',refdoc.paragraphs[29]);us=[e for e in ee if e['facts']['country']=='美国']
        add_para(doc,('当日核实的美国平台活动如下：'+'；'.join(item(e) for e in us) if us else '当日没有满足三源核验条件的美国三海平台活动')+'。这里只报告活动记载；同级舰艇名称不能识别单一舰体，现有记录不足以推断完整部署范围或确定意图。美国主要机舰活动见附图3。',model=refdoc.paragraphs[1])
        add_para(doc,'五、次日关注重点','heading',refdoc.paragraphs[31]);seeds=[e for e in ee if e['id'] in set(f['mainline']['events'])];seed=seeds[0]
        questions=load(ad/'analysis_snapshot.json')['next_focus']+f"[{seed['ref']}]"
        add_para(doc,questions,model=refdoc.paragraphs[1]);claims.append({'id':'V01','type':'inference','text':questions,'source_ids':[],'locator':[],'premise_claim_ids':[seed['ref']],'limits':['核查问题，不是已知未来事实']})
        p=add_para(doc,'F编号对应当日明细；附图1—4与正文采用同一当日核实集合。',model=refdoc.paragraphs[1]);p.paragraph_format.line_spacing=1.15
        for run in p.runs:run.font.size=Pt(12)
        for i,name in enumerate(charts.NAMES,1):
            caption=add_para(doc,f'附图{i}','heading',refdoc.paragraphs[40]);caption.paragraph_format.page_break_before=True
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            # The external PNG is authoritative; fit the complete graphic inside
            # the sample page without cropping or stretching.
            from PIL import Image
            w,h=Image.open(dest/'附件图表'/name).size;maxw=15.6;maxh=21.1;scale=min(maxw/w,maxh/h);p.add_run().add_picture(str(dest/'附件图表'/name),width=Cm(w*scale),height=Cm(h*scale))
        # Remove unreferenced October demonstration images from the copied
        # template package after replacing its body with v13 content.
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        embedded={x.get(qn('r:embed')) for x in doc._element.iter(qn('a:blip'))}
        for rid,rel in list(doc.part.rels.items()):
            if rel.reltype==RT.IMAGE and rid not in embedded:doc.part.drop_rel(rid)
        doc.save(dest/f'{n}日报.docx')
        workbook=[]
        for spec in structure['sheets']:
            rows=[]
            for e in ee:
                if e['kind']!=spec['name']:continue
                x=e['facts'];note=f"{e['ref']}；三源核验；辅助未知留空";desc=(f"{x['start']}—{x['end']}，" if x['start'] and x['end'] else '')+f"{x['country']}{x['platform']}在{x['region']}"+(f"执行{x['task']}" if x['task'] else '活动，任务未充分核实')
                # Unverified single-source localities stay in evidence audit.
                if e['kind']=='飞机':
                    if x['start'] and x['end']:note+=f"；观测窗口{x['start']}-{x['end']}，不作为起降时间"
                    rows.append([n,'星期'+'一二三四五六日'[date.weekday()],x['country'],x['platform'],e['batch'],1,x['quantity'],None,None,None,None,x['direction'],x['region'],desc,x['task'],note])
                else:rows.append([n,x['country'],x['platform'],e['batch'],x['quantity'],x['start'].replace(':','')+'-'+x['end'].replace(':','') if x['start'] and x['end'] else None,x['direction'],x['region'],x['task'],None,note])
            workbook.append({'name':spec['name'],'headers':spec['headers'],'rows':rows,'sample_widths':spec['widths']})
        save(ad/'workbook_input.json',workbook);specs.append({'node':n,'sheets':workbook})
        for e in ee:claims.append({'id':e['ref'],'type':'provided_fact','text':item(e),'source_ids':[e['id']],'locator':source_locator(e),'premise_claim_ids':[],'limits':['当日核实活动；辅助未知不补造'],'event_id':e['id']})
        audit={'as_of':f['as_of'],'mode':'historical_public_research','status':'generated_pending_current_agent_semantic_and_visual_review','read_files':load(ad/'input_manifest.json'),'frozen_fact_input':{'path':f'audits/{n}/frozen_ledger.json','sha256':before,'verification_scope':f['verification_scope']},'used_source_ids':used,'sources':[source_config.SOURCES[s] for s in used],'claims':claims,'stages':stages,'mainline_candidate_comparison':f['candidate_comparison'],'mainline_selection':f['mainline'],'missing_evidence':(['未采用能改变该节点论证的合格外部实证资料；基于比赛证据完成分析'] if not used else [])+['辅助基地、呼号、起降与母港不足两源支持；留空','缺少间隔日期完整观察及实体标识，不能证明持续驻留'],'excluded_sources': source_config.EXCLUDED,'rule_conflicts':['ZIP采用实际示例说明及本次用户入口的根日期目录','7/14/15天间隔不删除全赛段相关历史'],'suspected_upstream_issues':['宽松坐标相容性窗口是工程假设；未对所有记录取得外部专家复核'],'checks':{'source_metadata':'pending','semantic_support':'current_agent_authored_node_specific_analysis; final_review_pending','unchanged_fact_layer':'PASS' if before==sha(fp) else 'FAIL','visual_review':'pending'},'model_parameters':{'implementation':'原始证据重建＋当前代理撰写分析＋Python和artifact-tool生成；外部模型API调用0次','authoring_model':'Codex session; exact backend version and sampling unavailable','generation_seed':'not_applicable'},'future_guard':{'max_evidence_date':max(dt(e['date']).isoformat() for e in f['verified_events']),'as_of':f['as_of']}}
        save(ad/'audit.json',audit);assert sha(fp)==before
        print('BUILT',n,'events',len(ee),'mainline_dates',len(f['mainline']['dates']),flush=True)
    save(out/'workbooks.json',specs)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=BASE);args=ap.parse_args();build(args.output.resolve())
