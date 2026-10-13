"""RAW TXT/XLSX -> fresh eight fact ledgers -> hash-bound reviewed analysis replay."""
from common import *
import argparse,subprocess,shutil,sys,zipfile
from openpyxl import load_workbook
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,required=True);ap.add_argument('--node',required=True);a=ap.parse_args();out=a.output.resolve();assert not out.exists();out.mkdir(parents=True)
    log=out/'run.log'
    def run(command):
        with log.open('a',encoding='utf8') as f:
            result=subprocess.run(command,stdout=f,stderr=subprocess.STDOUT,encoding='utf8')
        if result.returncode:raise RuntimeError('reproduction failed; read '+str(log))
    run([sys.executable,'-X','utf8',str(BASE/'code/rebuild.py'),'--input',str(ROOT/'赛题2-初赛/数据'),'--output',str(out)])
    (out/'evidence').mkdir(exist_ok=True);shutil.copy2(BASE/'evidence/sample_structure.json',out/'evidence/sample_structure.json')
    facts=[]
    for n in NODES:
        p=out/f'audits/{n}/frozen_ledger.json';original=BASE/f'audits/{n}/frozen_ledger.json';snapshot=load(BASE/f'audits/{n}/analysis_snapshot.json');assert sha(p)==sha(original)==snapshot['frozen_sha256'],n+' raw facts differ'
        shutil.copy2(BASE/f'audits/{n}/analysis_snapshot.json',out/f'audits/{n}/analysis_snapshot.json');facts.append({'node':n,'raw_fact_hash':sha(p),'analysis_snapshot_hash':sha(out/f'audits/{n}/analysis_snapshot.json'),'fact_comparison':'identical','analysis_mode':'reviewed_snapshot_replay_not_new_reasoning'})
    run([sys.executable,'-X','utf8',str(BASE/'code/build.py'),'--output',str(out)])
    run([a.node,str(BASE/'code/workbooks.mjs'),str(out)])
    run([sys.executable,'-X','utf8',str(BASE/'code/verify_delivery.py'),'--output',str(out)])
    run([sys.executable,'-X','utf8',str(BASE/'code/package.py'),'--output',str(out)])
    run([sys.executable,'-X','utf8',str(BASE/'code/verify_delivery.py'),'--output',str(out),'--zip',str(out/'赛题2_初赛_v13_八节点正式提交包.zip')])
    products=[]
    for p in (out/'official').rglob('*'):
        if not p.is_file():continue
        rel=p.relative_to(out/'official');original=BASE/'official'/rel
        if p.suffix=='.docx':
            with zipfile.ZipFile(p) as x,zipfile.ZipFile(original) as y:equal=x.read('word/document.xml')==y.read('word/document.xml')
            scope='DOCX body OOXML; container metadata can vary'
        elif p.suffix=='.xlsx':
            x=load_workbook(p);y=load_workbook(original);equal=x.sheetnames==y.sheetnames and all([[(c.value,c.data_type,c.number_format) for c in r] for r in x[s].rows]==[[(c.value,c.data_type,c.number_format) for c in r] for r in y[s].rows] for s in x.sheetnames);x.close();y.close();scope='all values/types/number formats'
        else:equal=sha(p)==sha(original);scope='PNG byte hash'
        assert equal,str(rel);products.append({'path':rel.as_posix(),'equal':equal,'scope':scope,'reproduced_sha256':sha(p),'original_sha256':sha(original)})
    save(out/'reproduction_comparison.json',{'facts':facts,'products':products,'status':'PASS','raw_input_rebuilt':True,'old_fact_ledger_loaded_as_generation_input':False,'reviewed_analysis_replayed':True,'new_analysis_reasoning_executed':False})
    print('RAW REPRODUCTION PASS',out)
if __name__=='__main__':main()
