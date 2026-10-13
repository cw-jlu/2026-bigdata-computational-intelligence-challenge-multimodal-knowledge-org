"""Independent ZIP/OOXML/openpyxl/Pillow path; does not invoke builders."""
from common import *
import zipfile,io,collections,argparse
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from PIL import Image
NAMES=['01_舰机侦巡数量曲线图.png','02_中国主要机舰活动示意图.png','03_美国主要机舰活动示意图.png','04_其他国家主要机舰活动示意图.png']
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
def verify(out,archive=None,official_root=None):
    source_root=official_root or out/'official'
    expected=[];results=[]
    for n in NODES:expected += [f'{n}/{n}日报.docx',f'{n}/{n}机舰活动明细.xlsx']+[f'{n}/附件图表/{x}' for x in NAMES]
    if archive:
        z=zipfile.ZipFile(archive);assert sorted(z.namelist())==sorted(expected);assert z.testzip() is None
        read=lambda p:z.read(p)
    else:
        assert sorted(p.relative_to(source_root).as_posix() for p in source_root.rglob('*') if p.is_file())==sorted(expected)
        read=lambda p:(source_root/p).read_bytes()
    for n in NODES:
        f=load(out/f'audits/{n}/frozen_ledger.json');snap=load(out/f'audits/{n}/analysis_snapshot.json');assert snap['frozen_sha256']==sha(out/f'audits/{n}/frozen_ledger.json')
        raw=read(f'{n}/{n}日报.docx');parts=zipfile.ZipFile(io.BytesIO(raw));xml=ET.fromstring(parts.read('word/document.xml'));paras=[''.join(p.itertext()) for p in xml.findall('.//w:p',NS)];texts=[''.join(t.text or '' for t in p.findall('.//w:t',NS)) for p in xml.findall('.//w:p',NS)];text='\n'.join(texts)
        assert f'{dt(n).month}月{dt(n).day}日飞机舰船活动情况' in text
        for heading in ['一、重要军事动向','二、重点态势分析','三、外（台）机舰活动情况','四、美重要武器平台部署情况','五、次日关注重点']:assert heading in texts
        assert len(xml.findall('.//a:blip',NS))==4
        for num in range(1,5):assert f'附图{num}' in text
        assert '10月20日' not in text
        for e in f['today_events']:assert f"[{e['ref']}]" in text
        for b in snap['blocks']:assert b['text'] in text,b['id']
        wb=load_workbook(io.BytesIO(read(f'{n}/{n}机舰活动明细.xlsx')),data_only=False);assert wb.sheetnames==['飞机','舰艇'];cnt=0
        specifications=load(out/f'audits/{n}/workbook_input.json')
        for spec in specifications:
            s=wb[spec['name']];values=list(s.values);assert list(values[0])==spec['headers'];assert [list(r) for r in values[1:]]==spec['rows'];cnt+=len(spec['rows'])
            for row in list(s.rows)[1:]:
                assert row[0].data_type=='s' and row[0].value==n
                numeric=[5,6] if s.title=='飞机' else [4]
                for k in numeric:assert row[k].data_type=='n' and isinstance(row[k].value,int)
                assert not any(c.data_type=='f' for c in row)
            assert s.max_column==(16 if s.title=='飞机' else 11)
        assert cnt==len(f['today_events']);wb.close()
        chart=load(out/f'audits/{n}/chart_data.json');assert set(chart['event_ids'])=={e['id'] for e in f['today_events']}
        for name in NAMES:
            data=read(f'{n}/附件图表/{name}');im=Image.open(io.BytesIO(data));assert im.format=='PNG';im.verify()
        hashes={p:hashlib.sha256(read(p)).hexdigest() for p in expected if p.startswith(n+'/')}
        if archive:assert all(v==sha(out/'official'/p) for p,v in hashes.items())
        assert all(dt(e['date']).isoformat()<=f['as_of'] for e in f['verified_events'])
        results.append({'node':n,'official_files':6,'activity_records':cnt,'source_hashes':hashes,'structure_values_types_references':'PASS','semantic_truth':'not_established_by_this_code','reopened_from_archive':bool(archive),'reopened_from_physical_unpacked_directory':bool(official_root)})
    save(out/('independent_zip_verification.json' if archive else 'delivery_verification.json'),{'members':48,'nodes':results,'independent_code_path':True,'independent_human_reviewer':False})
    print('VERIFIED',48,'files',sum(x['activity_records'] for x in results),'records','ZIP' if archive else 'directory')
    return results
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=BASE);ap.add_argument('--zip',type=Path);a=ap.parse_args();verify(a.output.resolve(),a.zip)
