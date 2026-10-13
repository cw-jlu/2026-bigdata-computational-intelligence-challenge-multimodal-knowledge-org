"""Read-only historical intake; copy inspected engineering modules into v13."""
from pathlib import Path
import json, hashlib, shutil, subprocess, re, ast
from docx import Document
from openpyxl import load_workbook
BASE=Path(__file__).resolve().parents[1]
ROOT=BASE.parents[1]
OLD=ROOT/'artifact/20261012_163714_initial_eight_nodes_v12/code'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,o):
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(o,ensure_ascii=False,indent=2,default=str),encoding='utf8')
def main():
    files=['common.py','events.py','charts.py','workbooks.mjs','render_word.ps1','render_excel.ps1','raster_qa.py','raster_excel.py']
    for name in files: shutil.copy2(OLD/name,BASE/'code'/name)
    save(BASE/'audit/module_reuse.json',[{'source':(OLD/n).relative_to(ROOT).as_posix(),'sha256':sha(OLD/n),'new':f'code/{n}','purpose':'复用已检查的工程模块；events将在本轮修改，不加载历史台账'} for n in files])
    hist=json.loads((ROOT/'artifact/20261012_prompt_next_agent_v13/history_review.json').read_text(encoding='utf-8-sig'))
    save(BASE/'audit/history_review_parsed.json',hist)
    inventory=[];sections=[];scores=[]
    for folder in sorted((ROOT/'artifact').iterdir()):
        if not folder.is_dir() or folder==BASE:continue
        relevant=[]
        for sub in ['code','bdci','workflow_agent']:
            for p in (folder/sub).glob('*.py'):
                try:
                    txt=p.read_text(encoding='utf-8-sig');tree=ast.parse(txt)
                    relevant.append({'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'functions':[n.name for n in ast.walk(tree) if isinstance(n,ast.FunctionDef)],'reading':'program_parse_not_agent_full_read'})
                except Exception as e:relevant.append({'path':str(p),'error':str(e)})
        official=[]
        for sub in ['official','output','2025.5.10']:
            for p in (folder/sub).rglob('*.docx'):
                if 'inputs' in p.parts or 'qa' in p.parts:continue
                try:
                    doc=Document(p);paras=[x.text for x in doc.paragraphs];start=next((i for i,x in enumerate(paras) if x.startswith('二、')),None);end=next((i for i,x in enumerate(paras) if i>(start or 0) and x.startswith('三、')),len(paras))
                    sections.append({'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'section2':paras[start:end] if start is not None else [],'reading':'program_extracted_actual_docx'})
                    official.append(p.relative_to(folder).as_posix())
                except Exception as e:official.append({'path':str(p),'error':str(e)})
        for p in folder.rglob('*evaluation*.json'):
            if any(x in p.parts for x in ['inputs','node_modules']):continue
            try:scores.append({'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'data':json.loads(p.read_text(encoding='utf-8-sig'))})
            except Exception as e:scores.append({'path':str(p),'error':str(e)})
        inventory.append({'folder':folder.name,'source_code':relevant,'actual_docx':official,'readme':(folder/'README.md').read_text(encoding='utf-8-sig') if (folder/'README.md').exists() else None})
    save(BASE/'audit/history_inventory.json',inventory);save(BASE/'audit/actual_historical_section2.json',sections);save(BASE/'audit/score_archives.json',scores)
    ref=ROOT/'提交示例文件/日报/2025.10.20';doc=Document(ref/'2025.10.20日报.docx');wb=load_workbook(ref/'2025.10.20机舰活动明细.xlsx')
    save(BASE/'evidence/sample_structure.json',{'sections':[{k:getattr(s,k).cm for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']} for s in doc.sections],'paragraphs':[{'index':i,'text':p.text,'style':p.style.name} for i,p in enumerate(doc.paragraphs)],'sheets':[{'name':s.title,'headers':[c.value for c in s[1]],'types':[[c.data_type for c in r] for r in list(s.rows)[1:4]],'sample_values':[[c.value for c in r] for r in list(s.rows)[1:4]],'widths':{k:v.width for k,v in s.column_dimensions.items()}} for s in wb]})
    wb.close()
    instr=[]
    for root in [ROOT,ROOT.parent,ROOT.parent.parent]:
        if (root/'AGENTS.md').exists():instr.append(str(root/'AGENTS.md'))
    for root in [ROOT/'docs',ROOT/'artifact',ROOT/'赛题2-初赛',ROOT/'提交示例文件']:instr.extend(str(p) for p in root.rglob('AGENTS.md') if 'node_modules' not in p.parts)
    save(BASE/'audit/instruction_paths.json',instr)
    reading=[]
    for root in [ROOT/'docs',ROOT/'analysis_research_bundle',ROOT/'research/20261012_score_improvement_sources']:
        for p in root.rglob('*'):
            if p.is_file() and p.suffix in ['.md','.json']:
                reading.append({'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'reading':'program_full_read_hash; agent_reading_separate'})
    save(BASE/'audit/program_read_manifest.json',reading)
    print('historical folders',len(inventory),'actual sections',len(sections),'score archives',len(scores),'instructions',instr)
if __name__=='__main__':main()
