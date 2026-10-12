"""Read-only cross-version diagnosis. Version inclusion is never a truth label."""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
import json, hashlib, re, sys
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1]
ROOT = OUT.parents[1]
V9 = ROOT / 'artifact/20261012_134432_initial_eight_nodes_v9'
NODES = ['2025.5.10','2025.5.18','2025.5.30','2025.6.9','2025.6.17','2025.6.29','2025.7.9','2025.7.17']
VERSIONS = ['gpt_initial','claude','v2','v3','v4','v8','v9']
READS = {}

def read(p):
    data=p.read_bytes()
    READS[p.relative_to(ROOT).as_posix()]={'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data),'scope':'programmatic full-file read; not a claim of human sentence-by-sentence review'}
    return data

def load(p): return json.loads(read(p).decode('utf-8-sig'))
def save(name,obj):
    p=OUT/name;p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')

def risk(e):
    t=' '.join(e['issues'])
    if e['decision']=='纳入': return 'v9纳入'
    if e['decision']=='域外': return '域外'
    if '三类来源未全部关联' in t: return '关联缺口'
    if e['radar_conflicts']: return '雷达分支或记录异常'
    for f,label in [('quantity','数量冲突'),('country','主体冲突'),('task','任务冲突'),('start','时间冲突'),('end','时间冲突')]:
        if f+'明确冲突' in t:return label
    if '区域冲突' in t:return '区域冲突'
    return '关键字段支持不足'

def main():
    import io
    prior=load(V9/'evidence/historical_formal_candidate_union.json')['records']
    locs={(r['version'],r['node']):r['file'] for r in prior if r['version'] in VERSIONS}
    for n in NODES:locs['v9',n]=f'{V9.relative_to(ROOT).as_posix()}/official/{n}/{n}机舰活动明细.xlsx'
    tables={v:{} for v in VERSIONS};unmapped=[]
    candidates={e['id']:e for n in NODES for e in load(V9/f'audits/{n}/candidate_decisions.json')}
    for (v,n),rel in sorted(locs.items()):
        w=load_workbook(io.BytesIO(read(ROOT/rel)),data_only=True)
        for s in w:
            rows=list(s.values);headers=[str(x).replace('\n','') for x in rows[0]]
            for i,row in enumerate(rows[1:],2):
                if not any(x is not None for x in row):continue
                record=dict(zip(headers,row));eid=n+'-'+str(record['批号'])
                assert eid not in tables[v],(v,eid)
                tables[v][eid]={'file':rel,'sheet':s.title,'row':i,'values':record}
                if eid not in candidates:unmapped.append([v,eid])
        w.close()
    assert not unmapped,unmapped
    sets={v:set(t) for v,t in tables.items()}
    assert sets['v9']=={k for k,e in candidates.items() if e['decision']=='纳入'}
    raw_by_node={}; all_texts={}
    for n in NODES:
        fold=ROOT/'赛题2-初赛/数据'/n
        for p in sorted(fold.rglob('*.txt')):all_texts[p.relative_to(ROOT).as_posix()]=read(p).decode('utf-8-sig')
        p=next(fold.glob('*.xlsx'));w=load_workbook(io.BytesIO(read(p)),data_only=True);radar={}
        for s in w:
            rows=list(s.values);headers=rows[0]
            for i,row in enumerate(rows[1:],2):
                radar.setdefault(str(row[0]),[]).append({'row':i,**dict(zip(headers,row))})
        raw_by_node[n]=radar;w.close()
    def card(e):
        rows=raw_by_node[e['date']][e['batch']]
        assert rows==e['rows'],e['id']
        claims=[]
        for src,a in e['claims'].items():
            text=all_texts[a['path']]
            assert hashlib.sha256((ROOT/a['path']).read_bytes()).hexdigest()==a['sha256']
            claims.append({'source':src,'path':a['path'],'raw_text':text,'extracted_fields':{k:a.get(k) for k in ['country','key','region','regions','task','quantity','quantity_basis','start','end','observed']}})
        return {'id':e['id'],'node':e['date'],'versions':[v for v in VERSIONS if e['id'] in sets[v]],'v9_decision':e['decision'],'risk_bucket':risk(e),'issues':e['issues'],'facts':e['facts'],'claims':claims,'raw_radar_file':e['radar_file'],'raw_radar_rows':rows,'not_a_new_adjudication':True}
    cards=[card(e) for e in candidates.values()]
    extra=[c for c in cards if c['id'] in sets['claude']-sets['v4']]
    save('all_427_candidate_cards.json',cards)
    save('claude_extra_148.json',extra)
    differences=[];field_counts=Counter()
    for eid in sorted(sets['v4']&sets['v9']):
        a=tables['v4'][eid]['values'];b=tables['v9'][eid]['values'];diff={}
        for k in b:
            x,y=a.get(k),b.get(k)
            if x!=y:diff[k]={'v4':x,'v9':y};field_counts[k]+=1
        differences.append({'id':eid,'differences':diff})
    save('v4_v9_formal_field_differences.json',differences)
    fine=[]
    for d in differences:
        if '活动区域' not in d['differences']:continue
        e=candidates[d['id']];fine.append({'id':e['id'],'before_after':d['differences']['活动区域'],'explicit_region_values':e['field_audit']['region']['explicit_region_values'],'radar_check':e['field_audit']['region'].get('full_track_checks',[]),'action':'复核原文和地域精度；无两类充分支持时不得直接恢复v4细地名；可保留明确归属的一源线索','evidence_card':'all_427_candidate_cards.json#'+e['id']})
    save('fine_region_55_review.json',fine)
    branch=[]
    for e in candidates.values():
        if e['kind']!='飞机':continue
        groups=defaultdict(list)
        for r in raw_by_node[e['date']][e['batch']]:groups[str((r['开始经度'],r['开始纬度']))].append(r)
        if len(groups)<=1:continue
        branch.append({'id':e['id'],'groups':[{'raw_start_coordinate':k,'rows':rows} for k,rows in groups.items()],'v9_decision':e['decision'],'raw_text_regions':{s:a['region'] for s,a in e['claims'].items()},'action':'只生成分支假设；共同起点字段可拆候选，不足以证明哪支真实。两类正文不能唯一锚定、同刻实体冲突或同支异常未解时不放行。','status':'diagnostic_only_no_event_accepted'})
    save('radar_branch_hypotheses.json',branch)
    scores=load(ROOT/'artifact/20261012_all_versions_comparison/scores.json')
    feedback=load(V9/'evaluation/evaluation_initial_20261012.json')
    scores['v9']={'metrics':[feedback['metrics'][k] for k in ['completeness','accuracy','format','analysis']],'source':f'{V9.relative_to(ROOT).as_posix()}/evaluation/evaluation_initial_20261012.json','total_exact':feedback['weighted_total_exact_from_displayed_means'],'total_4dp':feedback['weighted_total_4dp'],'nodes':8}
    for v in VERSIONS:
        scores[v]={k:x for k,x in scores[v].items() if not k.startswith('v8_minus')}
        scores[v]['actual_formal_rows']=len(sets[v]);scores[v]['per_node_formal_rows']={n:sum(e.startswith(n+'-') for e in sets[v]) for n in NODES}
        assert sum(Decimal(x)*w for x,w in zip(scores[v]['metrics'],map(Decimal,['.25','.4','.15','.2'])))==Decimal(scores[v]['total_exact'])
    save('seven_version_scores.json',scores)
    all_history=set.union(*sets.values())
    summary={'scope':'56 actual XLSX; seven evaluated versions; eight target-day raw workbooks and 800 TXT read; no new platform evaluation','formal_rows_read':sum(len(x) for x in sets.values()),'all_raw_candidates':len(candidates),'all_version_union':len(all_history),'historical_union_absent':sorted(set(candidates)-all_history),'claude_contains_all_v4':sets['v4']<=sets['claude'],'claude_contains_all_v9':sets['v9']<=sets['claude'],'v9_contains_all_v4':sets['v4']<=sets['v9'],'claude_extra_to_v4':len(extra),'claude_extra_risk_buckets':dict(Counter(c['risk_bucket'] for c in extra)),'v4_v9_core_field_changes':{k:v for k,v in field_counts.items() if k not in ['日期','任务描述','备注','活动时间']},'v4_v9_all_column_changes':dict(field_counts),'v9_additions_to_v4':sorted(sets['v9']-sets['v4']),'v9_addition_quantity_basis':dict(Counter('正文显式数量至少一类' if any(a.get('quantity_basis')=='正文显式数量' for a in candidates[k]['claims'].values()) else '具名舰艇单一参加者推定' for k in sets['v9']-sets['v4'])),'radar_multiple_start_coordinate_batches':len(branch),'limits':['Risk buckets are mutually exclusive diagnosis order on v9 labels, not independently verified error counts.','Formal batch IDs identify comparison units; they do not prove same-event association.','No score is assigned to a new combination, and no inferred precision/recall is reported.','Current agent is nonblind and historical labels are visible.']}
    save('comparison_summary.json',summary)
    save('read_manifest.json',READS)
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
