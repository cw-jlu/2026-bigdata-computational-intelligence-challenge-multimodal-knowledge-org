from common import *
import io,zipfile
from docx import Document
from openpyxl import load_workbook
from PIL import Image

def main():
 rep=load(BASE/'reproduction_comparison.json');assert rep['status']=='PASS' and len(rep['nodes'])==8
 visual=load(BASE/'qa/current_agent_visual_receipt.json');assert visual['remaining_observed_defects']==[] and set(visual['docx_nodes'])==set(NODES)
 word={x['node']:x for x in load(BASE/'word_render.json')};raster={x['node']:x for x in load(BASE/'raster_records.json')};excel=load(BASE/'excel_render.json')
 manifest=load(BASE/'submission_manifest.json');assert manifest['file_count']==48
 revision=load(BASE/'evidence/v11_revision_audit.json');assert revision['restored_history_total']==14 and revision['daily_events']==205
 old=load(BASE/'evidence/v10_baseline/submission_manifest.json');oldhash={r['path']:r['sha256'] for r in old['files']}
 chart_same=[];xlsx_same=[]
 for n in NODES:
  assert sha(BASE/f'official/{n}/{n}日报.docx').lower()==word[n]['source_sha256'].lower()
  assert all(sha(BASE/f'official/{n}/{n}机舰活动明细.xlsx').lower()==r['source_sha256'].lower() for r in excel if r['node']==n)
  assert len([r for r in excel if r['node']==n])==4
  assert n in visual['xlsx_nodes'] and n in visual['chart_nodes']
  for p in (BASE/f'official/{n}/附件图表').glob('*.png'):
   rel=p.relative_to(BASE/'official').as_posix();assert sha(p)==oldhash[rel];chart_same.append(rel)
  # Workbook semantic equality against packaged v10 raw workbook input.
  baseline=load(BASE/f'evidence/v10_baseline/{n}/frozen_ledger.json')
  now=load(BASE/f'audits/{n}/frozen_ledger.json')
  assert baseline['today_events']==now['today_events']
  wa=load_workbook(BASE/f'evidence/v10_baseline/{n}/baseline.xlsx');wb=load_workbook(BASE/f'official/{n}/{n}机舰活动明细.xlsx')
  for sn in wa.sheetnames:
   assert list(wa[sn].values)==list(wb[sn].values)
   assert [[(c.data_type,c.number_format,str(c._style)) for c in r] for r in wa[sn]]==[[(c.data_type,c.number_format,str(c._style)) for c in r] for r in wb[sn]]
   assert {k:v.width for k,v in wa[sn].column_dimensions.items()}=={k:v.width for k,v in wb[sn].column_dimensions.items()}
   assert {k:v.height for k,v in wa[sn].row_dimensions.items()}=={k:v.height for k,v in wb[sn].row_dimensions.items()}
  wa.close();wb.close()
  xlsx_same.append(n)
 reopened=[]
 with zipfile.ZipFile(BASE/'赛题2_初赛_八节点_完整候选提交包.zip') as z:
  assert len(z.namelist())==48 and z.testzip() is None
  for row in manifest['files']:
   data=z.read(row['path']);assert hashlib.sha256(data).hexdigest()==row['sha256']
   suffix=Path(row['path']).suffix
   if suffix=='.docx':assert len(Document(io.BytesIO(data)).inline_shapes)==4
   elif suffix=='.xlsx':
    w=load_workbook(io.BytesIO(data));assert w.sheetnames==['飞机','舰艇'];w.close()
   else:Image.open(io.BytesIO(data)).verify()
   reopened.append(row['path'])
 save(BASE/'independent_package_check.json',{'status':'PASS','reopened_from_zip':reopened,'independence':'separate reader path, same agent, not blind semantic review'})
 rows=load(BASE/'node_status_matrix.json')
 for row in rows:
  n=row['date'];ad=BASE/f'audits/{n}';a=load(ad/'audit.json')
  a['status']='complete_targeted_revision_not_platform_evaluated'
  a['checks'].update(visual_review='all final DOCX pages and native workbook populated cells reviewed; unchanged chart hashes and prior review retained',semantic_support='three revised full-stage narratives and eight section scopes reviewed by current agent',raw_reproduction='PASS',independent_review='not_performed')
  a['revision_scope']='v11 retains v10 daily event adjudication, repairs full-history truncation and section-scope defects; no claim that all seven principles are fully solved'
  save(ad/'audit.json',a)
  save(ad/'visual_review.json',{'reviewer':'current agent','blind':False,'pages':raster[n]['pages'],'docx_sha256':sha(BASE/f'official/{n}/{n}日报.docx'),'xlsx_sha256':sha(BASE/f'official/{n}/{n}机舰活动明细.xlsx'),'current_receipt':'qa/current_agent_visual_receipt.json','unchanged_png_review':'v10 prior receipt plus identical SHA256; embedded charts viewed again in final DOCX','remaining_observed_defects':[]})
  row.update(visual_review='complete',raw_reproduction='PASS',independent_review='not_performed',docx_pages=raster[n]['pages'],platform_score=None,uploaded=False)
 save(BASE/'node_status_matrix.json',rows)
 stage={'status':'complete_local_targeted_revision','formal_file_count':48,'docx_count':8,'xlsx_count':8,'png_count':32,'docx_pages':sum(x['pages'] for x in raster.values()),'restored_historical_records':14,'daily_events':205,'daily_facts_changed':False,'mainline_history_coverage':'all related verified records from stage start, without gap truncation','relation_review':'source/object audit; no unsupported entry/exit assertion, no transit-corridor cases in supplied texts','raw_reproduction':'PASS','archive_reader_check':'PASS','platform_score':None,'uploaded':False,'independent_blind_review':False,'nodes':rows}
 save(BASE/'stage_audit.json',stage)
 save(BASE/'unchanged_daily_outputs.json',{'today_events_identical_to_v10':xlsx_same,'pngs_byte_identical_to_v10':chart_same,'note':'Same daily facts do not constitute proof of hidden truth or improved score.'})
 text='''# 初赛八节点 v11 修订说明

本版按“在v10基础上稍微修改”的要求，修复已确认的完整历史回溯和章节日期范围问题。保留v10及此前目录，新目录交付8份DOCX、8份XLSX、32张PNG，正式ZIP内为八日期目录、共48文件。工程ZIP另含代码、原始输入、审计、渲染与复现记录。

## 本次修订

取消主线“相隔15天即删除前段”的规则。缺证间隔仅限制连续性判断，不再删除早期相关可信记录。所有八节点主线都覆盖赛段起始日2025年5月1日至自身节点的全部相关可信事件，并保留其他未核实候选的审计索引；不把被排除的记录当作事实。

|节点与主线|补回历史记录|修订后的判断|
|---|---|---|
|6月17日 澳大利亚P-8A|5月4日东海反潜1条|全赛段14起活动，南海反潜重复与其他方向任务并存；15天间隔不证明连续驻留。|
|7月9日 加拿大哈利法克斯级|5月10、12、21、26日4条|全赛段11条；5月已有多方向多任务，6月下旬黄海警戒重复，不能把舰级记录写成单舰航程。|
|7月17日 法国戴高乐号|5—6月9条|全赛段13条；7月南海警戒是早期模式再现，不能写成首次转向南海或新建持续部署。|

这三节点的第二节完整改写为总体判断和三个证据阶段，五个其他节点沿用已有论证。全部外部背景与跨日证据索引移入第二节；第五节八份均改为仅由当日事实提出核查重点。H引用显示为连续编号范围以减少拥挤，审计仍逐条保存对应事实。

当日205条活动及其事实值与v10完全相同，8份XLSX保持相同填充值和格式，32张PNG与v10字节一致。14条补回的是跨日分析证据，不是新增14条节点当日活动。v10的数量证据修复、55条明确标注未独立核实的单源地点备注继续保留。

## 活动关系与通行区域复核

在已关联正文事实段中记录进入/离开的动词、对象、单一海域指代、来源原句与定位。观察窗口不改写为进出时刻；没有两类来源支持的进出关系不进入正式事实。节点当日可信事件仍按已核实区域活动纳入。审计见各节点 `relation_review.json`。

对各节点截止范围及全78天7,800篇原文检查台湾海峡、巴士海峡、巴林塘海峡、宫古水道、对马海峡，未检出这五类名称。本版没有基于通行区域邻接关系补纳活动，也不能用零案例检查声称通行路线判定能力已完整验证。

## 验证与限制

'''+f'八份日报共{stage["docx_pages"]}页，经WPS Office的Word.Application COM导出并转为PNG逐页检查；XLSX全体填充单元格经WPS四列带导出检查。打包后重新打开48个文件并核对哈希。内置render_docx.py因缺少LibreOffice失败，日志已保留，没有把WPS冒称Microsoft Office。\n\n'+'''已用 `run-all.ps1 -SkipNativeRender` 在新目录从原始输入复现全部八节点。冻结台账、分析文本、DOCX正文XML、XLSX单元格值和类型、32张PNG均一致。分析为绑定事实哈希的已审阅快照回放，不是重新调用模型独立撰文；Office容器时间戳允许不同。运行时与Office依赖见 `dependency_versions.json` 和 `requirements-lock.txt`，不随工程包分发。

本次属于局部修订，没有重写实体匹配、别名抽取、基地/港口关联或雷达噪声消解。事件仍以雷达批次组织候选，工程坐标窗仍有局限；进入/离开仅增加来源对象审计，尚不是完整路线重建系统。不能据此宣布七条原则已全部彻底满足，也没有独立盲审。

事实处理按各节点日期隔离；背景引用沿用v10同哈希历史资料，未新增联网检索，不用于覆盖赛内三源冲突。5月30日、6月9日、6月29日直接相关外部实证缺口仍保留。已知历史反馈为v4总分75.5445、v8总分71.1521、v9总分74.1754；v10和v11均没有新平台评测。本版未上传，不宣称准确性、完整性或总分提高。

复现入口：`run-all.ps1`，默认还生成WPS渲染；可指定 `-Python`、`-Node`、`-NodeModules`，`-Output`必须是新目录。全赛段主线选取见 `code/select_v11.py`，修订文字见 `code/revise_analysis_v11.py`，规则复核见 `code/review_v11.py`。历史v10快照保存在 `evidence/v10_baseline/`，v9/v10旧审计保留为历史材料，不冒充本轮新核验。

审计入口：`stage_audit.json`、`evidence/v11_revision_audit.json`、各节点 `three_module_changes.json`、`relation_review.json`、`visual_review.json`、`reproduction_comparison.json`、`independent_package_check.json`。
'''
 (BASE/'README.md').write_text(text,encoding='utf-8')
 print('Final audit complete',stage['docx_pages'],'pages',stage['formal_file_count'],'files')
if __name__=='__main__':main()
