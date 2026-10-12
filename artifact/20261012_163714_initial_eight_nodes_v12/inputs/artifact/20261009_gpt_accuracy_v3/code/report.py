"""Author DOCX and XLSX data only from adjudicated fact packs."""
import json
import sys
from pathlib import Path
from collections import Counter
from datetime import datetime
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from fusion import AIR_COLS,SHIP_COLS,NODES

CHARTS=['01_舰机侦巡数量曲线图.png','02_中国主要机舰活动示意图.png','03_美国主要机舰活动示意图.png','04_其他国家主要机舰活动示意图.png']
HEADINGS=['一、重要军事动向','二、重点态势分析','三、外（台）机舰活动情况','四、美重要武器平台部署情况','五、次日关注重点']

def hm(x): return f'{x//60:02d}:{x%60:02d}' if x is not None else '未明确'
def span(e): return hm(e['interval'][0])+'—'+hm(e['interval'][1])
def reference(c): return Path(c['file']).name.split('_')[0]
def text_event(e,label):
    qty=f"{e['quantity']}{'架次' if e['kind']=='飞机' else '艘次'}"
    mission='具体任务未得到两类来源明确支持' if e['task']=='任务未核实' else '执行'+e['task']
    return f"{e['country']}{e['platform']}（{qty}）于{span(e)}在{e['region']}活动，{mission}［{label}］。"
def table(doc,headers,rows,widths=None,keep_block=False):
    tb=doc.add_table(rows=1,cols=len(headers)); tb.style='Table Grid'; tb.autofit=False
    if widths:
        for col,width in zip(tb.columns,widths):col.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
    tb._tbl.tblPr.append(borders)
    for i,h in enumerate(headers): tb.rows[0].cells[i].text=h
    for row in rows:
        cc=tb.add_row().cells
        for i,v in enumerate(row): cc[i].text=str(v)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
        if ri==0: repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            if widths: cell.width=Cm(widths[ci])
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                if (keep_block or len(tb.rows)<=6) and ri<len(tb.rows)-1: p.paragraph_format.keep_with_next=True
                for run in p.runs: run.font.size=Pt(9); run.bold=ri==0
            if ri==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF3'); cell._tc.get_or_add_tcPr().append(shade)
    return tb

def phase_text(p,n):
    areas=sorted(p['regions'].items(),key=lambda x:(-x[1],x[0])); tasks=sorted(p['tasks'].items(),key=lambda x:(-x[1],x[0]))
    aa='、'.join(f'{k}{v}起' for k,v in areas)
    tt='、'.join(f'{k}{v}起' for k,v in tasks)
    gap=max(p['gaps'],default=0)
    tempo=f'阶段内相邻证据最大间隔{gap}天' if len(p['dates'])>1 else '本阶段仅有一个证据日期，不能推断持续驻留'
    boundary=p.get('boundary')
    reason=(f"边界依据：相邻核验日期{boundary['previous_date']}与{boundary['current_date']}之间出现{'、'.join(boundary['changes'])}；前一记录为{'、'.join(boundary['previous_signature'][0])}的{'、'.join(boundary['previous_signature'][1])}，后一记录为{'、'.join(boundary['current_signature'][0])}的{'、'.join(boundary['current_signature'][1])}。" if boundary else '以本条活动线索最早核验日期为观察起点，不认定此前不存在活动。')
    return f"第{['一','二','三'][n]}阶段（{p['start']}至{p['end']}）：已核实记录主要出现在{p.get('dominant_region',aa)}，主要核验任务为{p.get('dominant_task',tt)}。具体记录为{aa}；{tt}。{tempo}。{reason}阶段内仍含上述不同任务或细区，阶段是观察归纳，不是证明同一平台连续驻留。"

def document(pack,folder):
    day=pack['date']; date=datetime.strptime(day,'%Y.%m.%d'); events=sorted(pack['events'],key=lambda e:(e['interval'][0],e['country'],e['platform']))
    labels={e['event_id']:f'D{i+1:02d}' for i,e in enumerate(events)}
    doc=Document(); sec=doc.sections[0]
    # Bundled python-docx template has a blue Title bottom-border residue.
    for root in [doc.styles.element,doc.element]:
        for border in list(root.iter(qn('w:pBdr'))): border.getparent().remove(border)
    sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.9); sec.bottom_margin=Cm(1.8); sec.left_margin=sec.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name]; style.font.name='宋体' if name=='Normal' else '黑体'; style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
        style.font.size=Pt(10.5 if name=='Normal' else 18 if name=='Title' else 13)
    doc.styles['Normal'].paragraph_format.space_after=Pt(4); doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Normal'].paragraph_format.keep_together=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(11); doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
    footer=sec.footer.paragraphs[0]; footer.alignment=2
    footer.add_run('第 '); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field); footer.add_run(' 页')
    doc.core_properties.title=f'{date.month}月{date.day}日飞机舰船活动情况'; doc.core_properties.author='参赛工程师 GPT'
    doc.add_paragraph(doc.core_properties.title,'Title')
    counts=Counter(e['direction'] for e in events); foreign=[e for e in events if e['country']!='中国']; us=[e for e in events if e['country']=='美国']
    air=sum(e['quantity'] for e in events if e['kind']=='飞机'); ship=sum(e['quantity'] for e in events if e['kind']=='舰船')
    doc.add_paragraph(f"当日活动总览：{day}共纳入{len(events)}起三源核验核心海域活动，飞机{air}架次、舰艇{ship}艘次。其中东海{counts['东海']}起、南海{counts['南海']}起、黄海{counts['黄海']}起，外（台）活动{len(foreign)}起。统计以同一活动为单位，轨迹点不重复累计。")
    missions=Counter(e['task'] for e in events); ranked=missions.most_common(3)
    doc.add_paragraph('任务构成以'+ '、'.join(f'{k}（{v}起）' for k,v in ranked)+'为主。附图1仅统计已核实的侦巡任务，按活动开始小时计入；其他训练和任务未知事件不进入该曲线。附图2、附图3、附图4分别展示中国、美国、其他国家及台湾地区的当日核验活动。')
    doc.add_paragraph(f"核验边界：{len(pack['excluded'])}个雷达轨迹候选未进入当日明细，原因包括非核心海域、关键字段冲突或来源不足。文本缺少数量时不对舰级/机型默认填1；单一具名舰艇按实体数量1处理。飞机活动时间与起飞、降落时间分开，未核实的起降、基地与呼号留空。")
    doc.add_heading(HEADINGS[0],level=1)
    priorities=sorted(events,key=lambda e:(e['country']=='美国',any(k in e['platform'] for k in ['航空母舰','两栖攻击舰','轰炸机','电子侦察','预警']),e['quantity']),reverse=True)
    for e in priorities[:8]: doc.add_paragraph(text_event(e,labels[e['event_id']]))
    main=Counter(e['direction'] for e in foreign).most_common(1)
    if main:
        doc.add_paragraph(f"分析性判断：按当日已核验事件数，外（台）活动以{main[0][0]}方向相对集中（{main[0][1]}起）。该结论反映给定材料的核验覆盖，不能解释为实际总兵力或作战意图。中国活动时段与海域见附图2，美方与其他外（台）活动分布见附图3、附图4。")
    doc.add_heading(HEADINGS[1],level=1)
    tr=pack['trend']
    if tr:
        direction_counts=Counter(e['direction'] for e in tr['evidence']); task_counts=Counter(e['task'] for e in tr['evidence']); gap=max(tr['gaps'],default=0)
        posture='证据日期间呈连续线索' if gap<=7 else '含8至14天间隔，呈间歇性再出现线索'
        doc.add_paragraph(f"总体判断：{tr['country']}{tr['platform']}同型平台活动线索覆盖{tr['dates'][0]}至{tr['dates'][-1]}共{len(tr['dates'])}个不同日期，本节点仍有核验活动。记录涉及{'、'.join(f'{k}{v}起' for k,v in direction_counts.most_common())}，主要任务为{'、'.join(f'{k}{v}起' for k,v in task_counts.most_common(3))}。相邻证据最大间隔{gap}天，{posture}。")
        doc.add_paragraph(tr['selection_reason'])
        chosen=next(c for c in tr['candidate_comparison'] if c['country']==tr['country'] and c['platform_key']==tr['platform_key'] and c['direction']==tr['direction'])
        compared=[chosen]+[c for c in tr['candidate_comparison'] if c!=chosen][:4]
        table(doc,['当日主线候选','历史同区域/同任务比例','比较结果'],[[c['country']+c['platform']+' '+c['direction'],str(c['selection_basis']['same_region_ratio'])+' / '+str(c['selection_basis']['same_task_ratio']),('选用；具体变化支持'+str(len(tr['phases']))+'阶段' if c==chosen else '可选；综合排序次于选用线索' if c['eligible'] else c['rejection'])+'；'+str(c['date_count'])+'个日期'] for c in compared],[6,4,6.8])
        doc.add_paragraph('以下阶段只按具体区域、已核验任务变化或8—14日间隔划分。日期数量只满足至少3日的证据门槛，不决定三阶段。同型机舰在同一核心海域不同细区出现，不证明同一实体转移；区域未连续核实的日期不补写驻留或离开。')
        for i,p in enumerate(tr['phases']): doc.add_paragraph(phase_text(p,i))
        first,last=tr['phases'][0],tr['phases'][-1]
        if first['dominant_region']!=last['dominant_region']:
            change=f"较早阶段的记录主要集中于{first['dominant_region']}，节点所在阶段主要记录在{last['dominant_region']}，可观察到活动记录的具体区域分布改变。"
        elif first['dominant_task']!=last['dominant_task']:
            change=f"较早阶段主要记录{first['dominant_task']}，节点所在阶段主要记录{last['dominant_task']}，可观察到同区域活动的任务构成改变。"
        else:
            change=f"首末阶段主要区域与任务均为{last['dominant_region']}的{last['dominant_task']}，阶段间存在其他细区或任务记录；现有证据支持重复出现及内部变化，不能据此断言单向升级。"
        doc.add_paragraph('可观察变化：'+change+f"最大相邻证据间隔为{gap}日，缺证日期不等于停顿或撤离。")
        today_tasks=sorted({e['task'] for e in tr['evidence'] if e['date']==day and e['task']!='任务未核实'})
        doc.add_paragraph('受证据支持的解释：'+(f"区域与任务组合的改变与按不同任务海域组织活动相容；本节点{tr['direction']}仍有同型平台的{'、'.join(today_tasks)}记录，因此可以继续观察该方向任务组合的延续性。" if today_tasks else '本节点只能确认活动，具体任务未知，无法用平台用途解释任务变化。')+'这是对记录组合的有限解释，未证明兵力增减、协同指挥或特定战略意图。')
        doc.add_paragraph('其他可能与限制：记录变化也可能来自不同机体或舰体轮换、来源报道范围改变或三源核验覆盖差异。同一天同型平台可有多次活动，历史数量按事件计算；当前材料不足以判别这些解释，也不能把模型型号的重复出现直接归为同一实体航迹。')
        he=sorted(tr['evidence'],key=lambda e:(datetime.strptime(e['date'],'%Y.%m.%d'),e['interval'][0])); hr={e['event_id']:f'H{i+1:02d}' for i,e in enumerate(he)}
        doc.add_paragraph('跨日证据如下。军情与舆论编号分别定位该日期对应目录内的原始TXT；雷达列为原始工作表和Excel行号，完整文件名与字段支持关系保存在节点证据包。')
        rows=[]
        for e in he:
            cc=[pack['claims'][c] for c in e['claims']]
            rows.append([hr[e['event_id']],e['date'],e['region']+' '+span(e),e['task'],f"军情{reference(cc[0])}/舆论{reference(cc[1])}\n{e['kind']}行"+','.join(str(r) for r in e['radar_rows'])])
        table(doc,['证据','日期','区域与时间','任务','原始定位'],rows,[1,2.1,3.7,3.3,6.7])
    else: doc.add_paragraph('本节点未找到满足至少3个日期证据的外（台）三海主线，不补写确定趋势。')
    doc.add_heading(HEADINGS[2],level=1)
    for country in sorted(set(e['country'] for e in foreign)):
        doc.add_heading(country+'当日核验活动',level=2)
        for e in foreign:
            if e['country']==country: doc.add_paragraph(text_event(e,labels[e['event_id']]))
    if not foreign: doc.add_paragraph('当日无满足三源核验条件的外（台）核心海域事件。')
    doc.add_heading(HEADINGS[3],level=1)
    if us:
        for e in us: doc.add_paragraph(text_event(e,labels[e['event_id']]))
        doc.add_paragraph('以上部署描述限于当日核实的活动海域与时间，材料不足以确认驻地迁移或长期部署。美国活动示意见附图3。')
    else: doc.add_paragraph('当日没有通过三源核验的美国核心海域活动；不能据此认定美方不存在活动。附图3以空图说明该核验结果。')
    doc.add_heading(HEADINGS[4],level=1)
    if foreign:
        top=Counter(e['direction'] for e in foreign).most_common(2)
        for d,n in top:
            selected=[e for e in foreign if e['direction']==d]; platforms='、'.join(dict.fromkeys(e['country']+e['platform'] for e in selected))
            doc.add_paragraph(f"建议跟踪{d}方向的{platforms}。本日{n}起核验活动的时间分布为{hm(min(e['interval'][0] for e in selected))}至{hm(max(e['interval'][1] for e in selected))}；后续需同时核对区域、任务和时段，才能判断延续或间歇再出现。此为关注建议，不是对次日活动的预测。")
    else: doc.add_paragraph('建议优先补足当日未能消解的来源冲突，取得新的同事件证据后再判断延续。')
    index_heading=doc.add_heading('当日证据索引',level=1)
    index_heading.paragraph_format.page_break_before=False
    index_intro=doc.add_paragraph('以下编号仅用于定位，编号相同或相近不表示跨源对应。每个D编号对应明细备注中的事件ID；完整原文片段、支持及冲突理由另存中间证据。')
    index_intro.paragraph_format.keep_with_next=True
    rows=[]
    for e in events:
        cc=[pack['claims'][c] for c in e['claims']]
        rows.append([labels[e['event_id']],reference(cc[0]),reference(cc[1]),e['kind']+'行'+','.join(str(r) for r in e['radar_rows'])])
    table(doc,['事件','开源军情TXT','网上舆论TXT','雷达Excel定位'],rows,[1.5,2.8,2.8,9.7],keep_block=True)
    for i,file in enumerate(CHARTS,1):
        doc.add_page_break(); doc.add_heading(f'附图{i} '+file[3:-4],level=1)
        p=folder/'附件图表'/file
        with Image.open(p) as im: w,h=im.size
        width=min(16.5,20.5*w/h)
        doc.add_picture(str(p),width=Cm(width))
        doc.add_paragraph('本图仅使用'+day+'纳入事件。图中时间为材料活动窗口；未用示意坐标替代舰艇真实位置。')
    output=folder/(day+'日报.docx'); doc.save(output)
    return labels

def workbook_rows(pack,labels):
    day=pack['date']; date=datetime.strptime(day,'%Y.%m.%d'); air=[]; ship=[]
    for e in sorted(pack['events'],key=lambda e:(e['interval'][0],e['country'],e['platform'])):
        cc=[pack['claims'][x] for x in e['claims']]
        note=f"{labels[e['event_id']]} {e['event_id']}；活动{span(e)}；军情{reference(cc[0])}/舆论{reference(cc[1])}；雷达{e['kind']}行"+','.join(map(str,e['radar_rows']))
        if e['kind']=='飞机': note+='；起降/基地/呼号未双源核实'
        else: note+='；母港未双源核实'
        if e['trajectory_note']: note+='；竞争轨迹按两类文本限定分支'
        if e['task']=='任务未核实': note+='；具体任务仅单源，不作为事实'
        if e['kind']=='飞机': air.append([date.strftime('%Y-%m-%d'),'星期'+'一二三四五六日'[date.weekday()],e['country'],e['platform'],e['batch'],1,e['quantity'],'','','','',e['direction'],e['region'],e['task'],e['task'],note])
        else: ship.append([date.strftime('%Y-%m-%d'),e['country'],e['platform'],e['batch'],e['quantity'],span(e),e['direction'],e['region'],e['task'],'',note])
    return {'date':day,'sheets':[{'name':'飞机','headers':AIR_COLS,'rows':air},{'name':'舰艇','headers':SHIP_COLS,'rows':ship}]}

if __name__=='__main__':
    source,target=map(Path,sys.argv[1:3]); books=[]
    for day in NODES:
        pack=json.loads((source/f'factpack_{day}.json').read_text(encoding='utf-8')); folder=target/day; folder.mkdir(parents=True,exist_ok=True)
        labels=document(pack,folder); books.append(workbook_rows(pack,labels))
    (source/'workbook_rows.json').write_text(json.dumps(books,ensure_ascii=False,indent=2),encoding='utf-8')
    print('DOCX created:',len(books))
