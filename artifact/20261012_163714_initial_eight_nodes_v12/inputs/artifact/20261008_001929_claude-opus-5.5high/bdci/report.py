"""日报生成：{日期}日报.docx 与 {日期}机舰活动明细.xlsx。"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import lexicon as L
from .extract import parse_day
from .fusion import Event, conflict_summary
from .tracking import Mainline

AIR_COLS = ["日期", "星期", "国家地区", "机型", "批号", "批次", "架次", "起飞基地", "呼号", "起飞时间", "降落时间",
            "活动方向", "活动区域", "任务描述", "任务类型", "备注"]
SHIP_COLS = ["日期", "国家地区", "舰船", "批号", "艘次", "活动时间", "活动方向", "活动区域", "任务", "母港", "备注"]
DIR_ORDER = {"东海": 0, "南海": 1, "黄海": 2}
FIG_NAMES = ["01_舰机侦巡数量曲线图.png", "02_中国主要机舰活动示意图.png", "03_美国主要机舰活动示意图.png",
             "04_其他国家主要机舰活动示意图.png"]


def cn_date(day: str) -> str:
    d = parse_day(day)
    return f"{d.month}月{d.day}日"


def weekday(day: str) -> str:
    return L.WEEKDAYS[parse_day(day).weekday()]


def unit(e: Event) -> str:
    return "架" if e.kind == "飞机" else "艘"


def doc_no(m) -> str:
    return m.doc_id.split(":")[-1] if m else ""


def source_text(e: Event) -> str:
    parts = []
    if e.track:
        parts.append(f"雷达{e.track.sheet}表批号{e.track.batch}")
    if e.o:
        parts.append(f"开源军情{doc_no(e.o)}号")
    if e.p:
        parts.append(f"网上舆论{doc_no(e.p)}号")
    return "、".join(parts)


def remark(e: Event) -> str:
    items = [e.event_id, f"{len(e.sources)}源核验（{source_text(e)}）"]
    act = e.get("action")
    if act and act != "活动":
        items.append(f"{act}{e.direction}")
    items += conflict_summary(e)
    noise = [a for a in e.anomalies if "剔除" in a or "噪声" in a or "错配" in a]
    items += noise[:2]
    if e.kind == "舰船" and e.track and e.track.hull:
        items.append(f"舷号{e.track.hull}")
    return "；".join(items)


def task_sentence(e: Event) -> str:
    n = e.get("count") or 1
    start, end = e.get("start"), e.get("end")
    when = f"{start}至{end}" if start and end else (f"{start}起" if start else "当日")
    act = e.get("action")
    tail = {"进入": "，由外围航线进入该海域", "离开": "，任务后离开该海域", "进入并离开": "，进入该海域活动后离开"}.get(act, "")
    task = e.get("task") or "活动"
    return f"{e.get('country')}{e.get('platform')}{n}{unit(e)}于{when}在{e.get('region')}执行{task}{tail}。"


def sorted_events(events: list[Event]) -> list[Event]:
    return sorted(events, key=lambda e: (DIR_ORDER.get(e.direction, 9), e.get("start") or "99:99", e.event_id))


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def write_xlsx(node: str, events: list[Event], path: Path) -> None:
    d = parse_day(node)
    date_val = datetime(d.year, d.month, d.day)
    wb = Workbook()
    ws_air = wb.active
    ws_air.title = "飞机"
    ws_ship = wb.create_sheet("舰艇")
    ws_air.append(AIR_COLS)
    ws_ship.append(SHIP_COLS)
    for e in sorted_events(events):
        n = int(e.get("count") or 1)
        if e.kind == "飞机":
            t = e.track
            ws_air.append([date_val, weekday(node), e.get("country"), e.get("platform"), t.batch if t else "", 1, n,
                           (t.base if t else ""), e.get("callsign"),
                           e.get("start"), e.get("end"), e.direction, e.get("region"), task_sentence(e), e.get("task"),
                           remark(e)])
        else:
            t = e.track
            span = f"{e.get('start')}-{e.get('end')}" if e.get("start") and e.get("end") else e.get("start")
            ws_ship.append([date_val, e.get("country"), e.get("platform"), t.batch if t else "", n, span, e.direction,
                            e.get("region"), e.get("task"), (t.base if t else ""), remark(e)])
    head_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="999999")
    for ws, cols in ((ws_air, AIR_COLS), (ws_ship, SHIP_COLS)):
        widths = {"日期": 12, "星期": 8, "国家地区": 10, "机型": 26, "舰船": 26, "批号": 22, "批次": 6, "架次": 6, "艘次": 6,
                  "起飞基地": 22, "呼号": 12, "起飞时间": 9, "降落时间": 9, "活动时间": 13, "活动方向": 9, "活动区域": 16,
                  "任务描述": 52, "任务类型": 14, "任务": 14, "母港": 22, "备注": 70}
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[cell.column_letter].width = widths.get(c, 12)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                cell.alignment = Alignment(vertical="center", wrap_text=cols[cell.column - 1] in ("任务描述", "备注"))
            row[0].number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _set_font(run, size=12, bold=False, name="仿宋", color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def _para(doc, text="", size=12, bold=False, name="仿宋", align=None, indent=True, color=None, space_after=4):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if align:
        p.alignment = align
    if text:
        _set_font(p.add_run(text), size, bold, name, color)
    return p


def _tagged(doc, tag: str, text: str):
    """带“核实事实/分析推断”标签的段落。"""
    p = _para(doc)
    color = (0x1F, 0x4E, 0x78) if tag == "核实" else (0xA0, 0x40, 0x00)
    _set_font(p.add_run(f"【{tag}】"), 12, True, "黑体", color)
    _set_font(p.add_run(text), 12)
    return p


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    _set_font(p.add_run(text), 15, True, "黑体")
    return p


def _table(doc, header: list[str], rows: list[list], widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _set_font(c.paragraphs[0].add_run(str(h)), 10.5, True, "黑体")
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            _set_font(cells[i].paragraphs[0].add_run(str(v)), 10)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def importance(e: Event) -> float:
    s = 0.0
    c, p, task = e.get("country"), e.get("platform"), e.get("task")
    s += {"美国": 3.0, "日本": 2.0, "台湾": 2.0, "菲律宾": 1.5, "俄罗斯": 1.2}.get(c, 0.6 if c != "中国" else 1.0)
    if any(k in p for k in ("航空母舰", "两栖攻击舰", "直升机母舰")):
        s += 3.0
    if any(k in p for k in ("轰炸机", "核潜艇")):
        s += 2.0
    if any(k in p for k in ("侦察", "预警", "电子")) or "侦察" in task:
        s += 1.5
    if e.get("action") in ("进入", "离开", "进入并离开"):
        s += 1.0
    if (e.o and e.o.co_platforms) or int(e.get("count") or 1) >= 2:
        s += 0.8
    s += 0.3 * len(e.sources)
    return s


def write_docx(node: str, core: list[Event], all_events: list[Event], ml: Mainline | None, fig_dir: Path,
               path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.4)
    style = doc.styles["Normal"]
    style.font.name = "仿宋"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")

    title = f"{cn_date(node)}飞机舰船活动情况"
    _para(doc, title, 20, True, "黑体", WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=10)
    core = sorted_events(core)
    air = [e for e in core if e.kind == "飞机"]
    ship = [e for e in core if e.kind != "飞机"]
    foreign = [e for e in core if e.get("country") != "中国"]
    china = [e for e in core if e.get("country") == "中国"]
    us = [e for e in core if e.get("country") == "美国"]
    excluded = [e for e in all_events if not e.credible]
    other_dir = [e for e in all_events if e.credible and e.direction not in L.CORE_SEAS]
    n_air = sum(int(e.get("count") or 1) for e in air)
    n_ship = sum(int(e.get("count") or 1) for e in ship)
    resolved = [e for e in core if conflict_summary(e) or any("剔除" in a or "噪声" in a for a in e.anomalies)]

    # 当日活动总览 ---------------------------------------------------------
    _heading(doc, "当日活动总览")
    by_dir = Counter(e.direction for e in core)
    _tagged(doc, "核实", (
        f"{cn_date(node)}（{weekday(node)}），经开源军情、网上舆论、雷达信息表三类来源交叉核验，东海、南海、黄海方向共确认"
        f"可信机舰活动{len(core)}批，其中飞机{len(air)}批{n_air}架次、舰船{len(ship)}批{n_ship}艘次；"
        f"按方向统计，东海{by_dir.get('东海', 0)}批、南海{by_dir.get('南海', 0)}批、黄海{by_dir.get('黄海', 0)}批；"
        f"按主体统计，外（台）方{len(foreign)}批、中方{len(china)}批，其中美方{len(us)}批。"))
    _tagged(doc, "核实", (
        f"当日共形成候选事件{len(all_events)}个：可信且位于三海方向{len(core)}个（纳入日报），可信但位于其他海域"
        f"{len(other_dir)}个（按规则不纳入），因单源或关键冲突无法消解被排除{len(excluded)}个。纳入事件中有{len(resolved)}个"
        f"存在字段级冲突或雷达异常，已按“两类来源一致、第三类冲突可解释”原则处理，处理依据见明细表备注。"))
    rows = []
    for d in ("东海", "南海", "黄海"):
        a = [e for e in air if e.direction == d]
        s = [e for e in ship if e.direction == d]
        rows.append([d, len(a), sum(int(e.get("count") or 1) for e in a), len(s), sum(int(e.get("count") or 1) for e in s),
                     len([e for e in a + s if e.get("country") != "中国"])])
    rows.append(["合计", len(air), n_air, len(ship), n_ship, len(foreign)])
    _table(doc, ["方向", "飞机（批）", "飞机（架次）", "舰船（批）", "舰船（艘次）", "其中外（台）方（批）"], rows,
           [2.2, 2.4, 2.4, 2.4, 2.4, 3.6])
    _para(doc, "当日分时段在空/在海数量变化见附图1（01_舰机侦巡数量曲线图）。", 11, indent=True)
    _add_fig(doc, fig_dir / FIG_NAMES[0], "图1 舰机侦巡数量曲线图")

    # 一、重要军事动向 -------------------------------------------------------
    _heading(doc, "一、重要军事动向")
    groups = _formation_groups(core)
    ranked = sorted(groups, key=lambda g: -max(importance(e) for e in g))[:8]
    if not ranked:
        _tagged(doc, "核实", "当日三海方向未形成可信的重要军事动向。")
    for i, g in enumerate(ranked, 1):
        g = sorted(g, key=lambda e: e.get("start") or "")
        e0 = g[0]
        plats = "、".join(f"{e.get('platform')}{e.get('count') or 1}{unit(e)}" for e in g)
        start = min(e.get("start") for e in g if e.get("start")) if any(e.get("start") for e in g) else ""
        end = max(e.get("end") for e in g if e.get("end")) if any(e.get("end") for e in g) else ""
        act = e0.get("action")
        act_txt = {"进入": "，由外围航线进入", "离开": "，任务后离开", "进入并离开": "，进入后离开"}.get(act, "")
        tasks = "、".join(dict.fromkeys(e.get("task") for e in g if e.get("task"))) or "相关活动"
        ids = "、".join(e.event_id for e in g)
        srcs = "；".join(source_text(e) for e in g)
        where = e0.get("region") if e0.get("region").startswith(e0.direction) else f"{e0.direction}{e0.get('region')}"
        _tagged(doc, "核实", f"（{i}）{e0.get('country')}{plats}于{start}至{end}在{where}执行{tasks}"
                           f"{act_txt}。〔{ids}；证据：{srcs}〕")
    _para(doc, "中国、美国及其他国家机舰位置分布分别见附图2、附图3、附图4。", 11)

    # 二、重点态势分析 -------------------------------------------------------
    _heading(doc, "二、重点态势分析")
    if ml is None:
        _tagged(doc, "核实", "截至节点当日，节点当日出现的外（台）方三海活动中未发现覆盖3个及以上不同日期的可信主线。")
    else:
        _tagged(doc, "研判", "总体判断：" + ml.judgement)
        for k, ph in enumerate(ml.phases, 1):
            ds = ph.days
            day_txt = "、".join(f"{d.month}月{d.day}日" for d in ds[:12]) + ("等" if len(ds) > 12 else "")
            regs = "、".join(ph.regions[:4])
            tasks = "、".join(ph.tasks[:3]) or "相关活动"
            ids = "、".join(e.event_id for e in ph.events[:6]) + ("等" if len(ph.events) > 6 else "")
            _tagged(doc, "核实", f"第{_cn_num(k)}阶段（{ph.start.month}月{ph.start.day}日—{ph.end.month}月{ph.end.day}日）："
                                f"共{len(ds)}个日期出现（{day_txt}），主要在{'、'.join(ph.directions)}方向的{regs}活动，"
                                f"任务以{tasks}为主。〔证据事件：{ids}〕")
        gaps = [(b - a).days for a, b in zip(ml.days, ml.days[1:])]
        mx = max(gaps) if gaps else 0
        _tagged(doc, "研判", (f"该平台相邻出现间隔最长{mx}天（≤7天视为连续，8—14天视为同一主线的阶段间歇），"
                            f"阶段间{'存在方向转移' if ml.pattern == '转移' else '方向基本稳定'}；"
                            f"结合节点当日{'离开' if ml.pattern == '离开' else '仍在三海方向活动'}的事实，判断其"
                            f"{'短期内可能结束本轮部署，但不排除轮换返回' if ml.pattern == '离开' else '短期内仍将保持该方向存在'}。"))

    # 三、外（台）机舰活动情况 ------------------------------------------------
    _heading(doc, "三、外（台）机舰活动情况")
    if not foreign:
        _tagged(doc, "核实", "当日三海方向未发现外（台）方机舰可信活动。")
    for d in ("东海", "南海", "黄海"):
        evs = [e for e in foreign if e.direction == d]
        if not evs:
            continue
        a = [e for e in evs if e.kind == "飞机"]
        s = [e for e in evs if e.kind != "飞机"]
        cs = "、".join(f"{c}{n}批" for c, n in Counter(e.get("country") for e in evs).most_common())
        _para(doc, f"（{_cn_num(DIR_ORDER[d] + 1)}）{d}方向：飞机{len(a)}批{sum(int(e.get('count') or 1) for e in a)}架次，"
                   f"舰船{len(s)}批{sum(int(e.get('count') or 1) for e in s)}艘次（{cs}）。", 12, True, "楷体")
        for e in evs:
            _tagged(doc, "核实", task_sentence(e)[:-1] + f"。〔{e.event_id}；{source_text(e)}〕")

    # 四、美重要武器平台部署情况 -----------------------------------------------
    _heading(doc, "四、美重要武器平台部署情况")
    if not us:
        _tagged(doc, "核实", "当日东海、南海、黄海方向未发现经三源核验的美军重要武器平台活动。")
    else:
        kinds = Counter(e.kind for e in us)
        _tagged(doc, "核实", f"当日美军在三海方向共有{len(us)}批可信活动（飞机{kinds.get('飞机', 0)}批、舰船{kinds.get('舰船', 0)}批），"
                           f"位置分布见附图3。")
        rows = []
        for e in us:
            t = e.track
            rows.append([e.get("platform"), f"{e.get('count') or 1}{unit(e)}", f"{e.get('start')}-{e.get('end')}",
                         f"{e.direction}/{e.get('region')}", e.get("task"), (t.base if t else "")])
        _table(doc, ["平台", "数量", "时间", "方向/区域", "任务", "基地/母港"], rows, [4.2, 1.4, 2.4, 3.4, 2.4, 3.0])
        hv = [e for e in us if any(k in e.get("platform") for k in ("航空母舰", "两栖攻击舰", "侦察", "轰炸机", "预警"))]
        if hv:
            _tagged(doc, "研判", "其中" + "、".join(dict.fromkeys(e.get("platform") for e in hv)) +
                    "属于高价值平台，体现美方在相关方向保持侦察监视与前沿存在的意图。")

    # 五、次日关注重点 -------------------------------------------------------
    _heading(doc, "五、次日关注重点")
    k = 1
    if ml is not None:
        _tagged(doc, "研判", f"（{k}）关注{ml.country}{ml.platform}后续动向：该主线已持续{len(ml.days)}个日期，"
                            f"当前呈“{ml.pattern}”特征，次日重点核查其是否继续在{ml.phases[-1].directions[0]}方向活动。")
        k += 1
    entering = [e for e in foreign if e.get("action") in ("进入",)]
    if entering:
        _tagged(doc, "研判", f"（{k}）关注当日由外围进入三海的" + "、".join(
            dict.fromkeys(f"{e.get('country')}{e.get('platform')}" for e in entering[:5])) + "，判断其是否持续滞留或向纵深机动。")
        k += 1
    if us:
        _tagged(doc, "研判", f"（{k}）关注美军" + "、".join(dict.fromkeys(e.get("platform") for e in us[:5])) +
                "在相关海域的持续部署及与盟友协同情况。")
        k += 1
    big = [e for e in core if any(x in e.get("platform") for x in ("航空母舰", "两栖攻击舰"))]
    if big:
        _tagged(doc, "研判", f"（{k}）关注" + "、".join(dict.fromkeys(f"{e.get('country')}{e.get('platform')}" for e in big[:4])) +
                "等大型平台的编队行动与舰载机配合。")
        k += 1
    if k == 1:
        _tagged(doc, "研判", "（1）继续关注三海方向外（台）机舰常态化侦巡活动的变化。")
    _tagged(doc, "核实", f"（{k}）对当日因单源或冲突无法消解而排除的{len(excluded)}条线索持续跟踪，待获得第二类来源佐证后再行研判。")

    # 附图 -----------------------------------------------------------------
    _heading(doc, "附件图表")
    for i, (fn, cap) in enumerate(zip(FIG_NAMES[1:], ["图2 中国主要机舰活动示意图", "图3 美国主要机舰活动示意图",
                                                       "图4 其他国家主要机舰活动示意图"]), 2):
        _add_fig(doc, fig_dir / fn, cap, width=13)
    _para(doc, "说明：标注【核实】的内容为经至少两类来源相互支持且第三类来源无无法消解冲突的事实；标注【研判】的内容为基于"
               "核实事实的分析性推断。事件编号可在《机舰活动明细》备注及中间结果 events.jsonl 中反查原始文件与雷达记录。",
          10, indent=False, color=(0x55, 0x55, 0x55))
    doc.core_properties.title = title
    doc.core_properties.author = "联合制胜队"
    doc.save(path)


def _add_fig(doc, path: Path, caption: str, width: float = 15):
    if path.exists():
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, caption, 10.5, True, "黑体", WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def _formation_groups(events: list[Event]) -> list[list[Event]]:
    """同一篇开源军情报道的多平台编队合并为一条动向。"""
    groups: dict[str, list[Event]] = defaultdict(list)
    for e in events:
        key = e.o.doc_id if e.o and e.o.co_platforms else e.event_id
        groups[key].append(e)
    return list(groups.values())


def _cn_num(n: int) -> str:
    return "一二三四五六七八九十"[n - 1] if 1 <= n <= 10 else str(n)
