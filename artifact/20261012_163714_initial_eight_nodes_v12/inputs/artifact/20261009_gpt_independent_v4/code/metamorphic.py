"""Filename and Excel row-order invariance, plus first-node time cutoff.

These are engineering checks, not a truth-labelled accuracy measure.
Only task-local copies of raw files are modified.
"""
import json,random,shutil,argparse,hashlib
from pathlib import Path
from openpyxl import load_workbook
from events import process_day,choose_chain,date

def signature(events):
    return sorted((e['batch'],e['decision'],json.dumps(e['facts'],ensure_ascii=False,sort_keys=True)) for e in events)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--input',type=Path,required=True);p.add_argument('--artifact',type=Path,required=True);a=p.parse_args();inp=a.input.resolve();base=a.artifact.resolve();raw=inp/'2025.5.10';dest=base/'audit/metamorphic_input/2025.5.10'
    if dest.exists():raise SystemExit('Use an empty metamorphic_input directory; never delete an unknown target.')
    shutil.copytree(raw,dest)
    # Random names destroy numeric prefix and lexical ordering.
    rng=random.Random(31017)
    for source in ['开源军情','网上舆论']:
        for f in sorted((dest/source).glob('*.txt')):f.rename(f.with_name(f'{rng.getrandbits(64):016x}.txt'))
    x=next(dest.glob('*.xlsx'));w=load_workbook(x)
    for s in w:
        rr=list(s.values)[1:];rng.shuffle(rr);s.delete_rows(2,s.max_row)
        for r in rr:s.append(r)
    w.save(x);w.close()
    altered,_=process_day(dest,dest.parents[1]);original,_=process_day(raw,inp.parents[1])
    out=base/'output';saved=[json.loads(l) for l in (out/'evidence/2025.5.10.events.jsonl').read_text(encoding='utf-8').splitlines()]
    limited={}
    for d in sorted(inp.iterdir(),key=lambda p:date(p.name)):
        if date(d.name)>date('2025.5.10'):break
        limited[d.name]=process_day(d,inp.parents[1])[0]
    frozen=json.loads((out/'payloads/2025.5.10.json').read_text(encoding='utf-8'))['chain'];fresh=choose_chain('2025.5.10',limited)
    r={'seed':31017,'filename_and_row_shuffle_equal':signature(original)==signature(altered),'daily_first_node_equal_without_future_input':signature(original)==signature(saved),'trend_first_node_equal_without_future_input':fresh['events']==frozen['events'],
       'note':'变形测试证明不使用编号/首行并检查日界稳定性；不作为内容准确率。'}
    (base/'audit/metamorphic.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf-8');print(r)
    if not all(v for k,v in r.items() if k.endswith('equal') or k.endswith('input')):raise SystemExit(1)
