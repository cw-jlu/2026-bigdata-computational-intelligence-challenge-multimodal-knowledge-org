"""Package the complete GPT and Claude preliminary answers to the required layout."""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "submission"
ARTIFACT = ROOT / "artifact"
NODES = (
    "2025.5.10", "2025.5.18", "2025.5.30", "2025.6.9",
    "2025.6.17", "2025.6.29", "2025.7.9", "2025.7.17",
)
SOURCES = {
    "GPT": ARTIFACT / "20261008_001028_gpt-5.6solhigh" / "output",
    "Claude": ARTIFACT / "20261008_001929_claude-opus-5.5high" / "output",
}
CHARTS = (
    "01_舰机侦巡数量曲线图.png",
    "02_中国主要机舰活动示意图.png",
    "03_美国主要机舰活动示意图.png",
    "04_其他国家主要机舰活动示意图.png",
)
AIR_HEADERS = (
    "日期", "星期", "国家地区", "机型", "批号", "批次", "架次", "起飞基地",
    "呼号", "起飞时间", "降落时间", "活动方向", "活动区域", "任务描述", "任务类型", "备注",
)
SHIP_HEADERS = (
    "日期", "国家地区", "舰船", "批号", "艘次", "活动时间", "活动方向",
    "活动区域", "任务", "母港", "备注",
)
SECTIONS = (
    "当日活动总览", "一、重要军事动向", "二、重点态势分析",
    "三、外（台）机舰活动情况", "四、美重要武器平台部署情况", "五、次日关注重点",
)


def source_dir(model: str, node: str) -> Path:
    root = SOURCES[model]
    return root / f"{node}日报" if model == "GPT" else root / node / "日报"


def expected_files(node: str) -> tuple[str, ...]:
    return (
        f"{node}日报.docx", f"{node}机舰活动明细.xlsx",
        *(f"附件图表/{name}" for name in CHARTS),
    )


def check_day(model: str, node: str) -> dict[str, object]:
    folder = source_dir(model, node)
    expected = set(expected_files(node))
    actual = {p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file()}
    if actual != expected:
        raise ValueError(f"{model} {node}: files missing={expected - actual}, extra={actual - expected}")

    month, day = map(int, re.findall(r"\d+", node)[1:])
    doc = Document(folder / f"{node}日报.docx")
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    title = f"{month}月{day}日飞机舰船活动情况"
    if title not in paragraphs or any(section not in paragraphs for section in SECTIONS):
        raise ValueError(f"{model} {node}: DOCX title or section missing")
    explicit_chart_refs = all(f"附图{i}" in paragraphs for i in range(1, 5))
    appendix_chart_names = all(chart in paragraphs for chart in CHARTS)
    if not explicit_chart_refs and not appendix_chart_names:
        raise ValueError(f"{model} {node}: DOCX chart references missing")
    if len(doc.inline_shapes) < 4:
        raise ValueError(f"{model} {node}: DOCX has fewer than four embedded charts")

    book = load_workbook(folder / f"{node}机舰活动明细.xlsx", read_only=True, data_only=True)
    try:
        if book.sheetnames != ["飞机", "舰艇"]:
            raise ValueError(f"{model} {node}: sheet names mismatch")
        for sheet, headers in (("飞机", AIR_HEADERS), ("舰艇", SHIP_HEADERS)):
            row = next(book[sheet].iter_rows(min_row=1, max_row=1, values_only=True))
            if tuple(row) != headers:
                raise ValueError(f"{model} {node}: {sheet} headers mismatch: {row}")
        counts = {sheet: book[sheet].max_row - 1 for sheet in book.sheetnames}
    finally:
        book.close()

    for chart in CHARTS:
        with Image.open(folder / "附件图表" / chart) as image:
            if image.format != "PNG":
                raise ValueError(f"{model} {node}: chart is not PNG: {chart}")
            image.verify()
    return {"air_rows": counts["飞机"], "ship_rows": counts["舰艇"], "explicit_chart_refs": explicit_chart_refs}


def build(model: str) -> dict[str, object]:
    validations = json.loads((SOURCES[model] / "validation.json").read_text(encoding="utf-8"))
    if len(validations) != len(NODES) or any(not item["ok"] for item in validations):
        raise ValueError(f"{model}: source validation incomplete or failed")
    archive = OUTPUT / f"{model}_初赛答案.zip"
    rows: list[dict[str, object]] = []
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as output:
        for node in NODES:
            counts = check_day(model, node)
            folder = source_dir(model, node)
            for relative in expected_files(node):
                name = f"日报/{node}/{relative}"
                info = zipfile.ZipInfo(name, date_time=(2025, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                output.writestr(info, (folder / relative).read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
            rows.append({"node": node, **counts})

    wanted = [f"日报/{node}/{relative}" for node in NODES for relative in expected_files(node)]
    with zipfile.ZipFile(archive) as packaged:
        if packaged.namelist() != wanted or packaged.testzip() is not None:
            raise ValueError(f"{model}: archive contents or CRC check failed")
        if any(not info.flag_bits & 0x800 for info in packaged.infolist()):
            raise ValueError(f"{model}: archive names are not UTF-8")
        # Open files from the ZIP itself, so the delivered bytes are checked too.
        for node in NODES:
            prefix = f"日报/{node}/"
            Document(io.BytesIO(packaged.read(prefix + f"{node}日报.docx")))
            book = load_workbook(io.BytesIO(packaged.read(prefix + f"{node}机舰活动明细.xlsx")), read_only=True)
            book.close()
            for chart in CHARTS:
                with Image.open(io.BytesIO(packaged.read(prefix + "附件图表/" + chart))) as image:
                    image.verify()

    return {
        "model": model,
        "archive": archive.name,
        "sha256": hashlib.sha256(archive.read_bytes()).hexdigest(),
        "bytes": archive.stat().st_size,
        "nodes": rows,
        "files": len(wanted),
    }


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = [build(model) for model in ("GPT", "Claude")]
    (OUTPUT / "manifest.json").write_text(json.dumps(results, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("Built:", *(item["archive"] for item in results))


if __name__ == "__main__":
    main()
